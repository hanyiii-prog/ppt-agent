"""Repair Executor tests: real PPTX in, minimally mutated real PPTX out."""
import copy
import zipfile
from pathlib import Path

from ppt_agent.fidelity import extract_fidelity_dna
from ppt_agent.fidelity_diff import compare_dna
from ppt_agent.fidelity_repair import repair_plan_dict
from ppt_agent.fidelity_repair_executor import execute_repair_plan, verify_repair

from fidelity_fixtures import _inject_alpha, build_rich_pptx, mutate_pptx
from pptx import Presentation
from pptx.oxml.ns import qn


def _repair_roundtrip(tmp_path: Path, mutation: str) -> dict:
    reference = build_rich_pptx(tmp_path / "ref.pptx")
    candidate = mutate_pptx(reference, tmp_path / f"{mutation}.pptx", mutation)
    report = compare_dna(
        extract_fidelity_dna(reference, slide_index=3),
        extract_fidelity_dna(candidate, slide_index=3),
    )
    assert not report.passed
    plan = repair_plan_dict(report)
    result = execute_repair_plan(candidate, plan, tmp_path / f"{mutation}-repaired.pptx", slide_index=3)
    verification = verify_repair(reference, result["output"], slide_index=3)
    return {"result": result, "verification": verification, "reference": reference, "plan": plan}


def _repair_with_alpha_pair(tmp_path: Path) -> dict:
    reference = build_rich_pptx(tmp_path / "ref.pptx")
    prs = Presentation(str(reference))
    for shape in prs.slides[2].shapes:
        if shape.name.startswith("Rectangle") and (shape.fill.type is not None):
            _inject_alpha(shape, 65)
            break
    ref_path = tmp_path / "ref_alpha.pptx"
    prs.save(ref_path)

    candidate = mutate_pptx(ref_path, tmp_path / "cand.pptx", "geometry_changed")
    # now strip alpha on the candidate's same shape: 65 -> 40
    prs2 = Presentation(str(ref_path))
    prs3 = Presentation(str(candidate))
    for shape in prs3.slides[2].shapes:
        if shape.name.startswith("Rectangle") and (shape.fill.type is not None):
            srgb = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            alpha = srgb.find(qn("a:alpha"))
            if alpha is not None:
                alpha.set("val", "40000")
            break
    prs3.save(tmp_path / "cand_alpha.pptx")
    candidate = tmp_path / "cand_alpha.pptx"

    report = compare_dna(
        extract_fidelity_dna(ref_path, slide_index=3),
        extract_fidelity_dna(candidate, slide_index=3),
    )
    assert not report.passed
    plan = repair_plan_dict(report)
    result = execute_repair_plan(candidate, plan, tmp_path / "cand_alpha_repaired.pptx", slide_index=3)
    verification = verify_repair(ref_path, result["output"], slide_index=3)
    return {"result": result, "verification": verification}


def test_position_repair_restores_geometry(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "geometry_changed")
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]
    assert any(d["operation"] == "restore_geometry" for d in outcome["result"]["applied"])


def test_rotation_repair_restores_transform(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "rotation_changed")
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]
    assert any(d["operation"] == "restore_transform" for d in outcome["result"]["applied"])


def test_flip_repair_restores_transform(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "flip_changed")
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]


def test_zorder_repair_restores_layer_order(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "zorder_changed")
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]
    assert outcome["result"]["reorder_targets"].get("slide"), "reorder directives must be collected"


def test_crop_repair_removes_extra_crop(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "crop_changed")
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]
    assert any(d["operation"] == "restore_media_crop" for d in outcome["result"]["applied"])


def test_gradient_angle_repair(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "gradient_changed")
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]


def test_table_geometry_repair(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "table_changed")
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]


def test_alpha_repair(tmp_path: Path):
    outcome = _repair_with_alpha_pair(tmp_path)
    assert any(d["operation"] == "restore_alpha" for d in outcome["result"]["applied"]), outcome["result"]
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]


def test_font_repair(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "font_changed")
    assert outcome["verification"].passed, outcome["verification"].to_dict()["issues"][:3]


def test_non_repairable_directives_are_skipped_not_faked(tmp_path: Path):
    reference = build_rich_pptx(tmp_path / "ref.pptx")
    candidate = mutate_pptx(reference, tmp_path / "media.pptx", "media_changed")
    report = compare_dna(
        extract_fidelity_dna(reference, slide_index=3),
        extract_fidelity_dna(candidate, slide_index=3),
    )
    plan = repair_plan_dict(report)
    result = execute_repair_plan(candidate, plan, tmp_path / "media-repaired.pptx", slide_index=3)
    operations = {d["operation"] for d in plan["directives"]}
    assert "restore_media_asset" in operations
    skipped = {d["operation"] for d in result["skipped"]}
    assert "restore_media_asset" in skipped


def test_executor_never_mutates_source(tmp_path: Path):
    reference = build_rich_pptx(tmp_path / "ref.pptx")
    candidate = mutate_pptx(reference, tmp_path / "cand.pptx", "rotation_changed")
    before = candidate.read_bytes()
    report = compare_dna(
        extract_fidelity_dna(reference, slide_index=3),
        extract_fidelity_dna(candidate, slide_index=3),
    )
    execute_repair_plan(candidate, repair_plan_dict(report), tmp_path / "out.pptx", slide_index=3)
    assert candidate.read_bytes() == before


def test_repaired_package_is_a_valid_pptx(tmp_path: Path):
    outcome = _repair_roundtrip(tmp_path, "geometry_changed")
    output = Path(outcome["result"]["output"])
    with zipfile.ZipFile(output) as zf:
        assert zf.testzip() is None
    prs = Presentation(str(output))
    assert len(prs.slides) == 4
