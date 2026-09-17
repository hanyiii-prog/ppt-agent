# -*- coding: utf-8 -*-
"""Tests for ppt_agent.page_kits (Kimi v2-derived page-composition kits)."""
import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches  # noqa: E402

from ppt_agent import page_kits as K  # noqa: E402
from ppt_agent.clone_shell import audit_pages  # noqa: E402


@pytest.fixture()
def content_slide(tmp_path):
    """One slide on a '内容页 - 有标题' layout (matches copy_logos' default)."""
    prs = Presentation()
    lay = prs.slide_layouts[5]            # Title Only
    lay.name = "内容页 - 有标题"
    prs.slides.add_slide(lay)
    p = tmp_path / "t.pptx"
    prs.save(str(p))
    prs = Presentation(str(p))
    return prs, prs.slides[0]


def _prst(sh):
    g = sh._element.spPr.find(qn("a:prstGeom"))
    return g.get("prst") if g is not None else None


def _texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def test_content_header_draws_chrome_and_title(content_slide):
    prs, sl = content_slide
    top = K.content_header(sl, "二、建设计划 | 项目组织架构", "副标题一行")
    assert abs(top - K.TOP) < 1e-6
    ts = _texts(sl)
    assert any("项目组织架构" in t for t in ts)
    assert any("副标题一行" in t for t in ts)
    # chrome bar present (full-width top rect)
    bars = [s for s in sl.shapes if (s.width or 0) > 12 * 914400 and abs(s.top or 0) < 1000]
    assert bars, "content bar missing"


def test_four_role_cards_builds_cards_and_note(content_slide):
    prs, sl = content_slide
    cards = [{"name": f"主体{i}", "sub": "子标题", "role": "职责", "duties": ["a", "b"]}
             for i in range(4)]
    K.four_role_cards(sl, cards, "协同机制说明：xxx", title="多方协同", lead="四方协同")
    ts = _texts(sl)
    for i in range(4):
        assert any(f"主体{i}" in t for t in ts)
    assert any("协同机制说明" in t for t in ts)
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def test_org_chart_builds_nodes_and_depts(content_slide):
    prs, sl = content_slide
    groups = {"left": {"label": "院方工作组", "chip_head": "信息科", "chip_body": "统筹"},
              "right": {"label": "卫宁健康项目组",
                        "nodes": [{"head": "项目负责人：韩熠", "body": "统筹", "accent": True},
                                  {"head": "项目经理：贾俊俊", "body": "推进"},
                                  {"head": "研发负责人", "body": "研发"}]}}
    depts = [{"name": "种植科", "sub": "无牙颌", "lead": "王艳颖", "members": ["王庆福"]},
             {"name": "颌面外科", "sub": "恶性肿瘤", "lead": "严颖彬", "members": ["张岩"]},
             {"name": "正畸科", "sub": "骨性III类", "lead": "张淋坤", "members": ["林晨"]}]
    K.org_chart(sl, "院领导：马文盛", groups, depts, "病种牵头人职责：xxx",
                title="项目组织架构", lead="院领导挂帅")
    ts = _texts(sl)
    assert any("马文盛" in t for t in ts)
    assert any("韩熠" in t for t in ts)
    assert any("种植科" in t for t in ts)
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def test_two_panel_list_builds_two_panels(content_slide):
    prs, sl = content_slide
    left = {"title": "数据安全与合规",
            "rows": [("封闭部署", "院内局域网部署"), ("导出脱敏", "内置脱敏"),
                     ("留痕审计", "可追溯")]}
    right = {"title": "接口两步走",
             "rows": [("第一步", "打通基础通道"), ("第二步", "补影像接口")]}
    K.two_panel_list(sl, left, right, title="数据安全", lead="封闭合规")
    ts = _texts(sl)
    assert any("数据安全与合规" in t for t in ts)
    assert any("接口两步走" in t for t in ts)
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def test_progress_timeline_status_chips(content_slide):
    prs, sl = content_slide
    steps = [{"date": "8.31~9.4", "status": "done", "title": "双轨选型", "desc": "x"},
             {"date": "9.11~9.25", "status": "doing", "title": "收集标签", "desc": "y",
              "current": True},
             {"date": "9.30", "status": "plan", "title": "转开发", "desc": "z"}]
    K.progress_timeline(sl, steps, "材料收集与下发情况：xxx", title="推进纪实", lead="前期准备")
    ts = _texts(sl)
    assert any("双轨选型" in t for t in ts)
    assert any("当前阶段" in t for t in ts)
    # status chips carry gradient fills (done/doing/plan ramps)
    grads = [s for s in sl.shapes if s._element.spPr.find(qn("a:gradFill")) is not None]
    assert len(grads) >= 3
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def _rects(slide):
    """(L,T,W,H) in inches for every autoshape, rotation-aware."""
    from ppt_agent.clone_shell import rotated_bbox
    out = []
    for s in slide.shapes:
        if s.shape_type is not None and s.has_text_frame and s._element.tag.endswith("}sp"):
            out.append(rotated_bbox(s))
    return out


def test_quad_cards_matches_reference_grid(content_slide):
    """2x2 grid must land on the reference coordinates: cards 5.972x2.639 at
    (0.556,1.389) / (6.806,1.389) / (0.556,4.250) / (6.806,4.250)."""
    prs, sl = content_slide
    cards = [{"title": f"要点{i}",
              "body": [("不追求一步到位", "b"), ("以点带面跑通试点", "i")]}
             for i in range(4)]
    K.quad_cards(sl, cards, title="建设背景", lead="立足务实")
    rects = _rects(sl)
    hits = []
    for want in ((0.556, 1.389), (6.806, 1.389), (0.556, 4.250), (6.806, 4.250)):
        hit = [r for r in rects
               if abs(r[0] - want[0]) < 0.02 and abs(r[1] - want[1]) < 0.02
               and abs(r[2] - 5.972) < 0.02 and abs(r[3] - 2.639) < 0.02]
        hits.append(bool(hit))
    assert all(hits), "quad grid off reference: %s" % hits
    ts = _texts(sl)
    assert any("要点3" in t for t in ts)
    # emphasis run model: bold + inline-blue runs both survive
    styled = {r.text: (r.font.bold, r.font.color.rgb)
              for s in sl.shapes if s.has_text_frame
              for p in s.text_frame.paragraphs for r in p.runs}
    assert styled.get("以点带面跑通试点", (None, )) [0] is True
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def test_quad_cards_note_compresses_the_grid(content_slide):
    """With a note bar the card rows compress: the second row must clear the
    bar, and the audit's collision kind stays quiet (the note used to overlap
    the second row's body text by ~0.6in)."""
    prs, sl = content_slide
    cards = [{"title": f"要点{i}", "body": "以点带面跑通试点"} for i in range(4)]
    K.quad_cards(sl, cards, "说明：不追求一步到位。", title="建设背景", lead="立足务实")
    assert not [i for i in audit_pages(prs) if i["kind"] == "collision"], \
        [i for i in audit_pages(prs) if i["kind"] == "collision"]
    rects = _rects(sl)
    row2 = [r for r in rects
            if abs(r[1] - (K.TOP + K.QUAD_PITCH_Y)) < 0.02
            and abs(r[2] - K.QUAD_CARD[0]) < 0.02]
    assert row2, "second row missing"
    assert all(r[1] + r[3] <= K.NOTE_Y for r in row2), row2


def test_column_cards_reproduces_four_column_page(content_slide):
    """四大核心功能: 4 cols of 2.94x4.42 at y=1.33 + a 12.22x1.11 bar at 5.94."""
    prs, sl = content_slide
    cards = [{"title": f"功能{i}", "lines": ["说明文字" * 6], "grad": i == 0}
             for i in range(4)]
    K.column_cards(sl, cards, cols=4, head_h=0.89, top=1.33, card_h=4.42,
                   title="四大核心功能", lead="以智能标签为索引骨架",
                   bottom={"title": "预期效果", "body": "减少漏检", "h": 1.11,
                           "y": 5.94})
    rects = _rects(sl)
    cols = [r for r in rects
            if abs(r[1] - 1.33) < 0.02 and abs(r[2] - 2.94) < 0.03
            and abs(r[3] - 4.42) < 0.03]
    assert len(cols) >= 4, "4 column cards expected, got %d" % len(cols)
    assert abs(cols[0][0] - 0.556) < 0.02 and abs(cols[3][0] - 9.833) < 0.03
    bar = [r for r in rects if abs(r[2] - 12.22) < 0.03 and abs(r[1] - 5.94) < 0.03]
    assert bar, "bottom bar missing"
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def test_column_cards_three_columns_without_head(content_slide):
    """三步任务 / 三项承诺: 3 cols of 3.97 starting at x=0.556, y=1.33."""
    prs, sl = content_slide
    cards = [{"title": f"第{i + 1}步", "lines": ["任务一", "任务二"]}
             for i in range(3)]
    K.column_cards(sl, cards, cols=3, top=1.33, card_h=1.50, head_h=0.0,
                   title="推进安排", lead="分三步推进")
    rects = _rects(sl)
    cols = [r for r in rects
            if abs(r[1] - 1.33) < 0.02 and abs(r[2] - 3.97) < 0.03]
    assert len(cols) >= 3
    assert abs(cols[0][0] - 0.556) < 0.02 and abs(cols[2][0] - 8.806) < 0.03
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def test_stage_cards_homeplate_and_task_list(content_slide):
    """推进安排: 3 homePlate cards at y=1.33 (3.972x1.500) + a 2-column task
    list, each task ONE paragraph of styled runs (no / text / owner)."""
    prs, sl = content_slide
    stages = [{"head": f"第{i + 1}步：阶段", "desc": "说明文字"} for i in range(3)]
    tasks = [{"no": i + 1, "text": f"任务{i + 1}", "owner": "卫宁团队"}
             for i in range(10)]
    K.stage_cards(sl, stages, tasks, task_title="十项重点工作任务",
                  title="推进安排与重点工作任务", lead="分三步推进", top=1.333)
    plates = [s for s in sl.shapes if _prst(s) == "homePlate"]
    assert len(plates) == 3, "stage arrows must use homePlate geometry"
    from ppt_agent.clone_shell import rotated_bbox
    boxes = [rotated_bbox(s) for s in plates]
    for b, want_x in zip(sorted(boxes), (0.556, 4.681, 8.806)):
        assert abs(b[1] - 1.333) < 0.02 and abs(b[2] - 3.972) < 0.03
        assert abs(b[0] - want_x) < 0.03, (b, want_x)
    # one paragraph per task, not three
    tfs = [s.text_frame for s in sl.shapes
           if s.has_text_frame and "任务1（" in s.text_frame.text]
    assert tfs, "task list missing"
    paras = [p for p in tfs[0].paragraphs if p.runs]
    assert len(paras) == 5, "expected 5 tasks per column, got %d" % len(paras)
    assert len(paras[0].runs) == 3, "task row must be number+text+owner runs"
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def test_stage_timeline_homeplate_stages(content_slide):
    prs, sl = content_slide
    stages = [{"head": "10 月 · 出程序", "tasks": "上旬…\n中旬…", "milestone": "首批宽表", "strong": True},
              {"head": "11 月 · 磨合", "tasks": "上旬…", "milestone": "待验收版"},
              {"head": "12 月初 · 交付", "tasks": "上旬…", "milestone": "整体交付", "strong": True}]
    K.stage_timeline(sl, stages, "当前：标签收集（9 月，进行中）", "责任分工：xxx",
                     title="总体时间节点", lead="三阶段")
    plates = [s for s in sl.shapes if _prst(s) == "homePlate"]
    assert len(plates) >= 3, "stage arrows must use homePlate geometry"
    ts = _texts(sl)
    assert any("当前：标签收集" in t for t in ts)
    assert not [i for i in audit_pages(prs) if i["kind"] == "overflow"]


def test_toc_page_matches_reference_layout(content_slide):
    """目录页 -- the reference deck's own grid, and zero leftover placeholders.

    The defect this locks down: the page is built on the *content* shell, whose
    title placeholder is empty once cleared. An empty placeholder resolves back
    to its layout twin, so renderers painted the layout's skeleton heading
    ``单击此处编辑母版标题样式`` across the top of the contents page. ``toc_page``
    therefore drops it instead of blanking it.
    """
    prs, sl = content_slide
    K.toc_page(
        sl,
        [("建设背景与初步方案", "顶层方针 · 现状矛盾 · 三大试点科室诉求"),
         ("建设计划", "五层技术架构 · 四大核心功能 · 两套数据模型"),
         ("近期推进工作", "专题会议 · 临床调研 · 方案编制 · 模型设计"),
         ("下一步计划", "总体时间节点 · 推进安排 · 重点工作任务")],
        note=["立足务实原则，以点带面跑通试点；", "沉淀两套核心数据模型。"],
        prs=prs,
    )
    # the shell's empty title slot must NOT survive
    assert not any(s.is_placeholder for s in sl.shapes)
    ovals = [s for s in sl.shapes if _prst(s) == "ellipse"]
    deco = [s for s in ovals
            if abs((s.left or 0) / 914400 - K.TOC_DECO[0]) < 1e-3
            and abs((s.top or 0) / 914400 - K.TOC_DECO[1]) < 1e-3]
    assert deco, "deco circle missing"
    assert (deco[0].width or 0) / 914400 == pytest.approx(4.4444, abs=1e-3)
    stops = deco[0]._element.spPr.find(qn("a:gradFill")).findall(
        ".//" + qn("a:gs"))
    alphas = [int(g.find(qn("a:srgbClr")).find(qn("a:alpha")).get("val"))
              for g in stops]
    assert alphas == [7843, 1961], "deco must stay a near-invisible wash"
    for k in range(4):
        y = K.TOC_ITEM_Y0 + k * K.TOC_ITEM_PITCH
        hit = [s for s in ovals
               if abs((s.left or 0) / 914400 - K.TOC_ITEM_X) < 1e-3
               and abs((s.top or 0) / 914400 - y) < 1e-3
               and (s.width or 0) / 914400 == pytest.approx(0.75, abs=1e-3)]
        assert hit, "ordinal circle %d missing at y=%.3f" % (k + 1, y)
    ts = _texts(sl)
    assert "目录" in ts and "CONTENTS" in ts
    for num in ("01", "02", "03", "04"):
        assert num in ts, num
    assert any("建设计划" in t for t in ts)
    assert not [i for i in audit_pages(prs) if i["kind"] == "stale_placeholder"]
