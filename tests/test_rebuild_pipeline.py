import json

import pytest
from pptx import Presentation

from ppt_agent.rebuild_pipeline import (
    BUILTIN_TEMPLATES,
    DeckBrief,
    build_source_pack,
    build_v3_deck,
    evaluate_deck,
    select_template,
)


MARKDOWN = """# Quarterly Project Review

## Background

The project started in January and is tracked by the steering group.

## Results

| Item | Status |
| --- | --- |
| Data collection | Complete |
| Pilot review | Underway |

## Next steps

- Finalize integration testing
- Prepare operations handbook
"""


def test_source_pack_preserves_hashes_and_blocks(tmp_path):
    source = tmp_path / "report.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    document, pack = build_source_pack(source)
    assert pack.title == "Quarterly Project Review"
    assert pack.sources[0].format == "markdown"
    assert len(pack.blocks) == len(document.blocks)
    assert len(pack.sources[0].content_hash) == 64


def test_template_selection_uses_objective_and_tone():
    brief = DeckBrief.from_task({"objective": "compare platform data", "tone": "analytical"})
    assert select_template(brief)["id"] == "data-analysis"

    brief = DeckBrief.from_task({"objective": "project progress report", "tone": "formal"})
    assert select_template(brief)["id"] in {"project-progress", "executive-briefing"}


def test_build_v3_deck_writes_contract_artifacts_and_pptx(tmp_path):
    source = tmp_path / "report.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    output = tmp_path / "workspace"
    report = build_v3_deck(
        source,
        output,
        task={
            "objective": "project progress report for stakeholders",
            "audience": "steering group",
            "tone": ["formal", "credible"],
            "max_pages": 8,
        },
    )
    assert report["status"] == "passed"
    assert report["template"]["source"] == "builtin"
    assert report["quality"]["content_quality"]["source_coverage"] == 1.0

    for name in (
        "source-pack.json", "evidence-graph.json", "deck-brief.json",
        "outline.json", "template-profile.json", "page-design-ir.json",
        "final-report.json", "deck.pptx", "deck.html",
    ):
        assert (output / name).is_file()

    deck = Presentation(str(output / "deck.pptx"))
    assert len(deck.slides) == report["page_total"]
    assert report["quality"]["visual_quality"]["model_review"] == "not_run"


def test_quality_report_fails_on_missing_sources(tmp_path):
    source = tmp_path / "report.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    document, pack = build_source_pack(source)
    evidence = {"facts": [{"id": "fact-0001", "source_block_ids": [pack.blocks[0]["id"]]}]}
    outline = {
        "slides": [{"id": "slide-0001", "kind": "content", "title": "Title", "takeaway": "Takeaway"}],
        "omissions": [],
    }
    from ppt_agent.ir import Presentation

    presentation = Presentation(version="0.1", title="x", slides=[])
    report = evaluate_deck(presentation, outline, pack, evidence, {}, "rules")
    assert report["status"] == "failed"
    assert report["content_quality"]["missing_source_blocks"]


def test_build_v3_deck_rejects_impossible_page_budget(tmp_path):
    source = tmp_path / "report.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    with pytest.raises(ValueError, match="page_budget"):
        build_v3_deck(source, tmp_path / "out", task={"page_budget": {"min": 5, "max": 2}})


def test_model_planning_is_validated_and_labeled(tmp_path):
    source = tmp_path / "report.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    from ppt_agent.parsers import parse_markdown

    document = parse_markdown(MARKDOWN)
    ids = [block.id for block in document.blocks]
    response = json.dumps({
        "schema": "storyboard/v1",
        "design_profile": {"palette": {}, "typography": {}},
        "slides": [
            {
                "kind": "cover", "title": "Quarterly Project Review",
                "key_message": "", "source_block_ids": [ids[0]],
                "visual_form": "cover", "evidence_ids": [ids[0]],
            },
            {
                "kind": "content", "title": "Status and next steps",
                "key_message": "The pilot is underway and two actions remain.",
                "source_block_ids": ids[1:],
                "visual_form": "progress_sequence", "evidence_ids": ids[1:],
            },
        ],
        "omitted_source_blocks": [],
    }, ensure_ascii=False)

    output = tmp_path / "workspace"
    report = build_v3_deck(
        source, output,
        llm_fn=lambda _prompt: response,
        task={"objective": "project progress", "max_pages": 4},
    )
    assert report["planning"]["mode"] == "model"
    assert report["status"] == "passed"
    assert report["page_total"] == 3
    assert json.loads((output / "outline.json").read_text(encoding="utf-8"))["planning_mode"] == "model"
