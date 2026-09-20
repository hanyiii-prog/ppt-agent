"""End-to-end fidelity closed loop (md §18/§19) on real PPTX packages.

reference -> extract -> candidate (real mutations) -> diff -> repair ->
re-extract -> structural gate -> render both -> visual diff -> final gate,
plus the master/layout attribution regressions.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu

from ppt_agent.fidelity import extract_fidelity_dna
from ppt_agent.fidelity_diff import compare_dna
from ppt_agent.fidelity_gate import compare_decks
from ppt_agent.fidelity_pipeline import repair_deck
from ppt_agent.visual_regression import visual_status

from tests.fidelity_fixtures import build_rich_pptx, mutate_pptx


# --------------------------------------------------------------------------- #
# master / layout attribution regressions
# --------------------------------------------------------------------------- #
def _inheritance_deck(path: Path) -> Path:
    """A deck whose title style resolves through the master text styles."""
    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text_frame.paragraphs[0].add_run().text = "继承链标题"
    prs.save(str(path))
    return path


def test_master_typography_change_is_attributed_to_master(tmp_path: Path):
    reference = _inheritance_deck(tmp_path / "ref.pptx")

    prs = Presentation(str(reference))
    master = prs.slide_masters[0]
    title_style = master.element.find(qn("p:txStyles")).find(qn("p:titleStyle"))
    def_rpr = title_style.find(qn("a:lvl1pPr")).find(qn("a:defRPr"))
    def_rpr.set("sz", "3200")
    candidate = tmp_path / "master_changed.pptx"
    prs.save(str(candidate))

    ref_dna = extract_fidelity_dna(reference, slide_index=1)
    cand_dna = extract_fidelity_dna(candidate, slide_index=1)

    ref_title = next(r for r in ref_dna["placeholder_resolutions"]
                     if r["placeholder"].get("type") in ("title", "ctrTitle"))
    cand_title = next(r for r in cand_dna["placeholder_resolutions"]
                      if r["placeholder"].get("type") in ("title", "ctrTitle"))
    assert ref_title["source"]["font_size_pt"] == "master"
    assert cand_title["source"]["font_size_pt"] == "master"
    assert cand_title["resolved"]["font_size_pt"] == 32.0
    assert ref_title["resolved"]["font_size_pt"] != cand_title["resolved"]["font_size_pt"]

    report = compare_dna(ref_dna, cand_dna)
    assert not report.passed
    assert report.codes.get("text.font_size")


def test_layout_chrome_change_reports_style_fill(tmp_path: Path):
    reference = build_rich_pptx(tmp_path / "ref.pptx")
    prs = Presentation(str(reference))
    band_layout = prs.slide_layouts[6]
    changed = False
    for sp in band_layout.shapes._spTree.findall(qn("p:sp")):
        sp_pr = sp.find(qn("p:spPr"))
        if sp_pr is None:
            continue
        solid = sp_pr.find(qn("a:solidFill"))
        if solid is not None:
            solid.find(qn("a:srgbClr")).set("val", "FF0000")
            changed = True
            break
    assert changed, "the fixture must carry a filled layout shape to mutate"
    candidate = tmp_path / "layout_changed.pptx"
    prs.save(str(candidate))

    report = compare_dna(
        extract_fidelity_dna(reference, slide_index=3),
        extract_fidelity_dna(candidate, slide_index=3),
    )
    assert not report.passed
    assert report.codes.get("style.fill")
    # the defect is attributed to the inherited layer, not the slide content
    assert any(issue.path.startswith("layout.shapes") for issue in report.issues)


# --------------------------------------------------------------------------- #
# full closed loop
# --------------------------------------------------------------------------- #
def test_full_e2e_closed_loop(tmp_path: Path):
    reference = build_rich_pptx(tmp_path / "reference.pptx")
    # extract DNA of the reference (as the engine would for planning)
    ref_dna = extract_fidelity_dna(reference, slide_index=3)
    assert ref_dna["schema"] == "template-dna/fidelity/v2"

    # candidate: three chained single-property mutations on real PPTX
    candidate = mutate_pptx(reference, tmp_path / "c1.pptx", "rotation_changed")
    candidate = mutate_pptx(candidate, tmp_path / "c2.pptx", "gradient_changed")
    candidate = mutate_pptx(candidate, tmp_path / "c3.pptx", "zorder_changed")

    # structural diff first: it must fail and name the defects
    initial = compare_decks(reference, candidate)
    assert initial["passed"] is False

    # repair loop: diff -> repair -> re-extract -> re-diff -> render -> visual
    result = repair_deck(reference, candidate, tmp_path / "loop", max_iterations=3, render=True)
    assert result["passed"] is True
    assert result["iterations"] >= 1
    assert result["visual"]["status"] == "visual_pass"

    rebuilt = Path(result["candidate_final"])
    # re-extract: the rebuilt deck now matches structurally
    final_gate = compare_decks(reference, rebuilt)
    assert final_gate["passed"]

    # final visual gate on the rebuilt deck
    final_visual = visual_status(reference, rebuilt, tmp_path / "final-visual")
    assert final_visual["status"] == "visual_pass"
    assert final_visual["passed"] is True

    # extract again: rebuilt DNA equals reference DNA
    report = compare_dna(
        extract_fidelity_dna(reference, slide_index=3),
        extract_fidelity_dna(rebuilt, slide_index=3),
    )
    assert report.passed


def test_whole_deck_regression_covers_every_slide(tmp_path: Path):
    reference = build_rich_pptx(tmp_path / "ref.pptx")
    candidate = mutate_pptx(reference, tmp_path / "cand.pptx", "font_changed")
    gate = compare_decks(reference, candidate)
    assert not gate["passed"]
    failed_pages = [page["slide_index"] for page in gate["pages"] if not page["passed"]]
    assert failed_pages == [3], "only the mutated page may fail the deck gate"
