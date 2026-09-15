from __future__ import annotations

import json
from pathlib import Path

from ppt_agent.cli import main


def test_analyze_pptx_writes_nested_slide_count(tmp_path: Path) -> None:
    from pptx import Presentation

    source = tmp_path / "sample.pptx"
    dna = tmp_path / "dna.json"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(source)

    import sys
    old = sys.argv
    try:
        sys.argv = ["ppt-agent", "analyze-pptx", str(source), "-o", str(dna)]
        assert main() == 0
    finally:
        sys.argv = old

    payload = json.loads(dna.read_text(encoding="utf-8"))
    assert payload["presentation"]["slide_count"] == 1
    assert payload["slides"][0]["role"] == "first"
    assert payload["schema"] == "template-dna/v0.3"


def test_dna_to_ir_cli_preserves_master_and_shape_fidelity(tmp_path: Path) -> None:
    source = tmp_path / "dna.json"
    output = tmp_path / "ir.json"
    source.write_text(json.dumps({
        "schema": "template-dna/v0.3",
        "source": "fixture.pptx",
        "presentation": {"slide_size_inches": {"width": 13.333, "height": 7.5}},
        "theme": {"colors": {"dk1": "000000"}},
        "masters": [{"name": "Master 1"}],
        "slides": [{
            "slide": 1,
            "role": "first",
            "layout_name": "Title Slide",
            "shapes": [{
                "id": "7",
                "type": "TEXT_BOX",
                "geometry": {"left": 1.0, "top": 2.0, "width": 5.0, "height": 1.0},
                "style": {"fill": {"transparency": 0.25}},
                "z_index": 3,
                "parent_id": None,
                "fidelity": {"xml_sha256": "abc"},
            }],
        }],
        "special_surfaces": {"first": {}, "last": {}},
    }, ensure_ascii=False), encoding="utf-8")

    import sys
    old = sys.argv
    try:
        sys.argv = ["ppt-agent", "dna-to-ir", str(source), "-o", str(output)]
        assert main() == 0
    finally:
        sys.argv = old

    payload = json.loads(output.read_text(encoding="utf-8"))
    component = payload["slides"][0]["components"][0]
    assert payload["slides"][0]["purpose"] == "cover"
    assert payload["theme"]["masters"] == [{"name": "Master 1"}]
    assert component["style"]["fill"]["transparency"] == 0.25
    assert component["data"]["fidelity"]["z_index"] == 3
