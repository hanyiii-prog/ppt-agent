"""Real-PPTX fixture builders for the fidelity engine regression suite.

Every fixture is a genuine PPTX package produced by python-pptx (plus targeted
raw-XML injections for attributes python-pptx does not expose, e.g. flip,
group child-space scaling, connector arrows). No hand-written DNA dicts.
"""
from __future__ import annotations

import copy
import io
import zipfile
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

EMU_IN = 914400


def _png_bytes(width: int = 64, height: int = 48, color: str = "#2266AA") -> bytes:
    image = Image.new("RGB", (width, height), color)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _inject_flip(shape: Any, *, flip_h: bool = False, flip_v: bool = False) -> None:
    xfrm = shape._element.spPr.find(qn("a:xfrm"))
    if xfrm is None:
        return
    if flip_h:
        xfrm.set("flipH", "1")
    if flip_v:
        xfrm.set("flipV", "1")


def _inject_alpha(shape: Any, alpha_pct: int) -> None:
    """Add <a:alpha> to the shape's solid srgbClr (python-pptx cannot)."""
    solid = shape._element.spPr.find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    alpha = srgb.makeelement(qn("a:alpha"), {"val": str(alpha_pct * 1000)})
    srgb.append(alpha)


def _inject_connector_arrows(connector: Any) -> None:
    ln = connector._element.spPr.find(qn("a:ln"))
    if ln is None:
        return
    head = ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"})
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "arrow", "w": "sm", "len": "lg"})
    ln.append(head)
    ln.append(tail)


def _scale_group_child_space(group_shape: Any, factor: float) -> None:
    """Make the group render its children at ``factor`` scale (ext != chExt)."""
    grp_sp_pr = group_shape._element.find(qn("p:grpSpPr"))
    xfrm = grp_sp_pr.find(qn("a:xfrm"))
    ext = xfrm.find(qn("a:ext"))
    ch_ext = xfrm.find(qn("a:chExt"))
    ext.set("cx", str(int(int(ch_ext.get("cx")) * factor)))
    ext.set("cy", str(int(int(ch_ext.get("cy")) * factor)))


def _insert_layout_shape(slide_layout: Any) -> None:
    """Insert a real decorative band into the layout so inheritance is testable."""
    sp_xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="2001" name="LayoutBand"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="0" y="1496060"/><a:ext cx="10058400" cy="3571880"/></a:xfrm>'
        '<a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom>'
        '<a:solidFill><a:srgbClr val="1185FE"><a:alpha val="15000"/></a:srgbClr></a:solidFill>'
        "</p:spPr>"
        '<p:txBody><a:bodyPr/><a:p/></p:txBody>'
        "</p:sp>"
    )
    from lxml import etree

    sp_tree = slide_layout.shapes._spTree
    sp_tree.append(etree.fromstring(sp_xml))


def build_rich_pptx(path: Path) -> Path:
    """Four slides covering cover/toc/content/closing with every element kind."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # -- slide 1: cover -------------------------------------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x0C, 0x67, 0xBC)
    _inject_alpha(band, 65)
    title = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(1.2))
    run = title.text_frame.paragraphs[0].add_run()
    run.text = "天津市口腔医院专病数据库建设情况汇报"
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.name = "Arial"

    # -- slide 2: toc ---------------------------------------------------------
    toc = prs.slides.add_slide(prs.slide_layouts[6])
    toc_title = toc.shapes.add_textbox(Inches(1), Inches(0.6), Inches(6), Inches(0.9))
    toc_title.text_frame.text = "目录"
    for i, label in enumerate(("项目概况", "建设进展", "下阶段计划"), start=1):
        item = toc.shapes.add_textbox(Inches(1.2), Inches(2 + i * 0.9), Inches(8), Inches(0.6))
        item.text_frame.text = f"{i}. {label}"
        item.rotation = 0

    # -- slide 3: content (every element kind) --------------------------------
    content = prs.slides.add_slide(prs.slide_layouts[6])

    rotated = content.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    rotated.fill.solid()
    rotated.fill.fore_color.rgb = RGBColor(0x11, 0x85, 0xFE)
    rotated.rotation = 90

    flipped = content.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(1), Inches(1.5), Inches(1))
    flipped.fill.solid()
    flipped.fill.fore_color.rgb = RGBColor(0x09, 0x43, 0x7F)
    _inject_flip(flipped, flip_h=True, flip_v=True)

    gradient = content.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(1), Inches(2.5), Inches(1))
    gradient.fill.gradient()
    gradient.fill.gradient_angle = 45.0
    gradient.fill.gradient_stops[0].color.rgb = RGBColor(0x0C, 0x67, 0xBC)
    gradient.fill.gradient_stops[1].color.rgb = RGBColor(0x16, 0x87, 0xF1)

    themed = content.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10), Inches(1), Inches(2), Inches(1))
    themed.fill.solid()
    themed.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1

    group = content.shapes.add_group_shape()
    child_a = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(1), Inches(0.5))
    child_a.fill.solid()
    child_a.fill.fore_color.rgb = RGBColor(0x33, 0x33, 0x33)
    child_b = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(3), Inches(1), Inches(0.5))
    child_b.fill.solid()
    child_b.fill.fore_color.rgb = RGBColor(0x44, 0x44, 0x44)
    _scale_group_child_space(group, 2.0)

    picture = content.shapes.add_picture(io.BytesIO(_png_bytes()), Inches(4), Inches(3), Inches(1.2), Inches(0.9))
    picture.crop_left = 0.25

    connector = content.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6), Inches(3), Inches(8), Inches(4))
    connector.line.color.rgb = RGBColor(0x00, 0x00, 0x00)
    connector.line.width = Pt(1.5)
    _inject_connector_arrows(connector)

    freeform_builder = content.shapes.build_freeform(Emu(9 * EMU_IN), Emu(3 * EMU_IN))
    freeform_builder.add_line_segments(
        [(Emu(10 * EMU_IN), Emu(3 * EMU_IN)), (Emu(10 * EMU_IN), Emu(4 * EMU_IN))], close=True
    )
    freeform_builder.convert_to_shape()

    table_frame = content.shapes.add_table(3, 3, Inches(1), Inches(5), Inches(9), Inches(1.8))
    table = table_frame.table
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "抽取成功率"
    table.cell(1, 1).text = "99.2%"
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(2, 0).text = "评审"

    textbox = content.shapes.add_textbox(Inches(1), Inches(4.4), Inches(5), Inches(0.5))
    tf = textbox.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.08)
    paragraph = tf.paragraphs[0]
    paragraph.add_run().text = "正文示例"

    # -- slide 4: closing ------------------------------------------------------
    closing = prs.slides.add_slide(prs.slide_layouts[6])
    closing_shape = closing.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(9), Inches(2))
    closing_shape.fill.solid()
    closing_shape.fill.fore_color.rgb = RGBColor(0xF4, 0xF9, 0xFF)
    closing_run = closing.shapes.add_textbox(Inches(3), Inches(2.5), Inches(7), Inches(1)).text_frame.paragraphs[0].add_run()
    closing_run.text = "感谢聆听"

    # a real decorative band inherited by every slide from its layout
    _insert_layout_shape(prs.slide_layouts[6])

    prs.save(path)
    return path


def build_toc_pptx(path: Path, *, columns: int = 1, numbered: bool = True, hierarchical: bool = False, indicator: bool = False) -> Path:
    """TOC-only deck: one slide whose headline reads 目录."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    toc = prs.slides.add_slide(prs.slide_layouts[6])
    title = toc.shapes.add_textbox(Inches(1), Inches(0.6), Inches(6), Inches(0.9))
    title.text_frame.text = "目录 CONTENTS"

    entries = ["项目概况", "建设进展", "下阶段计划", "成果与展望"]
    per_column = 2
    for i, label in enumerate(entries):
        column = i // per_column if columns > 1 else 0
        row = i % per_column
        level_indent = 0.6 if (hierarchical and i % 2 == 1) else 0.0
        text = f"{i + 1}. {label}" if numbered else label
        item = toc.shapes.add_textbox(
            Inches(1.2 + column * 5.5 + level_indent), Inches(2 + row * 0.9), Inches(5), Inches(0.6)
        )
        item.text_frame.text = text
    if indicator:
        for i in range(len(entries)):
            dot = toc.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(2.1 + i * 0.9), Inches(0.12), Inches(0.12))
            dot.fill.solid()
            dot.fill.fore_color.rgb = RGBColor(0x11, 0x85, 0xFE)
            dot.line.fill.background()
    prs.save(path)
    return path


def mutate_pptx(source: Path, target: Path, mutation: str) -> Path:
    """Produce a single-property mutation of a real PPTX package."""
    prs = Presentation(str(source))
    content = prs.slides[2]
    shapes = list(content.shapes)

    def shape_by_name(prefix: str) -> Any:
        for shape in shapes:
            if shape.name.startswith(prefix):
                return shape
        raise KeyError(prefix)

    if mutation == "geometry_changed":
        shape_by_name("Rectangle").left = shape_by_name("Rectangle").left + Emu(100000)
    elif mutation == "rotation_changed":
        shape_by_name("Rectangle").rotation = 135.0
    elif mutation == "flip_changed":
        _inject_flip(shape_by_name("Rectangle"), flip_v=True)
    elif mutation == "alpha_changed":
        target_shape = None
        for shape in shapes:
            if shape.has_text_frame and shape.text_frame.text.startswith("正文"):
                target_shape = shape
                break
        # add solid fill with different alpha to a rectangle that has none via XML
        rect = shape_by_name("Oval") if target_shape is None else target_shape
        from lxml import etree

        sp_pr = rect._element.spPr
        solid = sp_pr.makeelement(qn("a:solidFill"), {})
        color = solid.makeelement(qn("a:srgbClr"), {"val": "1185FE"})
        alpha = color.makeelement(qn("a:alpha"), {"val": "40000"})
        color.append(alpha)
        solid.append(color)
        sp_pr.append(solid)
    elif mutation == "zorder_changed":
        sp_tree = shapes[0]._element.getparent()
        first = shapes[0]._element
        sp_tree.remove(first)
        sp_tree.append(first)
    elif mutation == "font_changed":
        for shape in shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size is not None:
                            run.font.size = Pt(run.font.size.pt + 2)
                        run.font.bold = not bool(run.font.bold)
    elif mutation == "gradient_changed":
        gradient = shape_by_name("Rounded Rectangle")
        gradient.fill.gradient_angle = 90.0
    elif mutation == "crop_changed":
        picture = None
        for shape in shapes:
            if shape.shape_type == 13:  # PICTURE
                picture = shape
                break
        picture.crop_right = 0.1
    elif mutation == "media_changed":
        picture = None
        for shape in shapes:
            if shape.shape_type == 13:
                picture = shape
                break
        import io as _io

        new_part, new_rid = content.part.get_or_add_image_part(_io.BytesIO(_png_bytes(color="#AA2266")))
        picture._element.blipFill.blip.set(qn("r:embed"), new_rid)
    elif mutation == "table_changed":
        for shape in shapes:
            if getattr(shape, "has_table", False):
                table = shape.table
                table.columns[0].width = Emu(int(table.columns[0].width) + 100000)
                table.rows[0].height = Emu(int(table.rows[0].height) + 50000)
                break
    else:
        raise ValueError(f"unknown mutation: {mutation}")

    prs.save(str(target))
    return target


def raw_slide_xml(path: Path, slide_index: int) -> str:
    with zipfile.ZipFile(path) as zf:
        name = f"ppt/slides/slide{slide_index}.xml"
        return zf.read(name).decode("utf-8")
