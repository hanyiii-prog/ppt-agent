from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from ppt_agent.page_validation import (
    infer_page_role,
    validate_rendered_pages,
    validate_structural_pages,
)
from PIL import Image


def test_page_roles_first_last_and_chapter():
    assert infer_page_role(1, 4, "封面") == "first"
    assert infer_page_role(2, 4, "目录") == "chapter"
    assert infer_page_role(3, 4, "正文") == "body"
    assert infer_page_role(4, 4, "结论") == "last"


def test_gate_records_text_and_image_shape_counts(tmp_path: Path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides[0]
    box = slide.shapes.add_textbox(0, 0, 1000000, 500000)
    box.text = "标题"
    pptx_path = tmp_path / "deck.pptx"
    prs.save(pptx_path)

    image = tmp_path / "slide-1.png"
    Image.new("RGB", (320, 180), "white").save(image)
    report = validate_rendered_pages(pptx_path, [image])
    assert report.pages[0].role == "first"
    assert report.pages[0].text_shapes == 1
    assert report.pages[0].image_shapes == 0


def test_gate_blocks_dark_text_on_dark_slide_background(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(10, 30, 40)
    box = slide.shapes.add_textbox(500000, 500000, 4000000, 800000)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Unreadable title"
    run.font.color.rgb = RGBColor(0, 0, 0)
    pptx_path = tmp_path / "dark-text.pptx"
    prs.save(pptx_path)

    image = tmp_path / "slide-1.png"
    Image.new("RGB", (320, 180), (10, 30, 40)).save(image)
    report = validate_rendered_pages(pptx_path, [image])

    assert not report.passed
    assert any("contrast" in issue.lower() for issue in report.pages[0].issues)


def test_gate_does_not_guess_unresolved_text_color_on_dark_background(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(10, 30, 40)
    box = slide.shapes.add_textbox(500000, 500000, 4000000, 800000)
    box.text_frame.text = "Color inherited from the theme"
    pptx_path = tmp_path / "inherited-color.pptx"
    prs.save(pptx_path)

    report = validate_structural_pages(pptx_path)

    assert report.passed
    assert not any("contrast" in issue.lower() for issue in report.pages[0].issues)


def test_gate_allows_inherited_decorative_bleed(tmp_path: Path):
    image = tmp_path / "background.png"
    Image.new("RGB", (100, 100), "blue").save(image)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(image),
        Inches(0),
        Inches(-0.07),
        width=prs.slide_width,
        height=Inches(4.72),
    )
    pptx_path = tmp_path / "decorative-bleed.pptx"
    prs.save(pptx_path)

    report = validate_structural_pages(pptx_path)

    assert report.passed
    assert not any("outside slide bounds" in issue for issue in report.pages[0].issues)
