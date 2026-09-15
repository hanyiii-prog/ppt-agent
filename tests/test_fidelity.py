from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from ppt_agent.fidelity import extract_fidelity_dna


def test_fidelity_dna_captures_stack_style_and_assets(tmp_path: Path):
    pptx_path = tmp_path / "fixture.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bottom = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(5), Inches(2))
    bottom.fill.solid()
    bottom.fill.fore_color.rgb = RGBColor(0x10, 0x20, 0x30)
    bottom.fill.transparency = 35

    top = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(3), Inches(1))
    top.fill.solid()
    top.fill.fore_color.rgb = RGBColor(0x30, 0x60, 0xA0)
    top.line.width = Inches(0.02)

    prs.save(pptx_path)
    dna = extract_fidelity_dna(pptx_path)

    assert dna["surface_role"] == "first"
    assert dna["presentation"]["slide_count"] == 1
    assert dna["layout"]["raw_xml"]
    assert dna["master"]["raw_xml"]
    assert len(dna["slide"]["shapes"]) >= 2

    shapes = dna["slide"]["shapes"]
    assert [shape["z_index"] for shape in shapes] == sorted(shape["z_index"] for shape in shapes)
    assert all(shape["raw_xml"] for shape in shapes)
    assert any("spPr_xml" in shape for shape in shapes)
