"""Parser tests: markdown / docx / pptx -> Content IR."""

from __future__ import annotations

from pathlib import Path

import pytest

from ppt_agent.content_ir import ContentDocument
from ppt_agent.parsers import parse_auto, parse_markdown

SAMPLE = """# 口腔医院互联互通项目上线总结

## 项目概况

- 覆盖 5 个院区
- 互联互通四甲评审通过

## 关键指标

| 指标 | 数值 |
| --- | --- |
| 数据抽取成功率 | 99.2% |
| 上线按期率 | 100% |

> 全院一盘棋，上线零事故。

1. 第一步完成数据治理
2. 第二步完成接口联调
"""


def test_markdown_full_vocabulary() -> None:
    document = parse_markdown(SAMPLE, source="sample.md")
    assert document.title == "口腔医院互联互通项目上线总结"
    kinds = [block.type for block in document.blocks]
    assert kinds[0] == "heading"
    assert "bullets" in kinds and "table" in kinds and "quote" in kinds and "ordered" in kinds
    table = next(block for block in document.blocks if block.type == "table")
    assert table.headers == ["指标", "数值"]
    assert table.rows == [["数据抽取成功率", "99.2%"], ["上线按期率", "100%"]]
    ordered = next(block for block in document.blocks if block.type == "ordered")
    assert ordered.items == ["第一步完成数据治理", "第二步完成接口联调"]
    assert all(block.id for block in document.blocks), "every block needs a stable id"


def test_auto_dispatch_markdown(tmp_path: Path) -> None:
    path = tmp_path / "deck-source.md"
    path.write_text(SAMPLE, encoding="utf-8")
    document = parse_auto(path)
    assert isinstance(document, ContentDocument)
    assert document.format == "markdown"
    assert document.source == str(path)


def test_auto_dispatch_rejects_unknown(tmp_path: Path) -> None:
    path = tmp_path / "mystery.xyz"
    path.write_text("whatever", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_auto(path)


def test_docx_parser(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    import docx

    path = tmp_path / "source.docx"
    document = docx.Document()
    document.add_heading("口腔医院上线总结", level=1)
    document.add_heading("项目概况", level=2)
    document.add_paragraph("覆盖 5 个院区，服务 92 人团队")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "数值"
    table.rows[1].cells[0].text = "抽取成功率"
    table.rows[1].cells[1].text = "99.2%"
    document.save(str(path))

    parsed = parse_auto(path)
    assert parsed.format == "docx"
    assert parsed.title == "口腔医院上线总结"
    headings = [block for block in parsed.blocks if block.type == "heading"]
    assert [block.level for block in headings] == [1, 2]
    tables = [block for block in parsed.blocks if block.type == "table"]
    assert tables and tables[0].rows == [["抽取成功率", "99.2%"]]


def test_pptx_text_parser(tmp_path: Path) -> None:
    pytest.importorskip("pptx")
    from pptx import Presentation

    path = tmp_path / "source.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "项目概况"
    body = slide.placeholders[1]
    body.text = "覆盖 5 个院区"
    body.text_frame.add_paragraph().text = "四甲评审通过"
    prs.save(str(path))

    parsed = parse_auto(path)
    assert parsed.format == "pptx"
    assert parsed.title == "项目概况"
    headings = [block for block in parsed.blocks if block.type == "heading"]
    bullets = [block for block in parsed.blocks if block.type == "bullets"]
    assert headings and headings[0].text == "项目概况"
    assert bullets and "覆盖 5 个院区" in bullets[0].items
