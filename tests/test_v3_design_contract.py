import json
import sys

import pytest

from ppt_agent.deck_brief import DeckBrief
from ppt_agent.evidence_graph import build_evidence_graph
from ppt_agent.outline_validator import validate_outline
from ppt_agent.source_pack import build_source_pack
from ppt_agent.outline_planner import Outline
from ppt_agent.composition_compiler import compile_slide
from ppt_agent.slide_author import author_slide_specs
from ppt_agent.repair_loop import apply_repair_requests
from ppt_agent.sdk import PptAgent


def test_multi_source_pack_keeps_every_block_and_traceable_facts(tmp_path):
    primary = tmp_path / "report.md"
    primary.write_text("# Report\n\nPrimary fact.\n", encoding="utf-8")
    supporting = tmp_path / "notes.html"
    supporting.write_text(
        "<html><body><h1>Notes</h1><p>HTML fact.</p></body></html>",
        encoding="utf-8",
    )
    image = tmp_path / "evidence.png"
    image.write_bytes(b"not a real image")

    document, pack = build_source_pack([primary, supporting, image])

    assert [item["format"] for item in pack.to_dict()["sources"]] == [
        "markdown", "html", "png",
    ]
    assert len(pack.blocks) == 5
    assert all(block["source_id"] in {"source-001", "source-002", "source-003"} for block in pack.blocks)

    evidence = build_evidence_graph(pack).to_dict()
    assert {fact["source_block_ids"][0] for fact in evidence["facts"]} == {
        block["id"] for block in pack.blocks
    }
    assert len({relation["id"] for relation in evidence["relations"]}) == len(evidence["relations"])
    assert all(
        set(relation["source_block_ids"]) <= {block["id"] for block in pack.blocks}
        for relation in evidence["relations"]
    )


def test_source_priority_can_promote_a_supporting_source(tmp_path):
    first = tmp_path / "first.md"
    second = tmp_path / "second.md"
    first.write_text("# First\n\nFirst.\n", encoding="utf-8")
    second.write_text("# Second\n\nSecond.\n", encoding="utf-8")

    _, pack = build_source_pack([first, second], source_priority=["second.md"])
    priorities = {item["id"]: item["priority"] for item in pack.to_dict()["sources"]}
    assert priorities["source-002"] == "primary"
    assert priorities["source-001"] == "supporting"


def test_outline_rejects_unknown_or_untraceable_evidence(tmp_path):
    source = tmp_path / "report.md"
    source.write_text("# Report\n\nFact.\n", encoding="utf-8")
    _, pack = build_source_pack(source)
    brief = DeckBrief.from_task({})
    outline = {
        "slides": [{
            "id": "slide-0001", "kind": "content", "title": "Title",
            "purpose": "purpose", "takeaway": "takeaway",
            "source_block_ids": [], "evidence_ids": ["not-a-fact"],
        }],
    }

    errors = validate_outline(outline, pack, brief, evidence={"facts": []})
    assert any("evidence" in error for error in errors)


def test_model_slide_authoring_is_validated_and_falls_back_per_slide(tmp_path):
    source = tmp_path / "report.md"
    source.write_text(
        "# Report\n\n## Results\n\n- Data collection is complete\n- 12 sites were tested\n",
        encoding="utf-8",
    )
    _, pack = build_source_pack(source)
    evidence = build_evidence_graph(pack)
    outline = Outline(
        schema="outline/v1", title="Report", story={},
        slides=[
            {"id": "slide-0001", "kind": "cover", "title": "Report", "purpose": "open"},
            {
                "id": "slide-0002", "kind": "content", "title": "Results",
                "purpose": "show results", "takeaway": "Testing is complete",
                "source_block_ids": [block["id"] for block in pack.blocks[1:]],
                "evidence_ids": [],
            },
        ],
    )

    def llm(prompt: str) -> str:
        if '"slide": 1' in prompt:
            return json.dumps({"id": "slide-0001", "speaker_notes": "Welcome."})
        return json.dumps({
            "id": "slide-0002",
            "items": ["  Collection is complete  "],
            "metrics": [{"label": "sites", "value": "12"}],
            "visual_intent": {"preferred_forms": ["metric_group"], "density": "standard"},
            "speaker_notes": "Emphasise the completed test coverage.",
        })

    specs = author_slide_specs(outline, pack, evidence, llm_fn=llm)
    assert specs[0].speaker_notes == "Welcome."
    assert specs[1].content["items"] == [{"text": "Collection is complete"}]
    assert specs[1].content["metrics"] == [{"label": "sites", "value": "12"}]
    assert specs[1].visual_intent["preferred_forms"] == ["metric_group"]
    assert specs[1].speaker_notes == "Emphasise the completed test coverage."
    assert all(item in {fact["id"] for fact in evidence.facts} for item in specs[1].evidence_ids)

    def broken_llm(prompt: str) -> str:
        if '"slide": 2' in prompt:
            return "not json"
        return "{}"

    fallback_specs = author_slide_specs(outline, pack, evidence, llm_fn=broken_llm)
    assert fallback_specs[1].content["authoring_warning"]
    assert fallback_specs[1].content["items"]


def test_user_template_uses_clone_shell_route(tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation

    from ppt_agent.pipeline import run_v3_pipeline

    template = tmp_path / "template.pptx"
    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    body_layout = prs.slide_layouts[5]
    title_layout.name = "标题幻灯片"
    body_layout.name = "内容页 - 有标题"
    prs.slides.add_slide(title_layout)
    for _ in range(3):
        prs.slides.add_slide(body_layout)
    prs.slides.add_slide(title_layout)
    prs.save(template)

    source = tmp_path / "report.md"
    source.write_text("# User Template Report\n\n- User fact one\n- User fact two\n", encoding="utf-8")
    output = tmp_path / "workspace"
    report = run_v3_pipeline(
        source,
        output,
        template_path=template,
        task={"objective": "present user supplied facts", "max_pages": 4},
    )
    assert report["template"]["source"] == "user"
    assert report["template"]["route"] == "user-template-clone-shell"
    assert len(Presentation(str(output / "deck.pptx")).slides) == report["page_total"]
    assert (output / "quality" / "user-template" / "delivery-report.json").is_file()


def test_v3_stage_surface_is_shared_by_sdk_cli_and_mcp(tmp_path):
    from ppt_agent.cli import main
    from ppt_agent.mcp.tools import list_tools

    source = tmp_path / "report.md"
    source.write_text("# Stage Report\n\n- One traceable fact\n- Another fact\n", encoding="utf-8")
    agent = PptAgent()
    analysis = agent.analyze_sources(source, task={"objective": "stage surface"})
    assert analysis["schema"] == "source-analysis/v1"
    assert analysis["content_relations"]["schema"] == "content-relations/v1"

    outline_stage = agent.plan_v3_outline(source, task={"objective": "stage surface", "max_pages": 4})
    assert outline_stage["passed"]
    assert outline_stage["outline"]["planning_mode"] == "rules"

    design_stage = agent.compile_v3_page_design(source, task={"objective": "stage surface", "max_pages": 4})
    assert design_stage["page_design_ir"]["schema"] == "page-design-ir/v1"
    assert agent.list_builtin_templates()["templates"]

    tool_names = {item["name"] for item in list_tools()}
    assert {
        "ppt_agent_analyze_sources", "ppt_agent_plan_outline", "ppt_agent_validate_outline",
        "ppt_agent_author_slides", "ppt_agent_list_builtin_templates", "ppt_agent_select_template",
        "ppt_agent_compile_page_design", "ppt_agent_evaluate_deck",
    } <= tool_names

    output = tmp_path / "stage.json"
    argv = ["ppt-agent", "analyze-sources", str(source), "-o", str(output)]
    original_argv = sys.argv
    sys.argv = argv
    try:
        assert main() == 0
    finally:
        sys.argv = original_argv
    assert json.loads(output.read_text(encoding="utf-8"))["schema"] == "source-analysis/v1"


def test_overflow_repair_changes_page_design_ir_and_rebuild_geometry():
    design_ir = {
        "schema": "page-design-ir/v1", "title": "Repaired",
        "theme": {
            "colors": {"text": "1F2937", "accent": "C6602D"},
            "fonts": {"body": "Test Sans"},
            "type_scale": {"body": 18},
        },
        "pages": [{
            "page_id": "slide-0001", "kind": "content", "title": "Title",
            "composition": "text_evidence", "visual_form": "text_evidence",
            "speaker_notes": "", "blocks": [],
            "slots": {
                "body": {
                    "id": "body", "kind": "body", "font": "body",
                    "x": 0.8, "y": 2.4, "w": 11.7, "h": 4.2,
                    "value": "Body text", "color": "text",
                },
            },
        }],
    }

    repaired = apply_repair_requests(design_ir, [{
        "page": 1, "rule": "overflow", "issues": ["text/overflow"],
    }])
    body = repaired["pages"][0]["slots"]["body"]
    assert repaired["pages"][0]["layout_variant"] == "wide"
    assert body["size_pt"] == pytest.approx(16.56)
    assert body["w"] == pytest.approx(12.333)

    slide = compile_slide(
        repaired["pages"][0], repaired["theme"]["colors"],
        repaired["theme"]["fonts"], repaired["theme"]["type_scale"],
        {"facts": []},
    )
    rendered_body = next(item for item in slide.components if item.id == "slide-0001-body")
    assert rendered_body.w == pytest.approx(12.333)
    assert rendered_body.style["font"]["size_pt"] == pytest.approx(16.56)
