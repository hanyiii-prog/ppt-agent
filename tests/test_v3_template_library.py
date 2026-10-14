import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402

from ppt_agent.composition_compiler import compile_presentation  # noqa: E402
from ppt_agent.evidence_graph import build_evidence_graph  # noqa: E402
from ppt_agent.pipeline import compile_page_design_stage  # noqa: E402
from ppt_agent.renderers import render_with  # noqa: E402
from ppt_agent.source_pack import build_source_pack  # noqa: E402
from ppt_agent.template_library import template_ids  # noqa: E402


MARKDOWN = """# Generic Template Fixture

## Objective

Explain the source in a calm, generic layout.

- First traceable point
- Second traceable point
"""


@pytest.mark.parametrize("template_id", template_ids())
def test_every_builtin_template_actually_renders(tmp_path, template_id):
    source = tmp_path / "source.md"
    source.write_text(MARKDOWN, encoding="utf-8")
    result = compile_page_design_stage(
        source,
        task={"objective": "generic template rendering check", "max_pages": 4, "template": template_id},
    )
    document, pack = build_source_pack(source)
    presentation = compile_presentation(
        result["page_design_ir"], build_evidence_graph(pack)
    )
    output = tmp_path / f"{template_id}.pptx"
    render_result = render_with(presentation, output)
    assert render_result.editable
    assert output.is_file() and output.stat().st_size > 0
    assert len(Presentation(str(output)).slides) == len(presentation.slides)
