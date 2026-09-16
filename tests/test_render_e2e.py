from pathlib import Path

from pptx import Presentation as PptxPresentation

from ppt_agent.dna_to_ir import template_dna_to_ir
from ppt_agent.ir import Presentation
from ppt_agent.renderer import render_presentation
from ppt_agent.template import analyze_pptx


def test_pptx_dna_ir_render_round_trip(tmp_path: Path):
    source_ir = Presentation.from_dict({
        "version": "0.1",
        "metadata": {"title": "Round Trip"},
        "theme": {"slide_size_inches": {"width": 13.333, "height": 7.5}},
        "slides": [
            {"id": "s1", "purpose": "cover", "components": [{"type": "title", "text": "Hello World"}]},
            {"id": "s2", "purpose": "content", "components": [
                {"type": "text", "text": "Point A", "x": 1.0, "y": 1.0, "w": 4.0, "h": 1.0},
                {"type": "text", "text": "Point B", "x": 1.0, "y": 2.5, "w": 4.0, "h": 1.0},
            ]},
        ],
    })

    first = render_presentation(source_ir, tmp_path / "first.pptx")
    assert len(PptxPresentation(str(first)).slides) == 2

    dna = analyze_pptx(first)
    assert dna["presentation"]["slide_count"] == 2
    assert dna["schema"] == "template-dna/v0.3"

    ir = template_dna_to_ir(dna, source="first.pptx")
    assert len(ir.slides) == 2

    second = render_presentation(ir, tmp_path / "second.pptx")
    assert len(PptxPresentation(str(second)).slides) == 2
