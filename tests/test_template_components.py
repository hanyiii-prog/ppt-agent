from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt

from ppt_agent.template_components import (
    build_template_component_catalog,
    rank_template_components,
)


def _template(tmp_path):
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "项目汇报"

    content = prs.slides.add_slide(prs.slide_layouts[5])
    content.shapes.title.text = "项目组织保障"
    for index, text in enumerate(("院方统筹", "信息科", "卫宁团队", "临床科室")):
        box = content.shapes.add_textbox(
            Inches(0.7 + index * 2.9),
            Inches(2.0),
            Inches(2.5),
            Inches(1.1),
        )
        box.text_frame.text = text
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(16)

    table_slide = prs.slides.add_slide(prs.slide_layouts[5])
    table_slide.shapes.title.text = "责任分工"
    table = table_slide.shapes.add_table(
        3, 3, Inches(0.8), Inches(1.8), Inches(11.8), Inches(3.2)
    ).table
    for row, values in enumerate(
        (("角色", "职责", "状态"), ("院方", "验收", "进行中"), ("卫宁", "开发", "计划"))
    ):
        for col, value in enumerate(values):
            table.cell(row, col).text = value

    path = tmp_path / "component-template.pptx"
    prs.save(path)
    return path


def test_catalog_keeps_single_page_components(tmp_path) -> None:
    catalog = build_template_component_catalog(_template(tmp_path))

    assert catalog["schema"] == "template-component-catalog/v1"
    assert catalog["component_count"] >= 3
    kinds = {item["component_kind"] for item in catalog["components"]}
    assert {"cover", "org_chart", "table"} <= kinds
    assert any(
        len(item["slots"]) >= 4
        for item in catalog["components"]
        if item["component_kind"] == "org_chart"
    )


def test_rank_returns_template_candidates_before_fallback(tmp_path) -> None:
    catalog = build_template_component_catalog(_template(tmp_path))

    candidates = rank_template_components(
        catalog,
        component_kind="org_chart",
        tags=("组织", "保障", "职责"),
        item_count=4,
    )

    assert candidates
    assert candidates[0]["source"] == "template"
    assert candidates[0]["component_kind"] == "org_chart"
    assert candidates[0]["score"] >= 0.5
