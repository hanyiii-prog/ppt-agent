import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from ppt_agent.clone_build import CloneBuildError  # noqa: E402
from ppt_agent.composition_compiler import compile_presentation  # noqa: E402
from ppt_agent.editability import audit_pptx_editability  # noqa: E402
from ppt_agent.ir import Presentation as IRPresentation  # noqa: E402
from ppt_agent.renderers import render_with  # noqa: E402


def test_compile_rejects_an_image_slot_as_non_editable_content():
    design_ir = {
        "schema": "page-design-ir/v1",
        "title": "Editable only",
        "theme": {
            "colors": {"text": "1F2937"},
            "fonts": {"body": "Test Sans"},
            "type_scale": {"body": 18},
        },
        "pages": [{
            "page_id": "slide-0001",
            "kind": "content",
            "title": "Reference screenshot",
            "slots": {
                "evidence": {
                    "id": "evidence",
                    "kind": "screenshot",
                    "x": 1.0,
                    "y": 1.0,
                    "w": 5.0,
                    "h": 3.0,
                    "value": "template screenshot",
                },
            },
        }],
    }

    with pytest.raises(ValueError, match="native and editable"):
        compile_presentation(design_ir, {"facts": []})


def test_saved_deck_audit_keeps_native_text_with_a_template_picture(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "Editable title"
    picture_path = tmp_path / "inherited-logo.png"
    from PIL import Image

    Image.new("RGB", (16, 16), "white").save(picture_path)
    slide.shapes.add_picture(str(picture_path), Inches(7), Inches(1), Inches(1), Inches(1))
    deck = tmp_path / "template-media.pptx"
    prs.save(deck)

    report = audit_pptx_editability(deck)

    assert report["passed"]
    assert report["pages"][0]["native_text_shape_count"] == 1
    assert report["pages"][0]["slide_local_picture_count"] == 1


def test_saved_deck_audit_blocks_a_picture_only_page(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture_path = tmp_path / "reference-screenshot.png"
    from PIL import Image

    Image.new("RGB", (80, 45), "white").save(picture_path)
    picture = slide.shapes.add_picture(
        str(picture_path), Inches(1), Inches(1), Inches(4), Inches(3)
    )
    picture.name = "Reference screenshot"
    deck = tmp_path / "screenshot-page.pptx"
    prs.save(deck)

    report = audit_pptx_editability(deck)

    assert not report["passed"]
    assert any("native editable" in issue for issue in report["pages"][0]["issues"])
    assert any("Reference screenshot" in issue for issue in report["pages"][0]["issues"])


def test_page_plan_rejects_screenshot_media_before_the_template_is_opened():
    with pytest.raises(CloneBuildError, match="non-editable media"):
        from ppt_agent.clone_build import render_clone_deck

        render_clone_deck(
            "unused-template.pptx",
            [{
                "role": "content",
                "kit": "column_cards",
                "title": "Results",
                "cards": [{"title": "Result", "lines": ["Native text"]}],
                "screenshot_path": "reference.png",
            }],
            "unused-output.pptx",
            audit=False,
            fidelity=False,
        )


def test_builtin_pipeline_writes_an_editability_gate(tmp_path):
    from pptx import Presentation as PptxPresentation

    from ppt_agent.pipeline import run_v3_pipeline

    source = tmp_path / "report.md"
    source.write_text(
        "# Editability Report\n\n- Native fact one\n- Native fact two\n",
        encoding="utf-8",
    )
    output = tmp_path / "workspace"
    report = run_v3_pipeline(
        source,
        output,
        task={"objective": "verify native editability", "max_pages": 3},
    )

    assert report["editability"]["passed"]
    deck = PptxPresentation(str(output / "deck.pptx"))
    assert len(deck.slides) == report["page_total"]
    for page in report["editability"]["pages"]:
        assert page["native_text_shape_count"] > 0
        assert page["native_editable_shape_count"] > 0


def test_mcp_audit_editability_tool(tmp_path):
    from ppt_agent.mcp.tools import ToolContext, tool_audit_editability
    from ppt_agent.sdk import PptAgent

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "Editable"
    deck = tmp_path / "ok.pptx"
    prs.save(deck)

    context = ToolContext(agent=PptAgent(), workspace=tmp_path)
    report = tool_audit_editability({"pptx": "ok.pptx"}, context)
    assert report["passed"]




def test_renderer_does_not_insert_a_picture_for_an_image_component(tmp_path):
    payload = {
        "version": "1.0",
        "metadata": {"title": "No raster"},
        "slides": [{
            "id": "s1",
            "purpose": "content",
            "components": [{
                "type": "image",
                "x": 1.0,
                "y": 1.0,
                "w": 3.0,
                "h": 2.0,
                "data": {"path": str(tmp_path / "reference.png")},
            }],
        }],
    }
    output = tmp_path / "guarded.pptx"
    render_with(IRPresentation.from_dict(payload), output)
    slide = Presentation(str(output)).slides[0]

    assert not any(shape.shape_type == 13 for shape in slide.shapes)
    assert any("[image]" in shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
