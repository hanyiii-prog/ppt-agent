# -*- coding: utf-8 -*-
"""Tests for ppt_agent.palette -- area-weighted, schemeClr-resolved palette.

The regression this guards: a deck painted mostly with THEME colours
(schemeClr accent1 over big header bars and cards) used to be misread as its
minority hard-coded colour, because the old logic counted literal srgbClr
occurrences. The 天津市口腔医院 template has blue accent1 painted across
~1300 sq in and purple hard-coded in 34 small shapes -- the eye says blue;
the old counter said purple, and generated decks fought their own template.
"""
import os
import re
import zipfile

import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches  # noqa: E402

from ppt_agent.palette import dominant_colors, hue_family, resolve_scheme  # noqa: E402

BRAND_BLUE = "2080FF"   # painted via schemeClr accent1, big areas
LITERAL_PURPLE = "7030A0"  # painted via srgbClr, many small shapes


def _deck_painted_with_theme(tmp_path):
    """Synthetic deck: 10 big schemeClr(accent1) banners + 12 tiny purple
    chips. Visually BLUE; occurrence-counting says purple."""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    for i in range(10):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.2 + i * 0.7),
                                Inches(12.9), Inches(0.6))
        # paint with schemeClr accent1 instead of a literal
        sh.fill.solid()
        spPr = sh._element.spPr
        for el in spPr.findall(qn('a:solidFill')):
            spPr.remove(el)
        fill = spPr.makeelement(qn('a:solidFill'), {})
        clr = fill.makeelement(qn('a:schemeClr'), {'val': 'accent1'})
        fill.append(clr)
        ln = spPr.find(qn('a:ln'))
        if ln is not None:
            ln.addprevious(fill)
        else:
            spPr.append(fill)
        sh.line.fill.background()
    for i in range(12):
        sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.3 + (i % 6) * 2.0),
                                Inches(0.3 + (i // 6) * 0.3), Inches(0.25), Inches(0.25))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(LITERAL_PURPLE)
        sh.line.fill.background()
    path = str(tmp_path / "themed.pptx")
    prs.save(path)
    # point theme accent1 at BRAND_BLUE by rewriting the theme part
    import shutil
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item.endswith('theme1.xml'):
                txt = data.decode('utf-8')
                txt = re.sub(r'(<a:accent1>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}(")',
                             r'\g<1>' + BRAND_BLUE + r'\g<2>', txt)
                data = txt.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)
    return path


def test_hue_family_buckets():
    assert hue_family("1185FE") == "blue"
    assert hue_family("7030A0") == "violet"
    assert hue_family("00DFFD") == "cyan"
    assert hue_family("FFFFFF") == "white/grey"
    assert hue_family("000000") == "near-black"


def test_resolve_scheme_lummod():
    theme = {"accent1": "1185FE"}
    assert resolve_scheme("accent1", theme) == "1185FE"
    out = resolve_scheme("accent1", theme, '<a:lumMod val="60000"/>')
    r, g, b = (int(out[i:i + 2], 16) for i in (0, 2, 4))
    src = (int("1185FE"[i:i + 2], 16) for i in (0, 2, 4))
    assert all(o <= s for o, s in zip((r, g, b), src))  # 60% lum darkens


def test_theme_painted_deck_reads_as_theme_colour(tmp_path):
    path = _deck_painted_with_theme(tmp_path)
    res = dominant_colors(path)
    assert res["primary"] == "blue"
    assert res["primary_rgb"] == BRAND_BLUE
    # the theme colour must outrank the literal purple on AREA
    areas = {c["rgb"]: c["area"] for c in res["colors"]}
    assert areas.get(BRAND_BLUE, 0) > areas.get(LITERAL_PURPLE, 0)
    # occurrence-counting would have picked purple (12 > 10 shapes) -- the
    # palette must NOT mirror that mistake
    counts = {c["rgb"]: c["n"] for c in res["colors"]}
    assert counts.get(LITERAL_PURPLE, 0) > counts.get(BRAND_BLUE, 0)


def test_theme_from_dna_prefers_area_palette(tmp_path):
    from ppt_agent.template import analyze_pptx
    from ppt_agent.theme import theme_from_dna
    path = _deck_painted_with_theme(tmp_path)
    dna = analyze_pptx(path)
    assert dna["dominant_palette"]["primary"] == "blue"
    theme = theme_from_dna(dna)
    assert theme.primary == BRAND_BLUE
