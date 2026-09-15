from pathlib import Path

import pytest

from ppt_agent.template import analyze_pptx


@pytest.fixture()
def sample_pptx(tmp_path: Path) -> Path:
    pptx = pytest.importorskip("pptx")
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn

    prs = pptx.Presentation()
    prs.slide_width = 13 * 914400
    prs.slide_height = 7 * 914400

    first = prs.slides.add_slide(prs.slide_layouts[6])
    first.background.fill.solid()
    first.background.fill.fore_color.rgb = RGBColor(10, 20, 30)
    bottom = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, 2000000, 1000000)
    bottom.fill.solid()
    bottom.fill.fore_color.rgb = RGBColor(255, 0, 0)
    top = first.shapes.add_shape(MSO_SHAPE.RECTANGLE, 1000000, 500000, 2000000, 1000000)
    top.fill.solid()
    top.fill.fore_color.rgb = RGBColor(0, 255, 0)
    solid = top._element.spPr.find(qn("a:solidFill"))
    color = next(iter(solid))
    alpha = color.makeelement(qn("a:alpha"), {"val": "65000"})
    color.append(alpha)
    top.line.width = 12700
    text = first.shapes.add_textbox(3000000, 1000000, 3000000, 1000000)
    text.text_frame.text = "Cover"
    text.text_frame.paragraphs[0].runs[0].font.name = "Aptos"
    text.text_frame.paragraphs[0].runs[0].font.size = pptx.util.Pt(28)

    body = prs.slides.add_slide(prs.slide_layouts[0])
    body.shapes.add_textbox(100000, 100000, 1000000, 500000).text_frame.text = "Body"
    last = prs.slides.add_slide(prs.slide_layouts[0])
    last.shapes.add_textbox(100000, 100000, 1000000, 500000).text_frame.text = "End"

    out = tmp_path / "sample.pptx"
    prs.save(out)
    return out


def test_template_dna_captures_semantics(sample_pptx: Path):
    dna = analyze_pptx(sample_pptx)

    assert dna["schema"] == "template-dna/v0.2"
    assert dna["presentation"]["slide_count"] == 3
    assert dna["special_surfaces"]["first"]["role"] == "first"
    assert dna["special_surfaces"]["last"]["role"] == "last"
    assert dna["special_surfaces"]["body_slide_count"] == 1
    assert dna["presentation"]["slide_size_inches"] == {"width": 13.0, "height": 7.0}

    first_shapes = dna["slides"][0]["shapes"]
    assert first_shapes[0]["z_index"] == 0
    assert first_shapes[1]["z_index"] == 1
    assert first_shapes[2]["z_index"] == 2
    assert first_shapes[1]["style"]["fill"]["rgb"] == "FF0000"
    assert first_shapes[2]["style"]["fill"]["transparency"] == pytest.approx(0.35, abs=0.01)
    assert first_shapes[2]["geometry"]["width"] > 0
    assert any(f[0] == "Aptos" for f in dna["global_style_statistics"]["fonts"])
    assert dna["theme"]["colors"]
    assert dna["masters"]
    assert dna["masters"][0]["layouts"]
