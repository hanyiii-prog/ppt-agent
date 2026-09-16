"""Template DNA -> theme derivation, and the `build --template` path."""
from __future__ import annotations

from collections import Counter

import pytest

from ppt_agent.theme import DEFAULT_THEME, theme_from_dna

BRAND = "7B1E3A"
ACCENT = "0E7C7B"


def _dna(*, fills, fonts, sizes, colors=None):
    return {
        "schema": "template-dna/v0.3",
        "theme": {"colors": colors or {}},
        "global_style_statistics": {
            "fills_rgb": fills,
            "fonts": fonts,
            "font_sizes_pt": sizes,
        },
    }


def _sample_dna():
    return _dna(
        fills=[[BRAND, 5], [ACCENT, 3]],
        fonts=[["SimSun", 5]],
        sizes=[[40.0, 1], [28.0, 1], [24.0, 1], [16.0, 1], [12.0, 1]],
    )


def test_theme_from_dna_takes_the_deck_colours():
    theme = theme_from_dna(_sample_dna(), name="t")
    assert theme.primary == BRAND
    assert theme.accent == ACCENT


def test_theme_from_dna_takes_the_deck_font():
    theme = theme_from_dna(_sample_dna())
    assert theme.font_title == "SimSun"
    assert theme.font_body == "SimSun"


def test_theme_from_dna_maps_the_type_ladder():
    theme = theme_from_dna(_sample_dna())
    assert theme.cover_title_pt == 40.0
    assert theme.section_title_pt == 28.0
    assert theme.slide_title_pt == 24.0
    assert theme.body_pt == 16.0
    assert theme.footer_pt == 12.0
    assert (
        theme.cover_title_pt
        > theme.section_title_pt
        > theme.slide_title_pt
        > theme.body_pt
        > theme.footer_pt
    )


def test_theme_from_dna_differs_from_the_builtin_preset():
    theme = theme_from_dna(_sample_dna())
    assert theme.name == "template"
    assert theme.primary != DEFAULT_THEME.primary
    assert theme.accent != DEFAULT_THEME.accent
    assert theme.font_title != DEFAULT_THEME.font_title


def test_theme_from_dna_ignores_a_stock_office_dark_slot():
    """A leftover Office `dk2` (blue) must not become the body ink."""
    dna = _sample_dna()
    dna["theme"]["colors"] = {"dk2": "1F497D"}
    assert theme_from_dna(dna).text != "1F497D"


def test_theme_from_dna_trusts_a_same_hue_dark_slot():
    dna = _sample_dna()
    dna["theme"]["colors"] = {"dk1": "3A0E1C"}
    assert theme_from_dna(dna).text == "3A0E1C"


def test_theme_from_dna_rotates_the_hue_when_no_accent_survives():
    dna = _dna(fills=[[BRAND, 6]], fonts=[["SimSun", 1]], sizes=[[24.0, 1]])
    theme = theme_from_dna(dna)
    assert theme.accent != theme.primary
    assert len(theme.accent) == 6


def test_theme_from_dna_degrades_on_junk_input():
    for junk in (None, 42, "not-dna", []):
        theme = theme_from_dna(junk)
        assert theme.name == "template"
        assert theme.primary == DEFAULT_THEME.primary


def test_theme_from_dna_clamps_absurd_metrics():
    dna = _dna(
        fills=[[BRAND, 1]],
        fonts=[["SimSun", 1]],
        sizes=[[999.0, 1], [0.5, 1]],
    )
    theme = theme_from_dna(dna)
    assert 26.0 <= theme.cover_title_pt <= 54.0
    assert 11.0 <= theme.body_pt <= 22.0
    assert 8.0 <= theme.footer_pt <= 13.0


def _require_pptx():
    pytest.importorskip("pptx", reason="python-pptx is required for template analysis")


def _reference_deck(path):
    from pptx import Presentation as Pptx
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Pptx()
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    for index in range(5):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.4 + index * 0.1), Inches(0.4 + index * 0.9), Inches(9.0), Inches(0.6),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(BRAND)
        shape.line.fill.background()
    for index in range(3):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(10.2), Inches(0.4 + index * 1.4), Inches(2.6), Inches(0.9)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
        shape.line.fill.background()

    text_slide = prs.slides.add_slide(blank)
    for index, (size, label) in enumerate([(40, "封面"), (28, "章节"), (24, "标题"), (16, "正文"), (12, "页脚")]):
        box = text_slide.shapes.add_textbox(Inches(0.8), Inches(0.4 + index * 1.2), Inches(9.0), Inches(0.9))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = label
        run.font.size = Pt(size)
        run.font.name = "SimSun"

    prs.save(str(path))
    return path


def _painted_fills(pptx_path):
    from pptx import Presentation as Pptx

    prs = Pptx(str(pptx_path))
    fills = Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            try:
                if int(shape.fill.type) == 1:
                    fills[str(shape.fill.fore_color.rgb)] += 1
            except (TypeError, ValueError, AttributeError):
                continue
    return fills


MARKDOWN = (
    "# 正畸专病库\n\n"
    "## 封面\n正畸专病库数据抽取方案\n2026 年 9 月\n\n"
    "## 字段取舍\n保留患者基本信息与就诊时间序列\n排除影像原始文件\n"
)


def test_dna_round_trips_into_a_theme(tmp_path):
    _require_pptx()
    from ppt_agent.template import analyze_pptx

    deck = _reference_deck(tmp_path / "reference.pptx")
    theme = theme_from_dna(analyze_pptx(deck), name="t")
    assert theme.primary == BRAND
    assert theme.accent == ACCENT
    assert theme.font_title == "SimSun"
    assert theme.cover_title_pt == 40.0


def test_build_paints_with_the_template_palette(tmp_path):
    _require_pptx()
    from ppt_agent.sdk import PptAgent

    deck = _reference_deck(tmp_path / "reference.pptx")
    outcome = PptAgent().build(
        out_dir=tmp_path / "out", markdown=MARKDOWN, template=deck, stem="deck",
        emit_html=False, gate=False,
    )
    assert outcome.slide_count > 0
    fills = _painted_fills(outcome.pptx_path)
    assert fills.get(BRAND, 0) > 0
    assert fills.get(ACCENT, 0) > 0
    assert fills.get(DEFAULT_THEME.primary, 0) == 0


def test_build_without_a_template_keeps_the_builtin_preset(tmp_path):
    _require_pptx()
    from ppt_agent.sdk import PptAgent

    outcome = PptAgent().build(
        out_dir=tmp_path / "out", markdown=MARKDOWN, stem="deck",
        emit_html=False, gate=False,
    )
    fills = _painted_fills(outcome.pptx_path)
    assert fills.get(DEFAULT_THEME.primary, 0) > 0
    assert fills.get(BRAND, 0) == 0


def test_build_records_the_template_theme_in_the_ir(tmp_path):
    _require_pptx()
    import json

    from ppt_agent.sdk import PptAgent

    deck = _reference_deck(tmp_path / "reference.pptx")
    outcome = PptAgent().build(
        out_dir=tmp_path / "out", markdown=MARKDOWN, template=deck, stem="deck",
        emit_html=False, gate=False,
    )
    theme = json.loads((tmp_path / "out" / "deck.json").read_text(encoding="utf-8"))["theme"]
    assert theme["primary"] == BRAND
    assert theme["name"] == "template:reference"
