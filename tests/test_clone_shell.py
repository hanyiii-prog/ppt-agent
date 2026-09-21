# -*- coding: utf-8 -*-
"""Tests for ppt_agent.clone_shell (template-page cloning renderer)."""
import os

import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from ppt_agent.clone_shell import (CloneShell, classify_shells,  # noqa: E402
                                   clear_body, set_title, ShellExhausted,
                                   audit_pages, drop_empty_placeholders)

ROLE_MAP = (("章节", "section"),
            ("标题幻灯片", "cover"),
            ("内容", "content"))


@pytest.fixture()
def mini_template(tmp_path):
    """Tiny 4-slide template built from default layouts, roles renamed to
    match the CJK classification map."""
    prs = Presentation()
    l0 = prs.slide_layouts[0]   # Title Slide
    l2 = prs.slide_layouts[2]   # Section Header
    l5 = prs.slide_layouts[5]   # Title Only
    # rename layouts so classify_shells' substring matching works
    l0.name = "标题幻灯片"
    l2.name = "章节标题页"
    l5.name = "内容页 - 有标题"
    for lay in (l0, l2, l5, l5):
        prs.slides.add_slide(lay)
    p = tmp_path / "tpl.pptx"
    prs.save(str(p))
    return str(p)


def test_classify_buckets(mini_template):
    prs = Presentation(mini_template)
    buckets = classify_shells(prs, ROLE_MAP)
    assert buckets["cover"] == [0]
    assert buckets["section"] == [1]
    assert buckets["content"] == [2, 3]


def test_take_clears_body_and_set_title(mini_template):
    d = CloneShell(mini_template, ROLE_MAP)
    idx, sl = d.take("content")
    assert idx == 2
    # C-route: no placeholder survives take(); all text is drawn by kits.
    assert not list(sl.placeholders)
    assert not set_title(sl, "hello")
    d.close()


def test_finish_prunes_and_reorders(mini_template):
    d = CloneShell(mini_template, ROLE_MAP)
    i0, _ = d.take("cover", cleared=False)
    i2, _ = d.take("content")
    out = os.path.join(os.path.dirname(mini_template), "out.pptx")
    d.finish(out, order=[i2, i0])
    prs = Presentation(out)
    assert len(prs.slides._sldIdLst) == 2
    assert prs.slides[0].slide_layout.name == "内容页 - 有标题"
    assert prs.slides[1].slide_layout.name == "标题幻灯片"
    os.remove(out)
    d.close()


def test_shell_exhausted(mini_template):
    d = CloneShell(mini_template, ROLE_MAP)
    d.take("cover", cleared=False)
    with pytest.raises(ShellExhausted):
        d.take("cover")
    d.close()


def test_clear_body_keeps_only_title(mini_template):
    prs = Presentation(mini_template)
    sl = prs.slides[0]
    sl.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    removed = clear_body(sl)
    assert removed >= 1
    assert all(sh.is_placeholder and sh.placeholder_format.idx == 0
               for sh in sl.shapes)


def test_audit_pages_detects_overflow(mini_template):
    prs = Presentation(mini_template)
    sl = prs.slides[2]
    tf = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.1)).text_frame
    tf.word_wrap = True
    tf.text = "这" * 120  # 120 CJK chars in a 0.1in-tall box -> overflow
    issues = audit_pages(prs)
    assert any(i["page"] == 3 and i["kind"] == "overflow" for i in issues)


def test_audit_flags_empty_placeholder_from_stock_layout(mini_template):
    """The fixture reuses python-pptx's stock layouts, whose title placeholder
    ships the skeleton heading 'Click to edit Master title style'. A shell that
    keeps its empty title slot therefore renders that string -- the exact class
    the 目录页 defect belonged to."""
    prs = Presentation(mini_template)
    issues = audit_pages(prs)
    assert {i["kind"] for i in issues} <= {"empty", "stale_placeholder"}
    # every page is flagged (a bare shell may leave more than one slot behind,
    # e.g. the title-slide layout ships title + subtitle)
    assert {i["page"] for i in issues if i["kind"] == "stale_placeholder"} == \
        set(range(1, len(prs.slides) + 1))
    assert all("Click to edit Master" in i["msg"]
               for i in issues if i["kind"] == "stale_placeholder")
    # the fix, not a suppression: drop the dead slots and the check goes quiet
    for sl in prs.slides:
        drop_empty_placeholders(sl)
    assert not [i for i in audit_pages(prs) if i["kind"] == "stale_placeholder"]


def test_audit_duplicate_ignores_card_field_labels(mini_template):
    """N cards sharing one short field label is by design, not a bug."""
    prs = Presentation(mini_template)
    sl = prs.slides[2]
    for x in (1.0, 3.0, 5.0):
        sl.shapes.add_textbox(Inches(x), Inches(1), Inches(2), Inches(0.3)) \
            .text_frame.text = "业务特征与难点"
    issues = audit_pages(prs)
    assert not any(i["page"] == 3 and i["kind"] == "duplicate" for i in issues)


def test_audit_duplicate_flags_text_drawn_twice(mini_template):
    """The same string painted exactly twice is the chip+beside bug pattern."""
    prs = Presentation(mini_template)
    sl = prs.slides[2]
    for x in (1.0, 3.0):
        sl.shapes.add_textbox(Inches(x), Inches(1), Inches(2), Inches(0.3)) \
            .text_frame.text = "临床科研应用端"
    issues = audit_pages(prs)
    assert any(i["page"] == 3 and i["kind"] == "duplicate" for i in issues)


def test_audit_duplicate_flags_long_repeated_sentence(mini_template):
    prs = Presentation(mini_template)
    sl = prs.slides[2]
    line = "复用卫宁底座与朗视影像资产打通数据链路"
    for x in (1.0, 3.0, 5.0):
        sl.shapes.add_textbox(Inches(x), Inches(1), Inches(4), Inches(0.4)) \
            .text_frame.text = line
    issues = audit_pages(prs)
    assert any(i["page"] == 3 and i["kind"] == "duplicate" for i in issues)


# ---------------------------------------------------------------------------
# V7 learnings: content chrome, repositioned title, geom patch, cover rebuild
# ---------------------------------------------------------------------------

def test_add_content_chrome_clears_template_decorations(mini_template):
    """chrome() must clear_body first so stray template shapes never leak
    through (the root cause of the V6->V7 residual-decoration bug)."""
    from pptx.enum.shapes import MSO_SHAPE
    from ppt_agent.clone_shell import (add_content_chrome, set_title,
                                       CHROME_BAR)
    d = CloneShell(mini_template, ROLE_MAP)
    _, sl = d.take("content", cleared=False)
    # simulate a leftover template decoration (stray photo + accent text)
    sl.shapes.add_picture(_blank_png(), Inches(1), Inches(2), Inches(2), Inches(2))
    sl.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(0.5)).text = "旧装饰"
    add_content_chrome(sl, prs=d.prs)
    # bar present at top
    bar = [s for s in sl.shapes
           if abs(s.top or 0) < 1000 and (s.width or 0) > 12 * 914400
           and s.fill.fore_color.rgb == __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor.from_string(CHROME_BAR)]
    assert bar, "content bar missing"
    # no leftover stray text
    assert not any("旧装饰" in (s.text_frame.text or "")
                   for s in sl.shapes if s.has_text_frame)
    # C-route: no placeholder survives; the chrome is layout-inherited.
    assert not list(sl.placeholders)
    d.close()


def test_set_geom_patches_preset(mini_template):
    from pptx.oxml.ns import qn
    from ppt_agent.clone_shell import box, set_geom, CHROME_BAR
    d = CloneShell(mini_template, ROLE_MAP)
    _, sl = d.take("content", cleared=False)
    pill = box(sl, 4.38, 4.81, 4.87, 0.61, CHROME_BAR, MSO_SHAPE.ROUNDED_RECTANGLE)
    set_geom(pill, "round2DiagRect")
    g = pill._element.spPr.find(qn("a:prstGeom"))
    assert g is not None and g.get("prst") == "round2DiagRect"
    d.close()


def test_rebuild_cover_z_order(mini_template):
    """Cover z-order must be: photo -> freeform bands -> logos -> pill -> title."""
    from pptx.oxml.ns import qn
    from ppt_agent.clone_shell import rebuild_cover
    d = CloneShell(mini_template, ROLE_MAP)
    idx, sl = d.take("cover", cleared=False)
    # give the cover a background photo + a freeform band to clone
    bg = sl.shapes.add_picture(_blank_png(), 0, -Inches(0.1), Inches(13.333), Inches(4.72))
    rebuild_cover(sl, prs=d.prs, title_text="专病数据库建设情况汇报",
                  meta_text="汇报部门：信息科")
    pics = [s for s in sl.shapes if s.shape_type == 13]
    assert pics, "background photo lost"
    # first picture sits at the very back (z-order 0)
    assert pics[0].top / 914400 < 0, "background photo not at back"
    # title text present and on top of chrome
    texts = [s.text_frame.text for s in sl.shapes if s.has_text_frame]
    assert any("专病数据库建设情况汇报" in t for t in texts)
    d.close()


# ---------------------------------------------------------------------------
# V9 learnings: rotation is a first-class attribute (python-pptx drops it)
# ---------------------------------------------------------------------------

def test_set_xfrm_roundtrips_rotation(mini_template):
    """`left/top/width/height` have nowhere to store rot/flip -> they were
    silently dropped on every redraw, rotating template decorations 90deg."""
    from ppt_agent.clone_shell import box, set_xfrm, shape_rot
    d = CloneShell(mini_template, ROLE_MAP)
    _, sl = d.take("content", cleared=False)
    band = box(sl, 3.55, -1.92, 3.91, 11.01, "1185FE")
    set_xfrm(band, 3.55, -1.92, 3.91, 11.01, rot=90.0, flip_h=True)
    rot, fh, fv = shape_rot(band)
    assert abs(rot - 90.0) < 0.01 and fh is True and fv is False
    # and clearing it again removes the attributes entirely
    set_xfrm(band, 3.55, -1.92, 3.91, 11.01)
    assert shape_rot(band) == (0.0, False, False)
    d.close()


def test_rotated_bbox_transposes_quarter_turn(mini_template):
    """A 3.91x11.01in frame rotated 90deg renders as 11.01x3.91in."""
    from ppt_agent.clone_shell import box, set_xfrm, rotated_bbox
    d = CloneShell(mini_template, ROLE_MAP)
    _, sl = d.take("content", cleared=False)
    band = box(sl, 3.55, -1.92, 3.91, 11.01, "1185FE")
    set_xfrm(band, 3.55, -1.92, 3.91, 11.01, rot=90.0, flip_h=True)
    L, T, W, H = rotated_bbox(band)
    assert abs(W - 11.01) < 0.02 and abs(H - 3.91) < 0.02, (W, H)
    # centre is preserved through the rotation
    assert abs((L + W / 2) - (3.55 + 3.91 / 2)) < 0.02
    assert abs((T + H / 2) - (-1.92 + 11.01 / 2)) < 0.02
    d.close()


def test_layout_chrome_detects_inherited_bar_and_logos(mini_template):
    """layout_chrome() is what tells the caller 'do NOT draw chrome here'."""
    from ppt_agent.clone_shell import layout_chrome
    prs = Presentation(mini_template)
    # plain layout: nothing to inherit on the fixture
    assert layout_chrome(prs.slides[2])["bar"] is False
    # simulate a template layout that paints a top bar + a logo
    _layout_bar(prs.slides[2].slide_layout)
    _layout_pic(prs.slides[2].slide_layout)
    inv = layout_chrome(prs.slides[2])
    assert inv["bar"] is True and inv["pictures"] >= 1


def test_add_content_chrome_inherits_instead_of_doubling(mini_template):
    """When the layout already paints chrome, the slide must draw NOTHING --
    redrawing doubled the logos and replaced the translucent wash with an
    opaque band (the '控件比 Kimi 差' root cause)."""
    from ppt_agent.clone_shell import add_content_chrome
    prs = Presentation(mini_template)
    _layout_bar(prs.slides[2].slide_layout)
    _layout_pic(prs.slides[2].slide_layout)
    sl = prs.slides[2]
    assert add_content_chrome(sl, prs=prs) == 0
    assert all(sh.is_placeholder for sh in sl.shapes), "chrome was redrawn on top"
    # force=True still supports shells whose layout has no chrome at all
    assert add_content_chrome(sl, prs=prs, force=True) > 0


def test_gradient_fill_writes_alpha_stops(mini_template):
    """The template 'bar' is 1185FE @15% -> @0%; alpha must survive."""
    from pptx.oxml.ns import qn
    from ppt_agent.clone_shell import box, gradient_fill
    d = CloneShell(mini_template, ROLE_MAP)
    _, sl = d.take("content", cleared=False)
    wash = box(sl, 0, 0, 13.333, 0.71, "1185FE")
    gradient_fill(wash, [(0, "1185FE", 15), (100, "1185FE", 0)],
                  angle=0, rot_with_shape=True)
    grad = wash._element.spPr.find(qn("a:gradFill"))
    assert grad is not None and grad.get("rotWithShape") == "1"
    alphas = [a.get("val") for a in grad.iter(qn("a:alpha"))]
    assert alphas == ["15000", "0"], alphas


def test_clone_shape_preserves_rotation(mini_template):
    """The correct way to reuse a layout decoration: deep-copy the XML."""
    from ppt_agent.clone_shell import box, set_xfrm, clone_shape, shape_rot
    d = CloneShell(mini_template, ROLE_MAP)
    _, src = d.take("section", cleared=False)
    band = box(src, 3.55, -1.92, 3.91, 11.01, "1185FE",
               MSO_SHAPE.ROUNDED_RECTANGLE)
    set_xfrm(band, 3.55, -1.92, 3.91, 11.01, rot=270.0)
    _, dst = d.take("content")
    copied = clone_shape(band, dst)
    assert shape_rot(copied)[0] == 270.0


_WAVE_XML = (
    '<p:sp %s><p:nvSpPr><p:cNvPr id="90" name="wave-1"/>'
    '<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
    '<a:xfrm><a:off x="0" y="0"/><a:ext cx="9144000" cy="2820000"/></a:xfrm>'
    '<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
    '<a:rect l="l" t="t" r="r" b="b"/><a:pathLst>'
    '<a:path w="9144000" h="2820000"><a:moveTo><a:pt x="0" y="0"/></a:moveTo>'
    '<a:lnTo><a:pt x="9144000" y="0"/></a:lnTo>'
    '<a:lnTo><a:pt x="0" y="2820000"/></a:lnTo><a:close/></a:path>'
    '</a:pathLst></a:custGeom>'
    '<a:gradFill rotWithShape="1"><a:gsLst>'
    '<a:gs pos="0"><a:srgbClr val="0061FA"><a:alpha val="75000"/></a:srgbClr>'
    '</a:gs>'
    '<a:gs pos="100000"><a:srgbClr val="1185FE"><a:alpha val="0"/></a:srgbClr>'
    '</a:gs></a:gsLst><a:lin ang="5400000" scaled="1"/></a:gradFill>'
    '<a:ln><a:noFill/></a:ln>'
    '</p:spPr></p:sp>')


def test_rebuild_closing_reuses_shell_media(mini_template):
    """Closing pages reuse the cover shell: photo + translucent wave bands are
    ALREADY on the slide and must be carried over, not redrawn.

    V9 cleared the body and drew a full-page opaque gradient plus an opaque
    rounded rect, so the photo vanished, the waves lost their alpha and the
    page rendered solid blue. This asserts the photo blob survives, the
    freeform band survives *with its alpha stops*, no full-page background is
    added, and the three closing lines are centred.
    """
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn
    from pptx.enum.text import PP_ALIGN
    from ppt_agent.clone_shell import rebuild_closing
    d = CloneShell(mini_template, ROLE_MAP)
    _, sl = d.take("cover", cleared=False)
    # shell media: an off-canvas background photo + a wave band + stale body
    sl.shapes.add_picture(_blank_png(), 0, Inches(-0.07),
                          Inches(13.333), Inches(4.72))
    sl.shapes._spTree.append(parse_xml(_WAVE_XML % nsdecls("p", "a")))
    stale = sl.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    stale.text_frame.text = "封面遗留正文"

    rebuild_closing(sl, prs=d.prs, title_text="打造口腔专病数据与临床科研标杆",
                    sub_text="恳请各位领导审议指正", meta_text="信息科 · 2026年9月")

    pics = [sh for sh in sl.shapes if sh.shape_type == 13]
    assert len(pics) == 1 and pics[0].top < 0, "background photo was dropped"
    waves = [sh for sh in sl.shapes if "FREEFORM" in str(sh.shape_type)]
    assert len(waves) == 1, "wave band was redrawn or lost"
    alphas = [a.get("val") for a in waves[0]._element.iter(qn("a:alpha"))]
    assert alphas == ["75000", "0"], alphas
    # no full-page opaque fill -- the V9 regression
    full = [sh for sh in sl.shapes
            if sh.width >= Inches(13.0) and sh.height >= Inches(7.0)]
    assert not full, "a full-page background was painted over the photo"
    texts = {sh.text_frame.text for sh in sl.shapes if sh.has_text_frame}
    assert "封面遗留正文" not in texts, "stale cover body survived"
    for want in ("打造口腔专病数据与临床科研标杆", "恳请各位领导审议指正",
                 "信息科 · 2026年9月"):
        hits = [sh for sh in sl.shapes
                if sh.has_text_frame and sh.text_frame.text == want]
        assert hits, want
        assert hits[0].text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER


def test_clone_shape_refuses_pictures(mini_template):
    """Pictures carry r:embed relationships -> must refuse loudly."""
    from ppt_agent.clone_shell import clone_shape
    d = CloneShell(mini_template, ROLE_MAP)
    _, dst = d.take("content", cleared=False)
    pic = dst.shapes.add_picture(_blank_png(), Inches(1), Inches(1),
                                 Inches(0.5), Inches(0.5))
    with pytest.raises(ValueError):
        clone_shape(pic, dst)
    d.close()


def test_audit_pages_flags_doubled_layout_decoration(mini_template):
    """The check that would have caught the whole V8 chrome-doubling class."""
    from ppt_agent.clone_shell import add_content_chrome
    prs = Presentation(mini_template)
    _layout_pic(prs.slides[2].slide_layout)
    sl = prs.slides[2]
    sl.shapes.add_picture(_blank_png(), Inches(12.23), Inches(0.09),
                          Inches(0.7), Inches(0.3))
    issues = audit_pages(prs)
    assert any(i["page"] == 3 and i["kind"] == "doubling" for i in issues)
    # the inherit route (add_content_chrome default) must stay clean
    prs2 = Presentation(mini_template)
    _layout_bar(prs2.slides[2].slide_layout)
    _layout_pic(prs2.slides[2].slide_layout)
    assert add_content_chrome(prs2.slides[2], prs=prs2) == 0
    assert not any(i["kind"] == "doubling" for i in audit_pages(prs2))


def _layout_bar(lay):
    """Append a raw full-width 0.71in bar to a layout shape tree.

    ``LayoutShapes`` is read-only in python-pptx (no ``add_shape``), so the
    decoration goes in as XML -- which is all ``layout_chrome`` reads anyway.
    """
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    lay.shapes._spTree.append(parse_xml(
        '<p:sp %s><p:nvSpPr><p:cNvPr id="90" name="bar"/>'
        '<p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="0" y="0"/>'
        '<a:ext cx="12192000" cy="649224"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:solidFill><a:srgbClr val="1185FE"/></a:solidFill></p:spPr>'
        '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
        % nsdecls("p", "a")))


def _layout_pic(lay):
    """Append a raw logo picture at (12.23, 0.09) 0.7x0.3in to a layout.

    The image relationship is registered for real (``get_or_add_image_part``),
    so ``shape_type`` / ``.image`` resolve instead of raising ``KeyError``.
    """
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
    _, rId = lay.part.get_or_add_image_part(_blank_png())
    lay.shapes._spTree.append(parse_xml(
        '<p:pic %s><p:nvPicPr><p:cNvPr id="91" name="logo"/>'
        '<p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
        '<p:blipFill><a:blip r:embed="%s"/><a:stretch><a:fillRect/>'
        '</a:stretch></p:blipFill>'
        '<p:spPr><a:xfrm><a:off x="11180160" y="82296"/>'
        '<a:ext cx="640080" cy="274320"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
        % (nsdecls("p", "a", "r"), rId)))


def _blank_png():
    """Valid 1x1 RGB PNG as a file-like object (built via zlib, no deps)."""
    import io
    import struct
    import zlib
    w = h = 1
    scan = b"\x00" + b"\x00" * (w * 3)
    sig = b"\x89PNG\r\n\x1a\n"

    def chunk(typ, data):
        c = typ + data
        return struct.pack(">I", len(data)) + c + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    idat = zlib.compress(scan)
    blob = sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")
    return io.BytesIO(blob)


# --- empty placeholders are dead DNA ----------------------------------------
def _layout_prompt(lay, text, idx=0):
    """Write skeleton text into the layout's own placeholder.

    The layout shape tree is read-only through the high-level API, so the run
    goes in as XML -- which is exactly what ``layout_placeholder_text`` reads.
    """
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn
    for sp in lay._element.iter(qn("p:sp")):
        nv = sp.find(qn("p:nvSpPr"))
        nvpr = nv.find(qn("p:nvPr")) if nv is not None else None
        ph = nvpr.find(qn("p:ph")) if nvpr is not None else None
        if ph is None or int(ph.get("idx") or 0) != idx:
            continue
        tx = sp.find(qn("p:txBody"))
        if tx is None:
            continue
        for p in tx.findall(qn("a:p")):
            tx.remove(p)
        tx.append(parse_xml('<a:p %s><a:r><a:t>%s</a:t></a:r></a:p>'
                            % (nsdecls("a"), text)))
        return True
    return False


def test_layout_placeholder_text_reads_the_twin(mini_template):
    """An empty slide placeholder resolves to its layout twin when rendered."""
    from ppt_agent.clone_shell import layout_placeholder_text
    prs = Presentation(mini_template)
    sl = prs.slides[2]
    # stock python-pptx layouts already carry a skeleton heading
    assert layout_placeholder_text(sl, 0) != ""
    assert _layout_prompt(sl.slide_layout, "单击此处编辑母版标题样式")
    assert layout_placeholder_text(sl, 0) == "单击此处编辑母版标题样式"
    # a layout placeholder that does not exist reads as empty, not as an error
    assert layout_placeholder_text(sl, 7) == ""


def test_drop_empty_placeholders_removes_the_dead_slot(mini_template):
    from ppt_agent.clone_shell import drop_empty_placeholders
    prs = Presentation(mini_template)
    empty = prs.slides[2]
    assert drop_empty_placeholders(empty) == 1
    assert not any(sh.is_placeholder for sh in empty.shapes)
    # a placeholder that carries text is left alone
    filled = prs.slides[3]
    set_title(filled, "二、建设计划 | 总体技术架构")
    assert drop_empty_placeholders(filled) == 0
    assert any(sh.is_placeholder for sh in filled.shapes)


def test_audit_flags_stale_placeholder_that_would_paint_the_prompt(mini_template):
    """The TOC-page defect: a blank PH0 whose layout twin still holds the
    skeleton heading -- renderers paint that string in the top slot."""
    from ppt_agent.clone_shell import drop_empty_placeholders
    prs = Presentation(mini_template)
    page3 = prs.slides[2]
    _layout_prompt(page3.slide_layout, "单击此处编辑母版标题样式")
    hits = [i for i in audit_pages(prs)
            if i["page"] == 3 and i["kind"] == "stale_placeholder"]
    assert hits, "empty placeholder backed by a layout prompt must be flagged"
    assert "单击此处编辑母版标题样式"[:6] in hits[0]["msg"]
    # a placeholder with text of its own is not stale
    set_title(prs.slides[3], "有标题")
    assert not any(i["page"] == 4 and i["kind"] == "stale_placeholder"
                   for i in audit_pages(prs))
    # dropping it is the fix -- the check goes quiet for that page
    drop_empty_placeholders(page3)
    assert not any(i["page"] == 3 and i["kind"] == "stale_placeholder"
                   for i in audit_pages(prs))
