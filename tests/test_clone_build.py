# -*- coding: utf-8 -*-
"""Tests for ppt_agent.clone_build (data-driven clone-shell builds) and the
three MCP tools that expose it (ppt_agent_clone_plan / clone_build / clone_audit)."""
import json

import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from ppt_agent.clone_build import (  # noqa: E402
    CloneBuildError,
    KITS,
    audit_deck,
    plan_template,
    render_clone_deck,
)


@pytest.fixture()
def mini_template(tmp_path):
    """2 cover + 1 section + 4 content shells, layouts renamed for the CJK
    role map. The double cover pool serves cover AND closing; the content pool
    serves content AND toc (the reference deck authors the TOC on it)."""
    prs = Presentation()
    l0 = prs.slide_layouts[0]
    l2 = prs.slide_layouts[2]
    l5 = prs.slide_layouts[5]
    l0.name = "标题幻灯片"
    l2.name = "章节标题页"
    l5.name = "内容页 - 有标题"
    for lay in (l0, l0, l2, l5, l5, l5, l5):
        prs.slides.add_slide(lay)
    p = tmp_path / "tpl.pptx"
    prs.save(str(p))
    return str(p)


def _plan():
    return [
        {"role": "cover", "pill": "天津市口腔医院",
         "title": "专病数据库建设情况汇报", "meta": "信息科 · 2026年9月"},
        {"role": "toc", "items": [("建设背景", "现状与诉求"), ("建设计划", "架构与功能")],
         "note": "立足务实原则，以点带面跑通试点。"},
        {"role": "section", "title": "一、建设背景", "lines": ["建设背景", "三大试点", "方案框架"]},
        {"role": "content", "kit": "quad_cards", "title": "建设背景 | 四个要点",
         "lead": "立足务实", "cards": [{"title": f"要点{i}", "body": "以点带面跑通试点"}
                                       for i in range(4)],
         "note": "说明：不追求一步到位。"},
        {"role": "content", "kit": "column_cards", "title": "建设计划 | 四大功能",
         "lead": "以智能标签为索引骨架",
         "cards": [{"title": f"功能{i}", "lines": ["说明文字"]} for i in range(3)],
         "cols": 3, "head_h": 0.6, "card_h": 2.2,
         "bottom": {"title": "预期效果", "body": "减少漏检", "h": 1.11, "y": 5.94}},
        {"role": "content", "kit": "progress_timeline", "title": "推进纪实",
         "lead": "前期准备",
         "steps": [{"date": "9.11~9.25", "status": "doing", "title": "收集标签",
                    "desc": "进行中", "current": True},
                   {"date": "9.30", "status": "plan", "title": "转开发", "desc": "待启动"}],
         "note": "材料收集与下发情况：正常。"},
        {"role": "closing", "title": "打造口腔专病数据与临床科研标杆",
         "sub": "恳请各位领导审议指正", "meta": "信息科 · 2026年9月"},
    ]


def test_plan_template_inventories_shells_and_kits(mini_template):
    info = plan_template(mini_template)
    assert info["shells"] == {"cover": 2, "section": 1, "content": 4}
    assert "cover" in (info["page_kind_counts"] or {})
    assert set(KITS) <= set(info["kits"])
    assert info["kind_summaries"], "per-kind DNA summaries missing"


def test_render_clone_deck_end_to_end(mini_template, tmp_path):
    out = tmp_path / "out.pptx"
    result = render_clone_deck(mini_template, _plan(), out)
    assert result["output"] == str(out)
    assert result["pages"] == 7
    assert result["audit"]["count"] == 0, result["audit"]["issues"]

    prs = Presentation(str(out))
    assert len(prs.slides._sldIdLst) == 7
    layouts = [s.slide_layout.name for s in prs.slides]
    assert layouts[0] == "标题幻灯片" and layouts[2] == "章节标题页"
    assert layouts[6] == "标题幻灯片"          # closing reuses the cover shell

    def texts(slide):
        return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]

    all_texts = [t for s in prs.slides for t in texts(s)]
    for want in ("专病数据库建设情况汇报", "目录", "CONTENTS", "01", "一、建设背景",
                 "要点3", "预期效果", "当前阶段", "打造口腔专病数据与临床科研标杆"):
        assert any(want in t for t in all_texts), want


def test_render_clone_deck_without_audit(mini_template, tmp_path):
    out = tmp_path / "out.pptx"
    result = render_clone_deck(mini_template, _plan()[:1], out, audit=False)
    assert "audit" not in result
    assert out.exists()


def test_header_only_content_page_warns(mini_template, tmp_path):
    out = tmp_path / "out.pptx"
    result = render_clone_deck(
        mini_template,
        [{"role": "content", "title": "只有页头"}],
        out, audit=False)
    assert any("only the header was drawn" in w for w in result["warnings"])


def test_unknown_role_rejected(mini_template, tmp_path):
    with pytest.raises(CloneBuildError, match="unknown page role"):
        render_clone_deck(mini_template, [{"role": "chapter"}], tmp_path / "o.pptx")


def test_unknown_kit_lists_valid_kits(mini_template, tmp_path):
    with pytest.raises(CloneBuildError, match="valid kits"):
        render_clone_deck(
            mini_template,
            [{"role": "content", "kit": "magic_cards", "title": "x"}],
            tmp_path / "o.pptx", audit=False)


def test_missing_kit_data_is_named(mini_template, tmp_path):
    with pytest.raises(CloneBuildError, match="missing required data: cards"):
        render_clone_deck(
            mini_template,
            [{"role": "content", "kit": "four_role_cards", "title": "x"}],
            tmp_path / "o.pptx", audit=False)


def test_short_shell_pools_reuse_shells(tmp_path, mini_template):
    """A plan with more pages than unique shells is NOT rejected -- content
    pages duplicate a template shell so a 13-slide template can serve an
    18-25 page deck (each page still inherits real DNA). Rejecting this was the
    bug that pushed real decks into the synthetic fallback route."""
    pages = [{"role": "content", "kit": "quad_cards", "title": f"p{i}",
              "cards": [{"title": "a", "body": "b"}]} for i in range(7)]
    result = render_clone_deck(mini_template, pages, tmp_path / "o.pptx",
                               audit=False, fidelity=False)
    assert result["pages"] == 7


def test_role_with_no_shell_fails_up_front(tmp_path):
    """A role the template cannot supply at all -- a cover on a content-only
    deck, where duplication has nothing to clone -- must still refuse up front."""
    from pptx import Presentation
    prs = Presentation()
    prs.slide_layouts[5].name = "内容页 - 列表版"
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[5])
    tpl = tmp_path / "no_cover.pptx"
    prs.save(str(tpl))
    with pytest.raises(CloneBuildError, match="template cannot serve this page plan"):
        render_clone_deck(str(tpl), [{"role": "cover", "title": "T"}],
                          tmp_path / "o.pptx", audit=False, fidelity=False)


def test_audit_deck_flags_overflow(mini_template, tmp_path):
    prs = Presentation(mini_template)
    sl = prs.slides[3]
    tf = sl.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.1)).text_frame
    tf.word_wrap = True
    tf.text = "这" * 120
    p = tmp_path / "bad.pptx"
    prs.save(str(p))
    report = audit_deck(p)
    assert report["count"] >= 1
    assert any(i["kind"] == "overflow" for i in report["issues"])


# --- MCP tool surface --------------------------------------------------------
def _context(tmp_path):
    from ppt_agent.mcp.tools import ToolContext
    from ppt_agent.sdk import PptAgent

    return ToolContext(agent=PptAgent(), workspace=tmp_path)


def test_tool_names_include_clone_tools():
    from ppt_agent.mcp.tools import tool_names

    names = tool_names()
    assert {"ppt_agent_clone_plan", "ppt_agent_clone_build",
            "ppt_agent_clone_audit"} <= set(names)
    assert len(names) == 19  # 13 base + 4 fidelity + narrative + generate (V2.1)


def test_clone_tools_round_trip(tmp_path):
    from ppt_agent.mcp.tools import call_tool

    context = _context(tmp_path)
    tpl = str(tmp_path / "tpl.pptx")
    _write_template(tpl)

    plan = call_tool("ppt_agent_clone_plan", {"template": tpl}, context)
    payload = json.loads(plan["content"][0]["text"])
    assert payload["shells"]["content"] == 4

    built = call_tool(
        "ppt_agent_clone_build",
        {"template": tpl, "pages": _plan(), "output": "decks/out.pptx"},
        context,
    )
    payload = json.loads(built["content"][0]["text"])
    assert payload["audit"]["count"] == 0
    assert (tmp_path / "decks" / "out.pptx").exists()

    audited = call_tool("ppt_agent_clone_audit",
                        {"pptx": str(tmp_path / "decks" / "out.pptx")}, context)
    payload = json.loads(audited["content"][0]["text"])
    assert payload["count"] == 0


def _write_template(path):
    """A template with the layout names the CJK role map expects + content DNA."""
    prs = Presentation()
    l0 = prs.slide_layouts[0]
    l2 = prs.slide_layouts[2]
    l5 = prs.slide_layouts[5]
    l0.name = "标题幻灯片"
    l2.name = "章节标题页"
    l5.name = "内容页 - 有标题"
    for _ in range(7):
        prs.slides.add_slide(l0 if len(prs.slides._sldIdLst) < 2 else
                            (l2 if len(prs.slides._sldIdLst) == 2 else l5))
    prs.save(path)
