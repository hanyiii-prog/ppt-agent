"""Programmatic mini-template builder shared by the batch-1 knowledge-layer tests.

Every fixture builds a real PPTX with python-pptx (repo convention: no binary
fixtures). The deck mirrors a typical corporate template:

* slide 1 (cover): display title + subtitle
* slides 2-3 (content): title + body + a *repeated* header bar and logo block
  at identical coordinates (slide-origin, so element/component detection has
  real recurring material)
* slide 4 (closing): display title
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from ppt_agent.page_dna import extract_deck_dna
from ppt_agent.semantic_role import annotate_deck_dna

BAR_RGB = "11506E"
LOGO_RGB = "C8861B"
TEXT_RGB = "1F2933"


def build_template_pptx(path: Path) -> Path:
    """Write the mini template deck and return the path."""
    pptx: Any = __import__("importlib").import_module("pptx")
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_bar(slide: Any) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.45))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(BAR_RGB)
        bar.name = "header-bar"

    def add_logo(slide: Any) -> None:
        logo = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.1), Inches(0.05), Inches(0.9), Inches(0.35))
        logo.fill.solid()
        logo.fill.fore_color.rgb = RGBColor.from_string(LOGO_RGB)
        logo.name = "logo-block"

    def add_text(slide: Any, left: float, top: float, width: float, height: float,
                 text: str, size_pt: int, rgb: str = TEXT_RGB) -> None:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        box.text_frame.text = text
        for paragraph in box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(size_pt)
                run.font.color.rgb = RGBColor.from_string(rgb)
                run.font.name = "Test Sans"

    # slide 1 -- cover
    cover = prs.slides.add_slide(blank)
    add_bar(cover)
    add_logo(cover)
    add_text(cover, 0.9, 2.6, 9.0, 1.4, "口腔医院互联互通项目上线总结", 32)
    add_text(cover, 0.9, 4.2, 7.0, 0.7, "2026 年度", 16)

    # slides 2-3 -- content pages with the same repeated header bar + logo
    for index, lines in enumerate((
        ["覆盖 5 个院区，服务 92 人团队", "互联互通四甲评审通过"],
        ["数据抽取成功率 99.2%", "推进专病数据库建设"],
    )):
        page = prs.slides.add_slide(blank)
        add_bar(page)
        add_logo(page)
        add_text(page, 0.9, 0.75, 8.0, 0.9, f"章节标题{index + 1}", 24)
        add_text(page, 0.9, 1.9, 8.0, 2.6, "\n".join(lines), 16)

    # slide 4 -- closing
    closing = prs.slides.add_slide(blank)
    add_bar(closing)
    add_text(closing, 0.9, 3.0, 8.0, 1.2, "谢谢", 32)

    prs.save(str(path))
    return path


def load_annotated_dna(path: Path) -> dict[str, Any]:
    """extract_deck_dna + semantic role annotation in one call."""
    return annotate_deck_dna(extract_deck_dna(str(path), include_raw_xml=False))
