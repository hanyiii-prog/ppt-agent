from pathlib import Path

from pptx import Presentation as PptxPresentation

from ppt_agent.ir import Presentation
from ppt_agent.renderer import render_presentation


def _ir_dict() -> dict:
    return {
        "version": "0.1",
        "metadata": {"title": "Renderer Demo"},
        "theme": {"slide_size_inches": {"width": 13.333, "height": 7.5}},
        "slides": [
            {
                "id": "slide-01",
                "purpose": "cover",
                "components": [{"type": "title", "text": "Cover Title"}],
            },
            {
                "id": "slide-02",
                "purpose": "content",
                "components": [
                    {"type": "text", "text": "Absolute box", "x": 1.0, "y": 1.0, "w": 5.0, "h": 1.0,
                     "style": {"fill": {"type": "solid", "rgb": "EEEEEE", "alpha": 0.5}}},
                    {"type": "paragraph", "text": "flows below"},
                    {"type": "shape", "x": 7.0, "y": 1.0, "w": 2.0, "h": 2.0,
                     "style": {"fill": {"type": "solid", "rgb": "336699"}, "line": {"rgb": "000000", "width_pt": 2}}},
                ],
            },
        ],
    }


def test_render_presentation_creates_expected_slides(tmp_path: Path):
    presentation = Presentation.from_dict(_ir_dict())
    output = render_presentation(presentation, tmp_path / "out.pptx")

    assert output.exists()
    prs = PptxPresentation(str(output))
    assert len(prs.slides) == 2
    assert round(prs.slide_width / 914400, 3) == 13.333
    first_text = " ".join(shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "Cover Title" in first_text
    second_text = " ".join(shape.text_frame.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert "Absolute box" in second_text
    assert "flows below" in second_text


def test_flow_layout_stacks_components_without_geometry():
    """`resolve_layout` stays the shared fallback for geometry-less components.

    The production render path now composes slides through the design layer, so
    this contract is asserted against the layout engine directly rather than
    through a rendered file.
    """
    from ppt_agent.ir import Component, Slide
    from ppt_agent.styling import resolve_layout

    slide = Slide(
        id="s1",
        purpose="content",
        components=[
            Component(type="paragraph", text="line one"),
            Component(type="paragraph", text="line two"),
        ],
    )
    boxes = resolve_layout(slide, 13.333, 7.5)
    assert len(boxes) == 2
    assert boxes[0].y < boxes[1].y
    assert not boxes[0].absolute and not boxes[1].absolute


def test_render_presentation_composes_a_designed_deck(tmp_path: Path):
    """A semantic slide must come out designed, not as a bare text box."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation.from_dict(_ir_dict())
    output = render_presentation(presentation, tmp_path / "designed.pptx")
    prs = PptxPresentation(str(output))

    cover = list(prs.slides[0].shapes)
    # spine + spine accent + base band + rule + title
    assert len(cover) >= 5
    assert any(shape.has_text_frame and "Cover Title" in shape.text_frame.text for shape in cover)
    # the deck carries real decoration: autoshapes, not just text boxes
    assert any(shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE for shape in cover)


def test_template_injected_slides_are_left_untouched(tmp_path: Path):
    """Slides with explicit geometry were designed elsewhere; keep them verbatim."""
    presentation = Presentation.from_dict(_ir_dict())
    output = render_presentation(presentation, tmp_path / "template.pptx")
    prs = PptxPresentation(str(output))

    texts = " ".join(
        shape.text_frame.text for shape in prs.slides[1].shapes if shape.has_text_frame
    )
    assert "Absolute box" in texts
    assert "flows below" in texts


def test_render_without_pptx_raises(monkeypatch, tmp_path: Path):
    from ppt_agent import renderer

    monkeypatch.setattr(renderer, "_PPTX_AVAILABLE", False)
    try:
        renderer.render_presentation(Presentation(version="0.1", title="x"), tmp_path / "x.pptx")
    except RuntimeError as exc:
        assert "python-pptx" in str(exc)
    else:  # pragma: no cover
        raise AssertionError("expected RuntimeError when python-pptx is unavailable")
