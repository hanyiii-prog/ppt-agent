"""Extractor hardening tests: real PPTX in, hardened fidelity DNA out."""
from pathlib import Path

from ppt_agent.fidelity import extract_fidelity_dna, rotated_bbox_emu

from fidelity_fixtures import EMU_IN, build_rich_pptx


def test_shape_identity_uses_per_kind_c_nv_pr(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    kinds = {shape["kind"] for shape in dna["slide"]["shapes"]}
    assert {"sp", "pic", "graphicFrame", "cxnSp", "grpSp"} <= kinds

    by_kind = {}
    for shape in dna["slide"]["shapes"]:
        by_kind.setdefault(shape["kind"], []).append(shape)
    assert by_kind["pic"][0]["shape_id"] is not None
    assert by_kind["graphicFrame"][0]["shape_id"] is not None
    assert by_kind["cxnSp"][0]["shape_id"] is not None
    assert by_kind["grpSp"][0]["children"], "group children must be captured"


def test_rotation_flip_and_rendered_bbox(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    rotated = next(s for s in dna["slide"]["shapes"] if s["geometry_emu"]["rotation"] == 90.0)
    geom = rotated["geometry_emu"]
    expected = rotated_bbox_emu(geom["x"], geom["y"], geom["cx"], geom["cy"], 90.0)
    assert geom["rendered_bbox"]["x"] == round(expected[0], 2)
    assert geom["rendered_bbox"]["cx"] == round(expected[2], 2)
    # a 2x1 rect rotated 90deg renders as 1x2
    assert geom["rendered_bbox"]["cx"] < geom["rendered_bbox"]["cy"]

    flipped = next(
        s for s in dna["slide"]["shapes"] if s["geometry_emu"]["flip_h"] and s["geometry_emu"]["flip_v"]
    )
    assert flipped["geometry_emu"]["flip_h"] is True
    assert flipped["geometry_emu"]["flip_v"] is True


def test_group_child_coordinate_mapping(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    group = next(s for s in dna["slide"]["shapes"] if s["kind"] == "grpSp")
    child = group["children"][0]
    mapped_child = child["geometry_emu"]
    assert mapped_child["cx"] > 0
    # parent group renders at 2x the child-space extent, so the mapped child
    # extent must be twice its declared extent (prefix-agnostic raw_xml scan)
    import re

    match = re.search(r'ext cx="(\d+)" cy="(\d+)"', child["raw_xml"])
    assert match, "child raw_xml must contain its declared extent"
    declared_cx = int(match.group(1))
    assert abs(mapped_child["cx"] - declared_cx * 2) <= 2.0


def test_structured_fill_alpha_gradient_and_theme(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    shapes = dna["slide"]["shapes"]

    solid = next(s for s in shapes if (s.get("fill") or {}).get("type") == "solid" and (s["fill"].get("color") or {}).get("rgb") == "1185FE")
    assert solid["fill"]["color"]["rgb"] == "1185FE"

    # cover band: transparency 35% -> alpha evidence 0.65 on slide 1
    cover = extract_fidelity_dna(path, slide_index=1)
    band = next(s for s in cover["slide"]["shapes"] if (s.get("fill") or {}).get("type") == "solid" and (s["fill"].get("color") or {}).get("rgb") == "0C67BC")
    assert band["fill"]["color"]["alpha"] == 0.65

    gradient = next(s for s in shapes if (s.get("fill") or {}).get("type") == "gradient")
    grad = gradient["fill"]
    assert len(grad["stops"]) >= 2
    # python-pptx gradient_angle=45 is stored as OOXML ang=315 (clockwise convention)
    assert grad["angle"] == 315.0

    themed = next(s for s in shapes if (s.get("fill") or {}).get("type") == "solid" and (s["fill"].get("color") or {}).get("scheme") == "accent1")
    assert themed["fill"]["color"]["resolved_rgb"], "accent1 must resolve via theme"
    assert dna["theme"]["colors"]["accent1"] == themed["fill"]["color"]["resolved_rgb"]


def test_typography_and_body_properties(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=1)
    title = next(
        s for s in dna["slide"]["shapes"]
        if s.get("text", "").startswith("天津市口腔医院")
    )
    run = title["typography"]["paragraphs"][0]["runs"][0]
    assert run["font_size_pt"] == 40.0
    assert run["bold"] is True
    assert run["font_latin"] == "Arial"

    dna3 = extract_fidelity_dna(path, slide_index=3)
    textbox = next(
        s for s in dna3["slide"]["shapes"] if s.get("text", "").startswith("正文")
    )
    body = textbox["typography"]["body"]
    assert body["wrap"] == "square"
    assert body.get("lIns") == 137160  # 0.15 inch (defaults are stripped by python-pptx)
    assert body["autofit"]["type"] == "noAutofit"


def test_media_crop_and_hash(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    pic = next(s for s in dna["slide"]["shapes"] if s["kind"] == "pic")
    media = pic["media"]
    assert media["sha256"], "media must be content-hashed"
    assert media["byte_size"] > 0
    assert media["crop"] == {"l": 25000}
    assert dna["assets"], "media manifest must list package assets"


def test_connector_and_custom_geometry(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    connector = next(s for s in dna["slide"]["shapes"] if s["kind"] == "cxnSp")
    assert connector["connector"]["preset"] in ("straightConnector1", "line")
    assert connector["connector"]["start_arrow"]["type"] == "triangle"
    assert connector["connector"]["end_arrow"]["type"] == "arrow"
    assert connector["connector"]["start"]["y"] < connector["connector"]["end"]["y"]

    custom = next(s for s in dna["slide"]["shapes"] if s.get("custom_geometry"))
    cust = custom["custom_geometry"]
    assert cust["paths"], "custGeom pathLst must be parsed"
    assert any(count.get("lnTo") for count in (p["command_counts"] for p in cust["paths"]))
    assert cust["gd"] is not None


def test_table_structure(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    frame = next(s for s in dna["slide"]["shapes"] if s["kind"] == "graphicFrame")
    table = frame["table"]
    assert len(table["column_widths_emu"]) == 3
    assert len(table["rows"]) == 3
    assert table["rows"][1]["cells"][1]["text"] == "99.2%"
    # cell(1,2).merge(cell(2,2)) is a vertical merge: rowSpan on the first cell
    assert any(
        cell.get("rowSpan") == 2 or cell.get("vMerge")
        for row in table["rows"]
        for cell in row["cells"]
    )
    assert table["structure_hash"]


def test_global_render_order_spans_master_layout_slide(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    orders = []
    for source in ("master", "layout", "slide"):
        for shape in dna[source]["shapes"]:
            orders.append((shape["global_render_order"], source))
    orders.sort()
    assert [source for _, source in orders] == sorted(
        [source for _, source in orders], key=lambda s: ["master", "layout", "slide"].index(s)
    )
    stacked = dna["slide"]["shapes"][0]["stack"]
    assert stacked["below"] == stacked["global_z"]
    assert stacked["above"] + stacked["below"] == stacked["total"] - 1 if "total" in stacked else True


def test_page_kind_classification(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    assert extract_fidelity_dna(path, slide_index=1)["page_kind"] == "cover"
    assert extract_fidelity_dna(path, slide_index=2)["page_kind"] == "toc"
    assert extract_fidelity_dna(path, slide_index=3)["page_kind"] == "content"
    assert extract_fidelity_dna(path, slide_index=4)["page_kind"] == "closing"


def test_toc_structure_extraction(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=2)
    assert dna["page_kind"] == "toc"
    toc = dna["toc"]
    assert toc["item_count"] == 3
    assert [item["numbering"] for item in toc["items"]] == [1, 2, 3]
    assert toc["numbering_ordered"] is True
    assert toc["structure_fingerprint"]


def test_background_and_inheritance_resolution(tmp_path: Path):
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    prs.slide_width = EMU_IN * 10
    prs.slide_height = EMU_IN * 7
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text_frame.paragraphs[0].add_run().text = "继承标题"
    for paragraph in slide.shapes.title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(40)
    path = tmp_path / "inherit.pptx"
    prs.save(path)

    dna = extract_fidelity_dna(path, slide_index=1)
    resolutions = dna["placeholder_resolutions"]
    assert resolutions, "placeholder inheritance chain must resolve"
    title_resolution = next(r for r in resolutions if r["placeholder"].get("type") in ("title", "ctrTitle"))
    assert title_resolution["resolved"]["font_size_pt"] == 40.0
    assert title_resolution["source"]["font_size_pt"] == "slide"
    assert title_resolution["layout_shape_id"] or title_resolution["master_shape_id"]
    # subtitle inherits from the master chain, never claims slide ownership
    subtitle = next((r for r in resolutions if r["placeholder"].get("type") in ("subTitle", "body")), None)
    if subtitle:
        assert subtitle["source"].get("font_size_pt") in ("layout", "master", None) or "font_size_pt" not in subtitle["source"]


def test_layout_injected_shape_appears_as_inherited_layer(tmp_path: Path):
    path = build_rich_pptx(tmp_path / "rich.pptx")
    dna = extract_fidelity_dna(path, slide_index=3)
    layout_shapes = dna["layout"]["shapes"]
    band = next(s for s in layout_shapes if s["name"] == "LayoutBand")
    assert band["fill"]["color"]["rgb"] == "1185FE"
    assert band["fill"]["color"]["alpha"] == 0.15
    assert band["preset_geometry"]["prst"] == "roundRect"
    # the band's rendered bbox is present and consistent with a non-rotated rect
    bbox = band["geometry_emu"]["rendered_bbox"]
    assert bbox["cx"] == band["geometry_emu"]["cx"]
    assert bbox["cy"] == band["geometry_emu"]["cy"]
