"""Page-kind DNA (``template-dna/v0.4``) regression tests.

These lock the three things v0.3 could not express and that silently wrecked
earlier rebuilds:

* the full rendered layer stack (master -> layout -> slide) with global paint
  order, so "what covers what" is answerable;
* rotation / flip with a rotation-aware bbox, so a 90-degree band is not
  mistaken for a tall thin box;
* transparency everywhere it lives -- per gradient stop, on run colours and on
  picture fills -- not just the first ``solidFill`` alpha.
"""

from __future__ import annotations

import struct
import zlib
from io import BytesIO
from pathlib import Path
import xml.etree.ElementTree as ET

import pytest

from ppt_agent.page_dna import (
    parse_color,
    parse_fill,
    rendered_bbox,
)
from ppt_agent.template import analyze_pptx

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# A 90-degree flipped band with a 15% -> 0% gradient: the exact decoration the
# template keeps in its "章节标题页" layout, and the exact shape the old
# extractor flattened into a vertical bar with no transparency.
BAND_XML = (
    '<p:sp %s><p:nvSpPr><p:cNvPr id="90" name="chapter-band"/>'
    '<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
    '<a:xfrm rot="5400000" flipH="1">'
    '<a:off x="3244850" y="-1751330"/><a:ext cx="3576320" cy="10064750"/></a:xfrm>'
    '<a:prstGeom prst="round2SameRect"><a:avLst>'
    '<a:gd name="adj1" fmla="val 5681"/><a:gd name="adj2" fmla="val 0"/>'
    '</a:avLst></a:prstGeom>'
    '<a:gradFill rotWithShape="1"><a:gsLst>'
    '<a:gs pos="0"><a:srgbClr val="1185FE"><a:alpha val="15000"/></a:srgbClr></a:gs>'
    '<a:gs pos="100000"><a:srgbClr val="1185FE"><a:alpha val="0"/></a:srgbClr></a:gs>'
    '</a:gsLst><a:lin ang="0" scaled="1"/></a:gradFill>'
    '</p:spPr><p:txBody><a:bodyPr/><a:p><a:r><a:rPr lang="zh-CN"/>'
    '<a:t>章节</a:t></a:r></a:p></p:txBody></p:sp>'
)

BAR_XML = (
    '<p:sp %s><p:nvSpPr><p:cNvPr id="91" name="content-bar"/>'
    '<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
    '<a:xfrm><a:off x="0" y="0"/><a:ext cx="12192000" cy="647700"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
    '<a:gradFill rotWithShape="1"><a:gsLst>'
    '<a:gs pos="0"><a:srgbClr val="1185FE"><a:alpha val="14902"/></a:srgbClr></a:gs>'
    '<a:gs pos="100000"><a:srgbClr val="1185FE"><a:alpha val="0"/></a:srgbClr></a:gs>'
    '</a:gsLst><a:lin ang="0" scaled="1"/></a:gradFill>'
    '</p:spPr><p:txBody><a:bodyPr/><a:p/></p:txBody></p:sp>'
)


def _png_bytes() -> BytesIO:
    """Valid 1x1 RGB PNG as a file-like object (zlib only, no deps)."""
    width = height = 1
    raw = b"\x00" + b"\x00" * (width * 3)
    signature = b"\x89PNG\r\n\x1a\n"

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return (struct.pack(">I", len(data)) + body
                + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF))

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return BytesIO(signature + chunk(b"IHDR", ihdr)
                   + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


@pytest.fixture()
def dna_deck(tmp_path: Path) -> Path:
    pptx = pytest.importorskip("pptx")
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn
    from pptx.util import Inches

    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]          # "Blank" -- no name signal at all
    title = prs.slide_layouts[0]          # carries the rotated divider band
    title.name = "分隔页"                  # no token in the name table: structure decides
    blank.shapes._spTree.append(parse_xml(BAR_XML % nsdecls("p", "a")))
    title.shapes._spTree.append(parse_xml(BAND_XML % nsdecls("p", "a")))

    # 1 cover
    prs.slides.add_slide(blank)
    # 2 toc -- headline text is the signal
    toc = prs.slides.add_slide(blank)
    toc.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(3.0), Inches(0.8)) \
        .text_frame.text = "目录"
    # 3 section -- the rotated band comes from the layout, and the slide carries
    #   enough shapes that the band is the ONLY usable signal
    section = prs.slides.add_slide(title)
    for i in range(10):
        section.shapes.add_shape(1, Inches(0.4 + i * 0.25), Inches(5.2),
                                 Inches(0.1), Inches(0.1))
    # 4/5 content -- header bar from the layout, plus one page-exclusive shape
    body = prs.slides.add_slide(blank)
    prs.slides.add_slide(blank)
    # 6 closing
    prs.slides.add_slide(blank)

    # run-level colour alpha (60%)
    box = body.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.0), Inches(0.5))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "半透明文字"
    run.font.color.rgb = RGBColor(0x11, 0x85, 0xFE)
    run.font.size = pptx.util.Pt(18)
    props = run._r.get_or_add_rPr()
    color = next(iter(props.find(qn("a:solidFill"))))
    color.append(color.makeelement(qn("a:alpha"), {"val": "60000"}))

    # picture transfer alpha (50%)
    picture = body.shapes.add_picture(_png_bytes(), Inches(10.0), Inches(0.1),
                                      Inches(0.8), Inches(0.7))
    picture._element.blipFill.blip.append(
        parse_xml('<a:alphaModFix %s amt="50000"/>' % nsdecls("a"))
    )

    # a shape that exists on this page only -- must never become an ornament
    only = body.shapes.add_shape(1, Inches(1.0), Inches(4.0), Inches(1.0), Inches(1.0))
    only.name = "only-here"
    only.fill.solid()
    only.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)

    out = tmp_path / "dna-deck.pptx"
    prs.save(out)
    return out


# --------------------------------------------------------------------------- #
# pure helpers
# --------------------------------------------------------------------------- #
def test_rendered_bbox_swaps_axes_for_a_quarter_turn():
    bbox = rendered_bbox(3.5479, -1.9153, 3.9111, 11.0069, 90.0)
    assert bbox == pytest.approx([0.0007, 1.6326, 11.0069, 3.9111], abs=1e-3)


def test_parse_color_keeps_alpha_on_the_native_scale():
    element = ET.fromstring(
        f'<a:srgbClr xmlns:a="{A_NS}" val="1185FE"><a:alpha val="14902"/></a:srgbClr>'
    )
    info = parse_color(element)
    assert info["rgb"] == "1185FE"
    assert info["alpha"] == 14902
    assert info["opacity"] == pytest.approx(0.149, abs=1e-4)
    assert info["transparency"] == pytest.approx(0.851, abs=1e-3)


def test_parse_fill_records_every_gradient_stop_alpha():
    sp_pr = ET.fromstring(
        f'<a:spPr xmlns:a="{A_NS}"><a:gradFill rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="1185FE"><a:alpha val="15000"/></a:srgbClr></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="1185FE"><a:alpha val="0"/></a:srgbClr></a:gs>'
        f'</a:gsLst><a:lin ang="5400000" scaled="1"/></a:gradFill></a:spPr>'
    )
    info = parse_fill(sp_pr)
    assert info["type"] == "gradient"
    assert info["rot_with_shape"] is True
    assert [stop["alpha"] for stop in info["stops"]] == [15000, 0]
    assert info["stops"][1]["transparency"] == 1.0
    assert info["linear"]["angle_deg"] == 90.0
    assert info["stop_count"] == 2


# --------------------------------------------------------------------------- #
# deck-level
# --------------------------------------------------------------------------- #
def test_page_kinds_are_classified_from_structure(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    kinds = {page["slide"]: page["kind"] for page in dna["slides"]}
    assert kinds == {1: "cover", 2: "toc", 3: "section", 4: "content",
                     5: "content", 6: "closing"}
    assert dna["presentation"]["page_kind_counts"]["content"] == 2


def test_layer_stack_is_master_then_layout_then_slide(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    page = dna["slides"][3]
    order = [(layer["origin"], layer["render_order"]) for layer in page["layers"]]
    assert [seq for _, seq in order] == sorted(seq for _, seq in order)
    origins = [origin for origin, _ in order]
    assert origins.index("layout") < origins.index("slide")


def test_rotated_band_keeps_rotation_prst_and_real_bbox(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    page = dna["slides"][2]
    band = next(layer for layer in page["layers"] if layer["name"] == "chapter-band")
    geometry = band["geometry"]
    assert geometry["rotation"] == 90.0
    assert geometry["flip_horizontal"] is True
    assert geometry["width"] == pytest.approx(3.9111, abs=1e-3)
    assert geometry["prst_geom"]["prst"] == "round2SameRect"
    assert geometry["prst_geom"]["adjust"]["adj1"] == "val 5681"
    # the painted footprint, not the unrotated frame
    assert geometry["rendered_bbox"] == pytest.approx(
        [0.0007, 1.6326, 11.0069, 3.9111], abs=1e-3)
    assert page["layer_summary"]["rotated"], "rotated band missing from the summary"


def test_translucent_fills_are_surfaced_per_page(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    page = dna["slides"][2]
    band = next(layer for layer in page["layers"] if layer["name"] == "chapter-band")
    fill = band["style"]["fill"]
    assert [stop["alpha"] for stop in fill["stops"]] == [15000, 0]
    assert any(entry["name"] == "chapter-band"
               for entry in page["layer_summary"]["translucent"])


def test_run_colour_alpha_is_preserved(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    page = dna["slides"][3]
    fonts = [
        font
        for layer in page["layers"]
        for font in ((layer.get("text") or {}).get("fonts") or [])
        if font.get("rgb") == "1185FE"
    ]
    assert fonts, "run colour not captured"
    assert fonts[0]["alpha"] == 60000
    assert fonts[0]["color_opacity"] == pytest.approx(0.6)


def test_picture_transfer_alpha_and_media_are_preserved(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    page = dna["slides"][3]
    picture = next(layer for layer in page["layers"] if layer["element"] == "pic")
    assert picture["picture"]["alpha"] == 50000
    assert picture["picture"]["transparency"] == pytest.approx(0.5)
    assert picture["picture"]["media"]["ext"] == "png"
    assert picture["style"]["fill"]["type"] == "picture"


def test_ornaments_are_shared_by_every_page_of_the_kind_and_keep_paint_order(
    dna_deck: Path,
):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    content = dna["page_kinds"]["content"]
    assert content["count"] == 2
    names = [ornament["name"] for ornament in content["ornaments"]]
    # the layout-provided bar is on both content pages -> it is chrome
    bar = next(o for o in content["ornaments"] if o["name"] == "content-bar")
    assert bar["origin"] == "layout"
    # a page-exclusive shape never becomes part of the kind's DNA
    assert "only-here" not in names
    orders = [ornament["render_order"] for ornament in content["ornaments"]]
    assert orders == sorted(orders)
    assert all(o["origin"] in ("layout", "master") for o in content["ornaments"])


def test_section_kind_owns_the_rotated_band(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    section = dna["page_kinds"]["section"]
    assert section["count"] == 1
    band = next(o for o in section["ornaments"] if o["name"] == "chapter-band")
    assert band["geometry"]["rotation"] == 90.0
    assert band["geometry"]["prst_geom"]["prst"] == "round2SameRect"
    # the divider kind must not inherit the body chrome
    assert all(o["name"] != "content-bar" for o in section["ornaments"])


def test_v03_contract_keys_still_present(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    for key in ("schema", "presentation", "theme", "global_style_statistics",
                "masters", "slides", "special_surfaces", "page_kinds"):
        assert key in dna
    assert dna["schema"] == "template-dna/v0.4"
    assert dna["special_surfaces"]["first"]["role"] == "first"
    assert dna["special_surfaces"]["last"]["role"] == "last"
    assert dna["special_surfaces"]["body_slide_count"] == 4
    assert "font_scheme" in dna["theme"]
    # v0.3 consumers read slide-local shapes and z_index
    first_shapes = dna["slides"][0]["shapes"]
    assert all(shape["origin"] == "slide" for shape in first_shapes)
    assert [shape["z_index"] for shape in first_shapes] == list(range(len(first_shapes)))


def test_kind_overrides_win_over_classification(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False, kind_overrides={4: "toc"})
    assert dna["slides"][3]["kind"] == "toc"
    assert dna["presentation"]["page_kind_counts"]["toc"] == 2


def test_include_raw_xml_false_drops_blobs_but_keeps_hashes(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    page = dna["slides"][0]
    assert page["raw_slide_xml"] is None
    assert all(layer["fidelity"]["raw_xml"] is None for layer in page["layers"])
    hashes = [layer["fidelity"]["xml_sha256"] for layer in page["layers"]
              if layer["fidelity"]]
    assert hashes and all(hashes)


def test_background_reports_declared_vs_inherited(dna_deck: Path):
    dna = analyze_pptx(dna_deck, include_raw_xml=False)
    assert dna["slides"][0]["background"]["declared"] is False
