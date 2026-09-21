"""Regression: the clone route must be reachable and slide text must be clean.

These two defects made every real template build look nothing like the
template: a supplied ``template_path`` never auto-extracted DNA, so the
pipeline silently fell through to the designed route with a hardcoded
theme; and markdown emphasis markers (``**标题``) leaked verbatim onto
slides.
"""
from __future__ import annotations

from pathlib import Path

import pytest

pytest.importorskip("pptx")

from ppt_agent.agent.pipeline import run_pipeline
from ppt_agent.blank_deck import _strip_md
from tests.design_dna_fixtures import build_template_pptx

MD = (
    "# 口腔医院专病库汇报\n\n"
    "## 建设背景\n\n"
    "- **政策驱动**：国家要求\n"
    "- 现状矛盾：数据孤岛\n\n"
    "## 建设计划\n\n"
    "- 架构设计三层\n"
    "- 四大功能\n"
)


class TestStripMd:
    def test_paired_and_unpaired_emphasis(self):
        assert _strip_md("**标题**") == "标题"
        assert _strip_md("**半开") == "半开"
        assert _strip_md("未闭**") == "未闭"
        assert _strip_md("`code`") == "code"
        assert _strip_md("普通文本") == "普通文本"


def test_clone_route_reachable_from_template_path(tmp_path: Path):
    """A template alone must trigger the clone route (auto DNA extraction)."""
    tpl = build_template_pptx(tmp_path / "template.pptx")
    report = run_pipeline(MD, out_dir=tmp_path / "out", template_path=tpl)
    assert report["route"] == "clone"
    assert report.get("fallback_reason") is None
    assert Path(report["artifacts"]["pptx"]).exists()


def test_clone_pages_have_no_markdown_emphasis(tmp_path: Path):
    """Every rendered run on a clone deck is free of raw ``**`` markers."""
    from pptx import Presentation
    tpl = build_template_pptx(tmp_path / "template.pptx")
    report = run_pipeline(MD, out_dir=tmp_path / "out", template_path=tpl)
    prs = Presentation(report["artifacts"]["pptx"])
    seen = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    seen += 1
                    assert "**" not in run.text and "__" not in run.text, run.text
    assert seen > 0
