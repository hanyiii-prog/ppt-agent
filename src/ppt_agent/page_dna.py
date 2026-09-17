"""Per-page-kind Template DNA extraction (schema ``template-dna/v0.4``).

Why this exists
---------------
``template-dna/v0.3`` recorded a flat shape list per slide with a coarse role
(``first`` / ``last`` / ``body``). That is not enough to *clone* a deck, and it
is not enough to even describe one:

* a cover, a TOC, a section divider, a body page and a closing page each have
  their own layer stack, their own ornaments and their own inheritance. Treating
  all of them as "body" is how you end up painting a full-bleed blue rectangle
  over a white template section page;
* the *rendered* stacking order is master shapes, then layout shapes, then slide
  shapes -- a v0.3 ``z_index`` only covers the last of the three, so two logos
  can overlap at identical coordinates and the dump still looks "clean";
* transparency lives on the **colour transform nodes** (``a:alpha``), per
  gradient stop, in the picture fill (``a:alphaModFix``) and on the line fill.
  Reading only the first ``solidFill`` alpha silently turns a 15%-alpha wash
  into an opaque band.

What this module produces
-------------------------
``extract_deck_dna`` returns a ``template-dna/v0.4`` dict that keeps every v0.3
key (so existing consumers keep working) and adds:

``page_kinds``
    ``{kind: {"count", "slides", "layout_names", "ornaments", ...}}`` for
    ``cover`` / ``toc`` / ``section`` / ``content`` / ``closing``. ``ornaments``
    is the set of shapes present on *every* page of that kind -- that set is the
    chrome, and it is the thing a clone must inherit rather than redraw.

``pages[i]["layers"]``
    The full rendered stack: every master shape, layout shape and slide shape
    with a global ``render_order``, its ``origin`` and its complete DNA.

``pages[i]["layers"][n]["dna"]``
    Preset geometry + ``avLst`` adjustments (or ``custGeom`` path census),
    rotation / flip with a rotation-aware ``rendered_bbox``, fill (solid /
    gradient with **per-stop alpha** / picture with ``alphaModFix`` + crop),
    line (width in pt, alpha, dash, arrow heads, join), effects, picture media
    reference, and run-level typography including colour alpha and the raw
    ``latin`` / ``ea`` / ``cs`` typeface references.

Nothing here guesses: every number comes from the OOXML.
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any, Iterable, Iterator
import hashlib
import math
import re
import zipfile
import xml.etree.ElementTree as ET

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
EMU_PER_INCH = 914400.0
EMU_PER_PT = 12700.0

SCHEMA = "template-dna/v0.4"

PAGE_KINDS = ("cover", "toc", "section", "content", "closing")

# Layout names are authoritative when they carry a signal; token order matters
# ("封底" must not fall through to the generic "内容" bucket).
_LAYOUT_TOKENS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("cover", ("封面", "标题幻灯片", "首页", "title slide", "cover")),
    ("closing", ("封底", "结束页", "谢谢", "尾页", "thank", "closing", "end slide")),
    ("toc", ("目录", "contents", "agenda")),
    ("section", ("章节", "分隔", "过渡", "节标题", "section", "divider")),
    ("content", ("内容", "正文", "content", "body")),
)

_TOC_TEXT = re.compile(r"(目\s*录|CONTENTS|Agenda)", re.IGNORECASE)
_SECTION_TEXT = re.compile(
    r"^\s*(第[一二三四五六七八九十百]+[章节部分]"
    r"|[一二三四五六七八九十]{1,3}\s*[、．.]\s*\S"
    r"|Part\s*\d+)"
)

_COLOR_TAGS = frozenset(
    {"srgbClr", "schemeClr", "prstClr", "sysClr", "scrgbClr", "hslClr"}
)
_LINEAR_SCALE_TAGS = frozenset(
    {
        "alphaMod", "alphaOff", "lumMod", "lumOff", "shade", "tint",
        "satMod", "satOff", "hueMod", "hueOff", "comp", "inv", "gray",
        "biLevel", "gamma", "redMod", "greenMod", "blueMod",
    }
)


# --------------------------------------------------------------------------- #
# small helpers
# --------------------------------------------------------------------------- #
def _q(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


def _local(element: Any) -> str:
    try:
        return element.tag.rsplit("}", 1)[-1]
    except AttributeError:
        return ""


def _int(value: Any, default: int | None = None) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _flag(value: Any, default: bool = False) -> bool:
    if value is None:
        return default
    return str(value) in ("1", "true", "True")


def _inches(emu: Any) -> float | None:
    value = _int(emu)
    return round(value / EMU_PER_INCH, 4) if value is not None else None


def _points(emu: Any) -> float | None:
    value = _int(emu)
    return round(value / EMU_PER_PT, 2) if value is not None else None


def _first_color(container: Any) -> Any | None:
    for child in container:
        if _local(child) in _COLOR_TAGS:
            return child
    return None


def _xml_of(element: Any) -> str:
    try:
        return ET.tostring(element, encoding="unicode")
    except (TypeError, ValueError, AttributeError):
        return ""


# --------------------------------------------------------------------------- #
# colour / fill / line / effects
# --------------------------------------------------------------------------- #
def parse_color(element: Any) -> dict[str, Any]:
    """Full colour record: base reference plus every transform child.

    ``alpha`` keeps PowerPoint's native 0..100000 scale; ``opacity`` and
    ``transparency`` are the 0..1 convenience twins.
    """
    out: dict[str, Any] = {"kind": _local(element)}
    tag = out["kind"]
    value = element.get("val")
    if tag == "srgbClr" and value:
        out["rgb"] = value.upper()
    elif tag == "schemeClr" and value:
        out["scheme"] = value
    elif tag == "prstClr" and value:
        out["prst"] = value
    elif tag == "sysClr":
        out["sys"] = value
        if element.get("lastClr"):
            out["rgb"] = element.get("lastClr").upper()
    elif tag == "scrgbClr":
        out["scrgb"] = [element.get("r"), element.get("g"), element.get("b")]
    elif tag == "hslClr":
        out["hsl"] = [element.get("hue"), element.get("sat"), element.get("lum")]

    for child in element:
        name = _local(child)
        raw = child.get("val")
        if name == "alpha":
            amount = _int(raw)
            if amount is not None:
                out["alpha"] = amount
                out["opacity"] = round(amount / 100000, 4)
                out["transparency"] = round(1 - amount / 100000, 4)
        elif name in _LINEAR_SCALE_TAGS:
            amount = _int(raw)
            if amount is not None:
                out[name] = amount
        elif name in ("hlinkClick", "hlinkMouseOver"):
            out[name] = child.get(_q(R, "id")) or child.get("id")
        elif raw is not None:
            out[name] = raw
    return out


def parse_fill(sp_pr: Any) -> dict[str, Any]:
    """Structured fill: solid / gradient (per-stop colour **and alpha**) /
    picture (``alphaModFix``, crop) / pattern / none.

    The v0.3 convenience keys (``rgb``, ``alpha``, ``opacity``,
    ``transparency``) are still emitted so older consumers keep working.
    """
    if sp_pr is None:
        return {"type": "unknown"}

    if sp_pr.find(_q(A, "noFill")) is not None:
        return {"type": "none"}
    if sp_pr.find(_q(A, "grpFill")) is not None:
        return {"type": "group"}

    solid = sp_pr.find(_q(A, "solidFill"))
    if solid is not None:
        info: dict[str, Any] = {"type": "solid"}
        color = _first_color(solid)
        if color is not None:
            info.update(parse_color(color))
        return info

    grad = sp_pr.find(_q(A, "gradFill"))
    if grad is not None:
        info = {
            "type": "gradient",
            # PowerPoint's default is rotWithShape="1"; keep it explicit.
            "rot_with_shape": _flag(grad.get("rotWithShape"), True),
            "stops": [],
        }
        stop_list = grad.find(_q(A, "gsLst"))
        if stop_list is not None:
            for stop in stop_list:
                if _local(stop) != "gs":
                    continue
                entry: dict[str, Any] = {"pos": _int(stop.get("pos"), 0)}
                color = _first_color(stop)
                if color is not None:
                    entry.update(parse_color(color))
                info["stops"].append(entry)
        linear = grad.find(_q(A, "lin"))
        if linear is not None:
            angle = _int(linear.get("ang"), 0) or 0
            info["linear"] = {
                "angle_deg": round(angle / 60000.0, 2),
                "scaled": _flag(linear.get("scaled")),
            }
        path = grad.find(_q(A, "path"))
        if path is not None:
            rect = path.find(_q(A, "fillToRect"))
            info["path"] = {
                "path": path.get("path"),
                "fill_to_rect": (
                    {k: _int(rect.get(k), 0) for k in ("l", "t", "r", "b")}
                    if rect is not None
                    else None
                ),
            }
        if grad.find(_q(A, "tileRect")) is not None:
            info["tile_rect"] = True
        # v0.3 compatibility: mirror the first stop.
        stops = info["stops"]
        if stops:
            for key in ("rgb", "scheme", "alpha", "opacity", "transparency"):
                if key in stops[0] and key not in info:
                    info[key] = stops[0][key]
            info["stop_count"] = len(stops)
        return info

    blip = (
        sp_pr if _local(sp_pr) == "blipFill" else sp_pr.find(_q(A, "blipFill"))
    )
    if blip is not None:
        info = {"type": "picture"}
        node = blip.find(_q(A, "blip"))
        if node is not None:
            embed = node.get(_q(R, "embed"))
            link = node.get(_q(R, "link"))
            if embed:
                info["r_embed"] = embed
            if link:
                info["r_link"] = link
            alpha_fix = node.find(_q(A, "alphaModFix"))
            if alpha_fix is not None:
                amount = _int(alpha_fix.get("amt"))
                if amount is not None:
                    info["alpha"] = amount
                    info["opacity"] = round(amount / 100000, 4)
                    info["transparency"] = round(1 - amount / 100000, 4)
            for name in ("alphaMod", "alphaOff", "lum", "grayscl", "biLevel",
                         "duotone", "blur"):
                child = node.find(_q(A, name))
                if child is not None:
                    info[name] = True
        src = blip.find(_q(A, "srcRect"))
        if src is not None:
            info["crop"] = {
                k: (_int(src.get(k), 0) or 0) / 1000.0 for k in ("l", "t", "r", "b")
            }
        if blip.find(_q(A, "stretch")) is not None:
            info["stretch"] = True
        if blip.find(_q(A, "tile")) is not None:
            info["tile"] = True
        return info

    pattern = sp_pr.find(_q(A, "pattFill"))
    if pattern is not None:
        info = {"type": "pattern", "prst": pattern.get("prst")}
        fg = pattern.find(_q(A, "fgClr"))
        bg = pattern.find(_q(A, "bgClr"))
        if fg is not None:
            color = _first_color(fg)
            if color is not None:
                info["foreground"] = parse_color(color)
        if bg is not None:
            color = _first_color(bg)
            if color is not None:
                info["background"] = parse_color(color)
        return info

    return {"type": "none"}


def parse_line(sp_pr: Any) -> dict[str, Any]:
    if sp_pr is None:
        return {"type": "none"}
    line = sp_pr.find(_q(A, "ln"))
    if line is None:
        return {"type": "inherited"}
    info: dict[str, Any] = {"type": "line", "width_pt": _points(line.get("w"))}
    for attr in ("cap", "cmpd", "algn"):
        if line.get(attr):
            info[attr] = line.get(attr)
    if line.find(_q(A, "noFill")) is not None:
        info["fill"] = {"type": "none"}
    else:
        fill = parse_fill(line)
        if fill.get("type") not in ("none", "unknown"):
            info["fill"] = fill
            if fill.get("rgb"):
                info["rgb"] = fill["rgb"]
            if fill.get("alpha") is not None:
                info["alpha"] = fill["alpha"]
                info["transparency"] = fill.get("transparency")
    dash = line.find(_q(A, "prstDash"))
    if dash is not None:
        info["dash"] = dash.get("val")
        info["dash_style"] = str(dash.get("val", "")).upper()
    if line.find(_q(A, "custDash")) is not None:
        info["dash"] = "custom"
    for end in ("headEnd", "tailEnd"):
        node = line.find(_q(A, end))
        if node is not None:
            info[end] = {
                "type": node.get("type"),
                "w": node.get("w"),
                "len": node.get("len"),
            }
    for join in ("round", "bevel", "miter"):
        if line.find(_q(A, join)) is not None:
            info["join"] = join
    return info


def parse_effects(sp_pr: Any) -> list[dict[str, Any]]:
    """Structured ``effectLst`` (shadow / glow / soft edge / reflection)."""
    if sp_pr is None:
        return []
    effects: list[dict[str, Any]] = []
    for container_name in ("effectLst", "effectDag"):
        container = sp_pr.find(_q(A, container_name))
        if container is None:
            continue
        for node in container:
            name = _local(node)
            if name in _COLOR_TAGS:
                continue
            entry: dict[str, Any] = {"type": name}
            if node.get("prst"):
                entry["prst"] = node.get("prst")
            for attr in ("blurRad", "dist", "dir", "rad", "stA", "endA", "sy",
                         "ky", "scaled", "algn", "rotWithShape", "fov", "k",
                         "sx"):
                raw = node.get(attr)
                if raw is None:
                    continue
                if attr in ("blurRad", "dist", "rad"):
                    entry[f"{attr}_pt"] = _points(raw)
                elif attr == "dir":
                    entry["dir_deg"] = round((_int(raw, 0) or 0) / 60000.0, 2)
                elif attr in ("scaled", "rotWithShape"):
                    entry[attr] = _flag(raw)
                else:
                    entry[attr] = _int(raw, raw)
            color = _first_color(node)
            if color is not None:
                entry["color"] = parse_color(color)
            effects.append(entry)
    return effects


# --------------------------------------------------------------------------- #
# geometry
# --------------------------------------------------------------------------- #
def rendered_bbox(
    left: float | None, top: float | None, width: float | None,
    height: float | None, rotation_deg: float,
) -> list[float] | None:
    """Axis-aligned box a rotated shape actually occupies, in inches.

    ``shape.left/top/width/height`` describe the *unrotated* frame; a 90-degree
    band therefore reports a tall thin box while painting a wide flat one.
    """
    if None in (left, top, width, height):
        return None
    cx = left + width / 2
    cy = top + height / 2
    rot = rotation_deg % 360
    if abs(rot - 90) < 1e-6 or abs(rot - 270) < 1e-6:
        rw, rh = height, width
    elif abs(rot) < 1e-6 or abs(rot - 180) < 1e-6:
        rw, rh = width, height
    else:
        rad = math.radians(rot)
        cos, sin = abs(math.cos(rad)), abs(math.sin(rad))
        rw, rh = width * cos + height * sin, width * sin + height * cos
    return [round(cx - rw / 2, 4), round(cy - rh / 2, 4), round(rw, 4), round(rh, 4)]


def parse_transform(element: Any) -> dict[str, Any]:
    """``a:xfrm`` -> offsets/extents (EMU and inches) + rotation/flip in degrees."""
    if element is None:
        return {}
    xfrm = element.find(_q(A, "xfrm"))
    if xfrm is None:
        return {}
    out: dict[str, Any] = {}
    off = xfrm.find(_q(A, "off"))
    if off is not None:
        out["off_emu"] = [_int(off.get("x"), 0), _int(off.get("y"), 0)]
        out["left_in"] = _inches(off.get("x"))
        out["top_in"] = _inches(off.get("y"))
    ext = xfrm.find(_q(A, "ext"))
    if ext is not None:
        out["ext_emu"] = [_int(ext.get("cx"), 0), _int(ext.get("cy"), 0)]
        out["width_in"] = _inches(ext.get("cx"))
        out["height_in"] = _inches(ext.get("cy"))
    ch_off = xfrm.find(_q(A, "chOff"))
    if ch_off is not None:
        out["child_off_emu"] = [_int(ch_off.get("x"), 0), _int(ch_off.get("y"), 0)]
    ch_ext = xfrm.find(_q(A, "chExt"))
    if ch_ext is not None:
        out["child_ext_emu"] = [_int(ch_ext.get("cx"), 0), _int(ch_ext.get("cy"), 0)]
    rot = _int(xfrm.get("rot"), 0) or 0
    out["rotation_deg"] = round(rot / 60000.0, 3)
    out["flip_horizontal"] = _flag(xfrm.get("flipH"))
    out["flip_vertical"] = _flag(xfrm.get("flipV"))
    return out


def parse_geometry(sp_pr: Any) -> dict[str, Any]:
    """Preset geometry + ``avLst`` adjustments, or a ``custGeom`` path census."""
    if sp_pr is None:
        return {}
    preset = sp_pr.find(_q(A, "prstGeom"))
    if preset is not None:
        out: dict[str, Any] = {"type": "preset", "prst": preset.get("prst")}
        adj: dict[str, str] = {}
        av_lst = preset.find(_q(A, "avLst"))
        if av_lst is not None:
            for node in av_lst:
                if _local(node) == "gd" and node.get("name"):
                    adj[node.get("name")] = node.get("fmla")
        if adj:
            out["adjust"] = adj
        return out
    custom = sp_pr.find(_q(A, "custGeom"))
    if custom is not None:
        commands: Counter[str] = Counter()
        paths = 0
        path_lst = custom.find(_q(A, "pathLst"))
        if path_lst is not None:
            for path in path_lst:
                if _local(path) != "path":
                    continue
                paths += 1
                for command in path:
                    commands[_local(command)] += 1
        return {
            "type": "custom",
            "path_count": paths,
            "commands": dict(commands),
        }
    return {}


# --------------------------------------------------------------------------- #
# text
# --------------------------------------------------------------------------- #
def _parse_run(run: Any) -> dict[str, Any]:
    props = run.find(_q(A, "rPr"))
    font: dict[str, Any] = {}
    out: dict[str, Any] = {"text": "", "font": font}
    if props is not None:
        size = _int(props.get("sz"))
        if size is not None:
            font["size_pt"] = round(size / 100.0, 2)
        for attr, key in (
            ("b", "bold"), ("i", "italic"), ("u", "underline"),
            ("strike", "strike"), ("cap", "cap"), ("dirty", "dirty"),
        ):
            value = props.get(attr)
            if value is not None:
                font[key] = value if attr == "u" else _flag(value)
        for attr, key in (("spc", "spacing_centipoints"), ("kern", "kern"),
                          ("baseline", "baseline_pct")):
            value = _int(props.get(attr))
            if value is not None:
                font[key] = value
        if props.get("lang"):
            font["lang"] = props.get("lang")
        typefaces: dict[str, str] = {}
        for node_name in ("latin", "ea", "cs", "sym"):
            node = props.find(_q(A, node_name))
            if node is not None and node.get("typeface"):
                typefaces[node_name] = node.get("typeface")
        if typefaces:
            font["typeface"] = typefaces
            if "latin" in typefaces:
                font["name"] = typefaces["latin"]
            elif "ea" in typefaces:
                font["name"] = typefaces["ea"]
        fill = props.find(_q(A, "solidFill"))
        if fill is not None:
            color = _first_color(fill)
            if color is not None:
                parsed = parse_color(color)
                if "rgb" in parsed:
                    font["rgb"] = parsed["rgb"]
                if "scheme" in parsed:
                    font["scheme"] = parsed["scheme"]
                if "alpha" in parsed:
                    font["alpha"] = parsed["alpha"]
                    font["color_opacity"] = parsed["opacity"]
                    font["color_transparency"] = parsed["transparency"]
        highlight = props.find(_q(A, "highlight"))
        if highlight is not None:
            color = _first_color(highlight)
            if color is not None:
                font["highlight"] = parse_color(color)
        link = props.find(_q(A, "hlinkClick"))
        if link is not None:
            font["hyperlink"] = link.get(_q(R, "id"))
    # run content: <a:t> text plus <a:br> line breaks, in document order
    chunks: list[str] = []
    for node in run:
        name = _local(node)
        if name == "t":
            chunks.append(node.text or "")
        elif name == "br":
            chunks.append("\n")
        elif name == "tab":
            chunks.append("\t")
    out["text"] = "".join(chunks)
    return out


def _parse_bullet(para_props: Any) -> dict[str, Any] | None:
    if para_props is None:
        return None
    if para_props.find(_q(A, "buNone")) is not None:
        return {"type": "none"}
    char = para_props.find(_q(A, "buChar"))
    if char is not None:
        return {"type": "char", "char": char.get("char")}
    auto = para_props.find(_q(A, "buAutoNum"))
    if auto is not None:
        return {
            "type": "auto_number",
            "scheme": auto.get("type"),
            "start_at": _int(auto.get("startAt"), 1),
        }
    if para_props.find(_q(A, "buBlip")) is not None:
        return {"type": "picture"}
    return None


def _parse_spacing(node: Any, prefix: str) -> dict[str, Any]:
    if node is None:
        return {}
    out: dict[str, Any] = {}
    pct = node.find(_q(A, "spcPct"))
    if pct is not None:
        value = _int(pct.get("val"))
        if value is not None:
            out[f"{prefix}_pct"] = round(value / 1000.0, 2)
    pts = node.find(_q(A, "spcPts"))
    if pts is not None:
        value = _int(pts.get("val"))
        if value is not None:
            out[f"{prefix}_pt"] = round(value / 100.0, 2)
    return out


def _parse_paragraph(paragraph: Any) -> dict[str, Any]:
    props = paragraph.find(_q(A, "pPr"))
    levels = _int(props.get("lvl"), 0) if props is not None else 0
    item: dict[str, Any] = {
        "level": levels or 0,
        "alignment": None,
        "runs": [],
        "fonts": [],
    }
    if props is not None:
        algn = props.get("algn")
        item["alignment"] = algn.upper() if algn else None
        for attr in ("marL", "marR", "indent"):
            value = _int(props.get(attr))
            if value is not None:
                item[f"{attr}_in"] = round(value / EMU_PER_INCH, 4)
        if props.get("rtl") is not None:
            item["rtl"] = _flag(props.get("rtl"))
        bullet = _parse_bullet(props)
        if bullet:
            item["bullet"] = bullet
        else:
            color_node = props.find(_q(A, "buClr"))
            if color_node is not None:
                color = _first_color(color_node)
                if color is not None:
                    item.setdefault("bullet", {})["color"] = parse_color(color)
            size_pct = props.find(_q(A, "buSzPct"))
            if size_pct is not None:
                item.setdefault("bullet", {})["size_pct"] = round(
                    (_int(size_pct.get("val"), 0) or 0) / 1000.0, 2)
            font_node = props.find(_q(A, "buFont"))
            if font_node is not None:
                item.setdefault("bullet", {})["font"] = font_node.get("typeface")
        item.update(_parse_spacing(props.find(_q(A, "lnSpc")), "line_spacing"))
        item.update(_parse_spacing(props.find(_q(A, "spcBef")), "space_before"))
        item.update(_parse_spacing(props.find(_q(A, "spcAft")), "space_after"))
    text_parts: list[str] = []
    for node in paragraph:
        name = _local(node)
        if name == "r":
            run = _parse_run(node)
            item["runs"].append(run)
            if run["font"]:
                item["fonts"].append(run["font"])
            text_parts.append(run["text"])
        elif name == "br":
            text_parts.append("\n")
        elif name == "fld":
            field_text = "".join(
                (child.text or "") for child in node if _local(child) == "t"
            )
            text_parts.append(field_text)
    item["text"] = "".join(text_parts)
    return item


def parse_text(shape: Any) -> dict[str, Any] | None:
    """Full text DNA: body properties, paragraph properties, run properties."""
    if not getattr(shape, "has_text_frame", False):
        return None
    element = getattr(shape, "_element", None)
    body = element.find(_q(P, "txBody")) if element is not None else None
    paragraphs_raw = (
        [node for node in body if _local(node) == "p"] if body is not None else []
    )
    if not paragraphs_raw:
        return None

    paragraphs = [_parse_paragraph(node) for node in paragraphs_raw]
    fonts: list[dict[str, Any]] = []
    for paragraph in paragraphs:
        fonts.extend(paragraph["fonts"])

    out: dict[str, Any] = {
        "text": "\n".join(p["text"] for p in paragraphs),
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "fonts": fonts,
        "runs": [run for p in paragraphs for run in p["runs"]],
    }

    frame = getattr(shape, "text_frame", None)
    if frame is not None:
        for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            try:
                out[attr] = round(getattr(frame, attr) / EMU_PER_INCH, 4)
            except Exception:
                pass
        try:
            out["word_wrap"] = bool(frame.word_wrap)
        except Exception:
            pass
        try:
            anchor = frame.vertical_anchor
            out["vertical_anchor"] = str(anchor).split(".")[-1] if anchor else None
        except Exception:
            pass
        try:
            out["auto_size"] = str(frame.auto_size).split(".")[-1]
        except Exception:
            pass

    if body is not None:
        body_props = body.find(_q(A, "bodyPr"))
        if body_props is not None:
            dna: dict[str, Any] = {}
            insets: dict[str, float | None] = {}
            for attr, key in (("lIns", "left"), ("tIns", "top"),
                              ("rIns", "right"), ("bIns", "bottom")):
                insets[key] = _inches(body_props.get(attr))
            dna["insets_in"] = insets
            for attr in ("anchor", "anchorCtr", "wrap", "rtlCol", "vert",
                         "upright", "spcFirstLastPara", "numCol"):
                value = body_props.get(attr)
                if value is not None:
                    dna[attr] = value
            rotation = _int(body_props.get("rot"))
            if rotation:
                dna["rot_deg"] = round(rotation / 60000.0, 2)
            for fit_name in ("noAutofit", "normAutofit", "spAutoFit"):
                fit = body_props.find(_q(A, fit_name))
                if fit is not None:
                    fit_info: dict[str, Any] = {"type": fit_name}
                    for attr in ("fontScale", "lnSpcReduction"):
                        value = _int(fit.get(attr))
                        if value is not None:
                            fit_info[attr] = value
                    dna["autofit"] = fit_info
                    break
            out["body"] = dna
    return out


# --------------------------------------------------------------------------- #
# shape-level DNA
# --------------------------------------------------------------------------- #
def _sp_pr(element: Any) -> Any:
    if element is None:
        return None
    for tag in ("spPr", "grpSpPr"):
        node = element.find(_q(P, tag))
        if node is not None:
            return node
    return None


def _shape_type(shape: Any) -> str:
    try:
        return str(shape.shape_type).split(".")[-1]
    except Exception:
        return "unknown"


def _shape_kind(element: Any) -> str:
    return _local(element) if element is not None else "unknown"


def _picture_dna(shape: Any, element: Any) -> dict[str, Any]:
    """Picture specifics: media reference, alpha, crop, decoded pixel size."""
    out: dict[str, Any] = {}
    blip_fill = element.find(_q(P, "blipFill")) if element is not None else None
    if blip_fill is not None:
        blip = blip_fill.find(_q(A, "blip"))
        if blip is not None:
            embed = blip.get(_q(R, "embed"))
            link = blip.get(_q(R, "link"))
            if embed:
                out["r_embed"] = embed
            if link:
                out["r_link"] = link
            alpha_fix = blip.find(_q(A, "alphaModFix"))
            if alpha_fix is not None:
                amount = _int(alpha_fix.get("amt"))
                if amount is not None:
                    out["alpha"] = amount
                    out["opacity"] = round(amount / 100000, 4)
                    out["transparency"] = round(1 - amount / 100000, 4)
        src = blip_fill.find(_q(A, "srcRect"))
        if src is not None:
            out["crop_pct"] = {
                k: (_int(src.get(k), 0) or 0) / 1000.0 for k in ("l", "t", "r", "b")
            }
    try:
        image = shape.image
        out["media"] = {
            "ext": image.ext,
            "px": list(image.size),
            "bytes": len(image.blob),
            "sha256": hashlib.sha256(image.blob).hexdigest()[:16],
        }
        try:
            out["media"]["dpi"] = list(image.dpi)
        except Exception:
            pass
    except Exception:
        pass
    return out


def _table_dna(element: Any) -> dict[str, Any] | None:
    graphic = element.find(_q(A, "graphic")) if element is not None else None
    if graphic is None:
        return None
    table = graphic.find(f"{_q(A, 'graphicData')}/{_q(A, 'tbl')}")
    if table is None:
        return None
    rows = [node for node in table if _local(node) == "tr"]
    out: dict[str, Any] = {
        "type": "table",
        "rows": len(rows),
        "cols": 0,
        "col_widths_in": [],
        "row_heights_in": [],
        "cells": [],
    }
    grid = table.find(_q(A, "tblGrid"))
    if grid is not None:
        widths = [_inches(node.get("w")) for node in grid if _local(node) == "gridCol"]
        out["col_widths_in"] = widths
        out["cols"] = len(widths)
    for row in rows:
        height = _int(row.get("h"))
        if height is not None:
            out["row_heights_in"].append(round(height / EMU_PER_INCH, 4))
        cells: list[str] = []
        for cell in row:
            if _local(cell) != "tc":
                continue
            text = "".join(node.text or "" for node in cell.iter(_q(A, "t")))
            cells.append(text)
        out["cells"].append(cells)
    return out


def extract_shape_dna(
    shape: Any,
    z_index: int,
    *,
    origin: str = "slide",
    render_order: int | None = None,
    parent_id: str | None = None,
    include_raw_xml: bool = True,
) -> dict[str, Any]:
    """One shape -> its complete DNA record (superset of the v0.3 record)."""
    element = getattr(shape, "_element", None)
    sp_pr = _sp_pr(element)
    xml = _xml_of(element)

    geometry: dict[str, Any] = {}
    transform = parse_transform(sp_pr)
    for attr in ("left", "top", "width", "height"):
        try:
            value = getattr(shape, attr)
            geometry[attr] = round(value / EMU_PER_INCH, 4) if value is not None else None
        except Exception:
            geometry[attr] = None
    if transform:
        if transform.get("left_in") is not None:
            geometry["left"] = transform["left_in"]
        if transform.get("top_in") is not None:
            geometry["top"] = transform["top_in"]
        if transform.get("width_in") is not None:
            geometry["width"] = transform["width_in"]
        if transform.get("height_in") is not None:
            geometry["height"] = transform["height_in"]
        geometry["rotation"] = transform.get("rotation_deg", 0.0)
        geometry["flip_horizontal"] = transform.get("flip_horizontal")
        geometry["flip_vertical"] = transform.get("flip_vertical")
        geometry["xfrm_emu"] = {
            "off": transform.get("off_emu"),
            "ext": transform.get("ext_emu"),
        }
        if transform.get("child_off_emu"):
            geometry["child_off_emu"] = transform["child_off_emu"]
            geometry["child_ext_emu"] = transform.get("child_ext_emu")
    else:
        try:
            geometry.setdefault("rotation", getattr(shape, "rotation", 0))
            geometry.setdefault("flip_horizontal", getattr(shape, "flip_horizontal", False))
            geometry.setdefault("flip_vertical", getattr(shape, "flip_vertical", False))
        except Exception:
            pass
    geometry["rendered_bbox"] = rendered_bbox(
        geometry.get("left"), geometry.get("top"),
        geometry.get("width"), geometry.get("height"),
        float(geometry.get("rotation") or 0.0),
    )

    # a picture's paint lives in p:blipFill, not in p:spPr
    fill_source = sp_pr
    if element is not None and _local(element) == "pic":
        blip_fill = element.find(_q(P, "blipFill"))
        if blip_fill is not None:
            fill_source = blip_fill

    style: dict[str, Any] = {
        "fill": parse_fill(fill_source),
        "line": parse_line(sp_pr),
    }
    prst = parse_geometry(sp_pr)
    if prst:
        geometry["prst_geom"] = prst
        style["geometry_kind"] = prst.get("type")

    record: dict[str, Any] = {
        "id": str(getattr(shape, "shape_id", "")),
        "name": getattr(shape, "name", None),
        "type": _shape_type(shape),
        "element": _shape_kind(element),
        "z_index": z_index,
        "z_order": z_index,
        "origin": origin,
        "render_order": render_order if render_order is not None else z_index,
        "parent_id": parent_id,
        "geometry": geometry,
        "style": style,
    }

    effects = parse_effects(sp_pr)
    if effects:
        style["effects"] = effects

    text = parse_text(shape)
    if text is not None:
        record["text"] = text

    placeholder = None
    if getattr(shape, "is_placeholder", False):
        try:
            fmt = shape.placeholder_format
            placeholder = {
                "idx": fmt.idx,
                "type": str(fmt.type).split(".")[-1],
                "name": getattr(shape, "name", None),
            }
        except Exception:
            placeholder = {"type": "unknown"}
        record["placeholder"] = placeholder

    if record["type"] == "PICTURE" or record["element"] == "pic":
        picture = _picture_dna(shape, element)
        if picture:
            record["picture"] = picture

    table = _table_dna(element)
    if table:
        record["table"] = table

    fidelity: dict[str, Any] = {}
    if xml:
        fidelity["xml_sha256"] = hashlib.sha256(xml.encode("utf-8")).hexdigest()
        fidelity["raw_xml"] = xml if include_raw_xml else None
        fidelity["custom_geometry"] = "custGeom" in xml
        fidelity["gradient_fill"] = "gradFill" in xml
        fidelity["alpha_transforms"] = bool(re.search(r"<a:alpha(?:ModFix|Off)?", xml))
        fidelity["blip_embeds"] = re.findall(r'r:embed="([^"]+)"', xml)
        fidelity["blip_links"] = re.findall(r'r:link="([^"]+)"', xml)
        fidelity["has_effects"] = any(
            token in xml for token in ("effectLst", "effectDag", "outerShdw", "glow")
        )
        fidelity["has_transform_2d"] = "xfrm" in xml
        fidelity["rotated"] = bool(
            re.search(r'<a:xfrm[^>]*(rot|flipH|flipV)=', xml)
        )
    record["fidelity"] = fidelity

    children = getattr(shape, "shapes", None)
    is_group = record["element"] == "grpSp"
    record["is_group"] = bool(is_group)
    if is_group and children is not None:
        record["children"] = [
            extract_shape_dna(
                child, index, origin=origin, render_order=index,
                parent_id=record["id"], include_raw_xml=include_raw_xml,
            )
            for index, child in enumerate(children)
        ]
    return record


# --------------------------------------------------------------------------- #
# layer stack
# --------------------------------------------------------------------------- #
def _safe_shapes(owner: Any) -> list[Any]:
    try:
        return list(owner.shapes)
    except Exception:
        return []


def extract_layer_stack(
    slide: Any, *, include_raw_xml: bool = True, include_master: bool = True
) -> list[dict[str, Any]]:
    """The *rendered* stack: master shapes, then layout shapes, then slide shapes.

    Within each layer shapes paint in XML document order, so a single global
    ``render_order`` is enough to answer "what covers what".
    """
    layers: list[dict[str, Any]] = []
    order = 0

    def _add(owner: Any, origin: str) -> None:
        nonlocal order
        for index, shape in enumerate(_safe_shapes(owner)):
            record = extract_shape_dna(
                shape, index, origin=origin, render_order=order,
                include_raw_xml=include_raw_xml,
            )
            layers.append(record)
            order += 1

    layout = None
    try:
        layout = slide.slide_layout
    except Exception:
        layout = None

    if include_master and layout is not None:
        master = None
        try:
            master = layout.slide_master
        except Exception:
            master = None
        if master is not None:
            _add(master, "master")

    if layout is not None:
        _add(layout, "layout")

    _add(slide, "slide")
    return layers


def layer_stack_summary(layers: list[dict[str, Any]]) -> dict[str, Any]:
    counts = Counter(str(layer.get("origin")) for layer in layers)
    return {
        "total": len(layers),
        "by_origin": dict(counts),
        "rotated": [
            {
                "name": layer.get("name"),
                "origin": layer.get("origin"),
                "rotation_deg": (layer.get("geometry") or {}).get("rotation"),
                "rendered_bbox": (layer.get("geometry") or {}).get("rendered_bbox"),
                "prst": ((layer.get("geometry") or {}).get("prst_geom") or {}).get("prst"),
            }
            for layer in layers
            if abs(float((layer.get("geometry") or {}).get("rotation") or 0.0)) > 1e-6
            or (layer.get("geometry") or {}).get("flip_horizontal")
            or (layer.get("geometry") or {}).get("flip_vertical")
        ],
        "translucent": [
            {
                "name": layer.get("name"),
                "origin": layer.get("origin"),
                "fill_type": ((layer.get("style") or {}).get("fill") or {}).get("type"),
                "alpha": ((layer.get("style") or {}).get("fill") or {}).get("alpha"),
                "stops": [
                    {"pos": stop.get("pos"), "rgb": stop.get("rgb"),
                     "alpha": stop.get("alpha")}
                    for stop in (((layer.get("style") or {}).get("fill") or {}).get("stops") or [])
                ],
            }
            for layer in layers
            if ((layer.get("style") or {}).get("fill") or {}).get("alpha") is not None
            or (((layer.get("style") or {}).get("fill") or {}).get("stops"))
        ],
    }


# --------------------------------------------------------------------------- #
# page-kind classification
# --------------------------------------------------------------------------- #
def _headline_text(slide: Any) -> str:
    """Best-effort headline: title placeholder first, else the largest run.

    Sizes are read from ``a:rPr/@sz`` in the XML because most real decks inherit
    their font size from the layout, leaving ``run.font.size`` as ``None``.
    """
    best, best_size = "", -1.0
    for shape in _safe_shapes(slide):
        try:
            if shape.is_placeholder and shape.placeholder_format.idx == 0:
                text = (shape.text_frame.text or "").strip()
                if text:
                    return text
        except Exception:
            pass
        element = getattr(shape, "_element", None)
        if element is None:
            continue
        body = element.find(_q(P, "txBody"))
        if body is None:
            continue
        text = "".join(node.text or "" for node in body.iter(_q(A, "t"))).strip()
        if not text:
            continue
        sizes = [int(s) for s in re.findall(r'\bsz="(\d+)"', _xml_of(body))]
        size = max(sizes) / 100.0 if sizes else 0.0
        if size > best_size:
            best_size, best = size, text
    return best


# A divider/section page is structurally bare: a headline and not much else.
_SECTION_MAX_SHAPES = 8


def page_signals(slide: Any) -> dict[str, Any]:
    """Cheap structural fingerprint used to classify ``Blank``-layout pages.

    Deck-generated files (and most LLM-built decks) put everything on the slide
    and name every layout ``Blank``, so layout names carry no signal and
    headline regexes are fooled by numbered content headings. Structure does
    not lie: a divider draws a rotated band and almost nothing else, a body
    page draws a full-width header bar and then lots of content.
    Inherited decorations (layout and master shapes) count: a cloned-template
    body page may be distinguishable only by the chrome its layout paints.
    """
    slide_shapes = _safe_shapes(slide)
    inherited: list[Any] = []
    try:
        layout = slide.slide_layout
    except Exception:
        layout = None
    if layout is not None:
        inherited = list(_safe_shapes(layout))
        try:
            inherited += _safe_shapes(layout.slide_master)
        except Exception:
            pass

    has_header_bar = False
    has_rotated_band = False
    for shape in list(slide_shapes) + inherited:
        element = getattr(shape, "_element", None)
        sp_pr = _sp_pr(element)
        if sp_pr is None:
            continue
        transform = parse_transform(sp_pr)
        prst = (parse_geometry(sp_pr) or {}).get("prst")
        width = float(transform.get("width_in") or 0.0)
        height = float(transform.get("height_in") or 0.0)
        top = float(transform.get("top_in") or 0.0)
        rotation = abs(float(transform.get("rotation_deg") or 0.0))
        if rotation > 1.0 and prst in ("round2SameRect", "roundRect", "rect"):
            if width >= 2.0 and height >= 2.0:
                has_rotated_band = True
        if not has_header_bar and width >= 10.0 and abs(top) < 0.2 and 0.4 <= height <= 1.2:
            has_header_bar = True
    return {
        "shape_count": len(slide_shapes),
        "inherited_shape_count": len(inherited),
        "has_header_bar": has_header_bar,
        "has_rotated_band": has_rotated_band,
    }


def classify_page_kind(
    slide: Any, index: int, total: int, *, layout_name: str | None = None
) -> str:
    """``cover`` / ``toc`` / ``section`` / ``content`` / ``closing``.

    Resolution order, most trustworthy first:

    1. position -- the bookends. A last page reusing the *cover* layout is a
       closing page, not a second cover;
    2. headline text -- ``目录`` / ``CONTENTS`` means TOC. This runs *before*
       the layout name because a TOC routinely reuses the body layout;
    3. layout name -- authored intent (``内容页 - 有标题``, ``章节标题页``);
    4. structure -- a rotated band with almost nothing else is a divider, a
       full-width top bar plus content is a body page.
    """
    if total <= 1:
        return "cover"
    if index == 1:
        return "cover"
    if index == total:
        return "closing"

    headline = _headline_text(slide)
    if headline and _TOC_TEXT.search(headline):
        return "toc"

    name = (layout_name or "").strip().lower()
    if name and name != "blank":
        for kind, tokens in _LAYOUT_TOKENS:
            if any(token in name for token in tokens):
                return kind

    signals = page_signals(slide)
    if signals["has_rotated_band"]:
        return "section"
    if not signals["has_header_bar"] and signals["shape_count"] <= _SECTION_MAX_SHAPES:
        return "section"
    if headline and _SECTION_TEXT.match(headline) and not signals["has_header_bar"]:
        return "section"
    return "content"


# --------------------------------------------------------------------------- #
# page-kind aggregation: the ornaments a clone must inherit, not redraw
# --------------------------------------------------------------------------- #
def _r2(value: Any) -> str:
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return "-"


def _ornament_key(record: dict[str, Any]) -> str:
    """Identity of a decoration: where it paints, what it is, how it is styled."""
    geometry = record.get("geometry") or {}
    fill = (record.get("style") or {}).get("fill") or {}
    prst = (geometry.get("prst_geom") or {}).get("prst")
    return "|".join(
        str(part)
        for part in (
            record.get("origin"),
            record.get("element"),
            record.get("type"),
            prst,
            _r2(geometry.get("left")),
            _r2(geometry.get("top")),
            _r2(geometry.get("width")),
            _r2(geometry.get("height")),
            geometry.get("rotation"),
            geometry.get("flip_horizontal"),
            geometry.get("flip_vertical"),
            fill.get("type"),
            fill.get("rgb") or fill.get("scheme") or "",
            fill.get("alpha"),
            len(fill.get("stops") or []),
        )
    )


def aggregate_page_kinds(pages: list[dict[str, Any]]) -> dict[str, Any]:
    """Group pages by kind and surface what every page of that kind shares.

    ``ornaments`` = shapes present on *all* pages of the kind. Those are the
    template's per-kind chrome; a clone inherits them and must not paint over.
    """
    grouped: dict[str, list[dict[str, Any]]] = {kind: [] for kind in PAGE_KINDS}
    for page in pages:
        grouped.setdefault(str(page.get("kind")), []).append(page)

    result: dict[str, Any] = {}
    for kind, members in grouped.items():
        if not members:
            result[kind] = {
                "count": 0,
                "slides": [],
                "layout_names": [],
                "ornaments": [],
                "frequent": [],
                "layer_origins": {},
                "transparency": [],
                "rotation": [],
            }
            continue
        seen: dict[str, dict[str, Any]] = {}
        counts: Counter[str] = Counter()
        for page in members:
            for layer in page.get("layers") or []:
                key = _ornament_key(layer)
                counts[key] += 1
                seen.setdefault(key, layer)
        ornaments = []
        frequent = []
        total_members = len(members)
        for key, count in counts.items():
            record = seen[key]
            entry = {
                "name": record.get("name"),
                "origin": record.get("origin"),
                # global paint order across master -> layout -> slide
                "render_order": record.get("render_order"),
                "z_index": record.get("z_index"),
                "type": record.get("type"),
                "element": record.get("element"),
                "present_on": count,
                "placeholder": record.get("placeholder"),
                "geometry": record.get("geometry"),
                "style": {
                    "fill": (record.get("style") or {}).get("fill"),
                    "line": (record.get("style") or {}).get("line"),
                },
                "text": (record.get("text") or {}).get("text"),
                "picture": record.get("picture"),
            }
            if count >= total_members:
                ornaments.append(entry)
            elif count >= max(2, total_members // 2):
                frequent.append({**entry, "total": total_members})
        # layer order is the point: keep the global paint order, never sort by name
        ornaments.sort(key=lambda item: item.get("render_order") or 0)
        frequent.sort(key=lambda item: (-item["present_on"], item.get("render_order") or 0))
        origins: Counter[str] = Counter()
        for page in members:
            for layer in page.get("layers") or []:
                origins[str(layer.get("origin"))] += 1
        layout_names = sorted(
            {str(page.get("layout_name")) for page in members if page.get("layout_name")}
        )
        result[kind] = {
            "count": len(members),
            "slides": [page.get("slide") for page in members],
            "layout_names": layout_names,
            "ornaments": ornaments,
            "frequent": frequent[:20],
            "layer_origins": dict(origins),
            "transparency": [
                {
                    "slide": page.get("slide"),
                    **entry,
                }
                for page in members
                for entry in ((page.get("layer_summary") or {}).get("translucent") or [])
            ],
            "rotation": [
                {"slide": page.get("slide"), **entry}
                for page in members
                for entry in ((page.get("layer_summary") or {}).get("rotated") or [])
            ],
        }
    return result


# --------------------------------------------------------------------------- #
# deck-level entry point
# --------------------------------------------------------------------------- #
def _theme_font_scheme(path: Path) -> dict[str, Any]:
    """``a:fontScheme`` -> major/minor typefaces for latin / ea / cs."""
    out: dict[str, Any] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            names = [n for n in archive.namelist() if re.search(r"theme\d+\.xml$", n)]
            if not names:
                return out
            root = ET.fromstring(archive.read(names[0]))
    except (OSError, ET.ParseError, zipfile.BadZipFile):
        return out
    scheme = root.find(f".//{_q(A, 'fontScheme')}")
    if scheme is None:
        return out
    out["name"] = scheme.get("name")
    for node_name in ("majorFont", "minorFont"):
        node = scheme.find(_q(A, node_name))
        if node is None:
            continue
        fonts: dict[str, Any] = {}
        latin = node.find(_q(A, "latin"))
        if latin is not None:
            fonts["latin"] = latin.get("typeface")
        ea = node.find(_q(A, "ea"))
        if ea is not None:
            fonts["ea"] = ea.get("typeface")
        cs = node.find(_q(A, "cs"))
        if cs is not None:
            fonts["cs"] = cs.get("typeface")
        script_fonts = [
            {"script": f.get("script"), "typeface": f.get("typeface")}
            for f in node.findall(_q(A, "font"))
        ]
        if script_fonts:
            fonts["script"] = script_fonts
        out[node_name] = fonts
    return out


def _theme_colors(path: Path) -> dict[str, str]:
    out: dict[str, str] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            names = [n for n in archive.namelist() if re.search(r"theme\d+\.xml$", n)]
            if not names:
                return out
            root = ET.fromstring(archive.read(names[0]))
    except (OSError, ET.ParseError, zipfile.BadZipFile):
        return out
    scheme = root.find(f".//{_q(A, 'clrScheme')}")
    if scheme is None:
        return out
    for child in scheme:
        value = next(iter(child), None)
        if value is not None and value.get("val"):
            out[child.tag.rsplit("}", 1)[-1]] = value.get("val")
    return out


def _background_dna(owner: Any) -> dict[str, Any]:
    """Background incl. ``a:bgRef`` inheritance, not just the resolved colour."""
    element = getattr(owner, "_element", None)
    out: dict[str, Any] = {}
    bg = element.find(_q(P, "bg")) if element is not None else None
    if bg is None:
        out["declared"] = False
    else:
        out["declared"] = True
        pr = bg.find(_q(P, "bgPr")) or bg.find(_q(P, "bgRef"))
        if pr is not None and _local(pr) == "bgPr":
            out["type"] = str(
                (parse_fill(pr) or {}).get("type")
            )
            out.update(parse_fill(pr))
        ref = bg.find(_q(P, "bgRef"))
        if ref is not None:
            out["bg_ref_index"] = _int(ref.get("idx"))
            color = _first_color(ref)
            if color is not None:
                out["bg_ref_color"] = parse_color(color)
    try:
        fill = owner.background.fill
        out["resolved_type"] = str(fill.type).split(".")[-1].lower()
        try:
            rgb = fill.fore_color.rgb
            if rgb:
                out["resolved_rgb"] = str(rgb)
        except Exception:
            pass
    except Exception:
        pass
    return out


def _layout_dna(layout: Any, *, include_raw_xml: bool) -> dict[str, Any]:
    shapes = _safe_shapes(layout)
    return {
        "name": layout.name,
        "type": str(getattr(layout, "type", "")).split(".")[-1],
        "background": _background_dna(layout),
        "shapes": [
            extract_shape_dna(s, i, origin="layout", render_order=i,
                              include_raw_xml=include_raw_xml)
            for i, s in enumerate(shapes)
        ],
        "shape_count": len(shapes),
        "raw_layout_xml": _xml_of(getattr(layout, "_element", None))
        if include_raw_xml
        else None,
    }


def extract_deck_dna(
    path: str | Path,
    *,
    include_raw_xml: bool = True,
    kind_overrides: dict[int, str] | None = None,
) -> dict[str, Any]:
    """Extract ``template-dna/v0.4`` -- per-page-kind, full-attribute DNA.

    Every v0.3 key is still produced (``slides``, ``special_surfaces``,
    ``masters``, ``theme``, ``global_style_statistics``), so existing consumers
    keep working; the added value is ``page_kinds`` and, per page, ``layers``.
    """
    try:
        from pptx import Presentation as PptxPresentation
    except ImportError as exc:  # pragma: no cover - optional extra
        raise RuntimeError(
            "python-pptx is required for PPTX analysis; "
            "install with pip install 'ppt-agent[pptx]'"
        ) from exc

    source = Path(path)
    prs = PptxPresentation(str(source))
    slide_count = len(prs.slides._sldIdLst) if hasattr(prs.slides, "_sldIdLst") else 0
    slides_seq = list(prs.slides)
    total = len(slides_seq)

    fonts: Counter[str] = Counter()
    font_sizes: Counter[float] = Counter()
    fills: Counter[str] = Counter()
    shape_types: Counter[str] = Counter()
    kind_counter: Counter[str] = Counter()

    pages: list[dict[str, Any]] = []
    for slide_no, slide in enumerate(slides_seq, 1):
        layout_name = None
        try:
            layout_name = slide.slide_layout.name
        except Exception:
            layout_name = None
        override = (kind_overrides or {}).get(slide_no)
        kind = override or classify_page_kind(
            slide, slide_no, total, layout_name=layout_name
        )
        kind_counter[kind] += 1

        layers = extract_layer_stack(slide, include_raw_xml=include_raw_xml)
        slide_shapes = [layer for layer in layers if layer.get("origin") == "slide"]

        for layer in layers:
            shape_types[str(layer.get("type"))] += 1
            text = layer.get("text") or {}
            for font in text.get("fonts", []):
                if font.get("name"):
                    fonts[font["name"]] += 1
                if font.get("size_pt"):
                    font_sizes[font["size_pt"]] += 1
            fill_rgb = (layer.get("style") or {}).get("fill", {}).get("rgb")
            if fill_rgb:
                fills[fill_rgb] += 1

        role = "first" if slide_no == 1 else "last" if slide_no == total else "body"
        pages.append({
            "slide": slide_no,
            "kind": kind,
            "role": role,
            "purpose": {"cover": "cover", "closing": "closing"}.get(kind, "content"),
            "layout_name": layout_name,
            "background": _background_dna(slide),
            "layout_background": _background_dna(slide.slide_layout)
            if getattr(slide, "slide_layout", None) is not None
            else {},
            "inheritance": {
                attr: _flag(getattr(slide, attr, None))
                for attr in (
                    "follow_master_graphics", "follow_master_background",
                    "show_master_shapes", "preserve",
                )
                if hasattr(slide, attr)
            },
            "layout_signature": {
                "shape_types": dict(
                    Counter(str(layer.get("type")) for layer in slide_shapes)
                ),
                "placeholders": dict(
                    Counter(
                        str((layer.get("placeholder") or {}).get("type"))
                        for layer in slide_shapes
                        if layer.get("placeholder")
                    )
                ),
                "layer_origins": dict(
                    Counter(str(layer.get("origin")) for layer in layers)
                ),
            },
            "layers": layers,
            "layer_summary": layer_stack_summary(layers),
            "shapes": slide_shapes,
            "raw_slide_xml": _xml_of(getattr(slide, "_element", None))
            if include_raw_xml
            else None,
        })

    masters = []
    for master in prs.slide_masters:
        layouts = [
            _layout_dna(layout, include_raw_xml=include_raw_xml)
            for layout in master.slide_layouts
        ]
        master_shapes = _safe_shapes(master)
        masters.append({
            "name": master.name,
            "background": _background_dna(master),
            "shapes": [
                extract_shape_dna(s, i, origin="master", render_order=i,
                                  include_raw_xml=include_raw_xml)
                for i, s in enumerate(master_shapes)
            ],
            "layouts": layouts,
            "raw_master_xml": _xml_of(getattr(master, "_element", None))
            if include_raw_xml
            else None,
        })

    page_kinds = aggregate_page_kinds(pages)

    return {
        "schema": SCHEMA,
        "source": str(source),
        "presentation": {
            "slide_size_inches": {
                "width": round(prs.slide_width / EMU_PER_INCH, 3),
                "height": round(prs.slide_height / EMU_PER_INCH, 3),
            },
            "slide_count": slide_count or total,
            "first_slide_role": "first" if pages else None,
            "last_slide_role": "last" if pages else None,
            "page_kind_counts": dict(kind_counter),
        },
        "theme": {
            "colors": _theme_colors(source),
            "font_scheme": _theme_font_scheme(source),
        },
        "dominant_palette": _dominant_palette(source),
        "global_style_statistics": {
            "fonts": fonts.most_common(20),
            "font_sizes_pt": font_sizes.most_common(20),
            "fills_rgb": fills.most_common(20),
            "shape_types": dict(shape_types),
        },
        "page_kinds": page_kinds,
        "masters": masters,
        "slides": pages,
        "special_surfaces": {
            "first": pages[0] if pages else None,
            "last": pages[-1] if pages else None,
            "body_slide_count": max(0, len(pages) - 2),
        },
    }


def _dominant_palette(path: Path) -> dict[str, Any]:
    try:
        from .palette import dominant_colors
        return dominant_colors(str(path))
    except Exception:
        return {}


def summarize_kind(kind_dna: dict[str, Any], *, max_ornaments: int = 12) -> str:
    """Human-readable layer-stack dump for a single page kind (CLI/debug)."""
    lines = [
        f"# {kind_dna.get('count', 0)} page(s) -- layouts: "
        f"{', '.join(kind_dna.get('layout_names') or []) or '-'}"
    ]
    for ornament in (kind_dna.get("ornaments") or [])[:max_ornaments]:
        geometry = ornament.get("geometry") or {}
        fill = (ornament.get("style") or {}).get("fill") or {}
        bits = [
            f"#{ornament.get('render_order')}",
            f"[{ornament.get('origin')}] {ornament.get('name')}",
            f"{geometry.get('left')},{geometry.get('top')} "
            f"{geometry.get('width')}x{geometry.get('height')}",
        ]
        if abs(float(geometry.get("rotation") or 0.0)) > 1e-6:
            bits.append(f"rot={geometry.get('rotation')}")
        if geometry.get("flip_horizontal") or geometry.get("flip_vertical"):
            bits.append(
                f"flip={'H' if geometry.get('flip_horizontal') else ''}"
                f"{'V' if geometry.get('flip_vertical') else ''}"
            )
        if geometry.get("rotation") or geometry.get("flip_horizontal") or geometry.get("flip_vertical"):
            bits.append(f"bbox={geometry.get('rendered_bbox')}")
        prst = (geometry.get("prst_geom") or {}).get("prst")
        if prst:
            bits.append(f"prst={prst}")
        if fill.get("type") == "gradient":
            bits.append(
                "grad="
                + ",".join(
                    f"{stop.get('pos')}:{stop.get('rgb') or stop.get('scheme')}"
                    f"@{stop.get('alpha')}"
                    for stop in fill.get("stops") or []
                )
                + f" ang={((fill.get('linear') or {}).get('angle_deg'))}"
            )
        elif fill.get("rgb"):
            bits.append(f"fill={fill.get('rgb')}@{fill.get('alpha')}")
        if ornament.get("picture", {}) and ornament["picture"].get("transparency"):
            bits.append(f"pic_alpha={ornament['picture'].get('alpha')}")
        if ornament.get("present_on") != kind_dna.get("count"):
            bits.append(f"on {ornament.get('present_on')}/{kind_dna.get('count')}")
        lines.append("  " + " ".join(bits))
    frequent = kind_dna.get("frequent") or []
    if frequent:
        lines.append("  -- variants (present on most pages) --")
        for ornament in frequent[:6]:
            geometry = ornament.get("geometry") or {}
            lines.append(
                f"  #{ornament.get('render_order')} {ornament.get('name')} "
                f"{geometry.get('left')},{geometry.get('top')} "
                f"{geometry.get('width')}x{geometry.get('height')} "
                f"({ornament.get('present_on')}/{ornament.get('total')})"
            )
    return "\n".join(lines)


def iter_deck_shapes(dna: dict[str, Any]) -> Iterator[dict[str, Any]]:
    """Every shape record across every page of a deck DNA dict."""
    for page in dna.get("slides") or dna.get("pages") or []:
        for layer in page.get("layers") or []:
            yield layer
            stack: list[dict[str, Any]] = list(layer.get("children") or [])
            while stack:
                child = stack.pop()
                yield child
                stack.extend(child.get("children") or [])


__all__ = [
    "SCHEMA",
    "PAGE_KINDS",
    "classify_page_kind",
    "extract_deck_dna",
    "extract_layer_stack",
    "extract_shape_dna",
    "aggregate_page_kinds",
    "layer_stack_summary",
    "parse_color",
    "parse_fill",
    "parse_line",
    "parse_effects",
    "parse_geometry",
    "parse_text",
    "parse_transform",
    "rendered_bbox",
    "summarize_kind",
    "iter_deck_shapes",
]
