"""Format parsers: source file -> Content IR.

Every parser is deterministic and returns a ``ContentDocument``; heavy
dependencies (python-docx / python-pptx) are imported lazily inside their
functions so the core import chain stays dependency-free (CI
``dependency-free-core`` gate).

``parse_auto`` dispatches on file extension.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from ..content_ir import ContentBlock, ContentDocument

_MD_HEADING_PREFIXES = ("#",)


def _next_id(prefix: str, counter: list[int]) -> str:
    counter[0] += 1
    return f"{prefix}-{counter[0]:04d}"


# --------------------------------------------------------------------------- #
# markdown
# --------------------------------------------------------------------------- #
def parse_markdown(text: str, *, source: str | None = None) -> ContentDocument:
    """Markdown -> ContentDocument (stdlib only).

    Supports ATX headings, bullets (- * +), ordered lists, pipe tables,
    blockquotes and paragraphs; fenced code blocks become ``quote`` blocks.
    """
    blocks: list[ContentBlock] = []
    counter = [0]
    lines = text.replace("\r\n", "\n").split("\n")
    index = 0
    title: str | None = None

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1  # skip closing fence
            blocks.append(ContentBlock(
                type="quote", id=_next_id("quote", counter),
                text="\n".join(code), source="markdown",
            ))
            continue

        if stripped.startswith(_MD_HEADING_PREFIXES):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading_text = stripped.lstrip("#").strip()
            if level == 1 and title is None:
                title = heading_text
            blocks.append(ContentBlock(
                type="heading", id=_next_id("h", counter),
                text=heading_text, level=min(level, 6), source="markdown",
            ))
            index += 1
            continue

        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip().lstrip(">").strip())
                index += 1
            blocks.append(ContentBlock(
                type="quote", id=_next_id("quote", counter),
                text="\n".join(quote_lines), source="markdown",
            ))
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and set(lines[index + 1].replace("|", "").replace("-", "").replace(":", "").strip()) <= {""} and "-" in lines[index + 1]:
            header_cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            blocks.append(ContentBlock(
                type="table", id=_next_id("tbl", counter),
                headers=header_cells, rows=rows, source="markdown",
            ))
            continue

        if stripped[:2] in ("- ", "* ") or stripped.startswith("+ "):
            items: list[str] = []
            while index < len(lines) and lines[index].strip()[:2] in ("- ", "* ", "+ "):
                items.append(lines[index].strip()[2:].strip())
                index += 1
            blocks.append(ContentBlock(
                type="bullets", id=_next_id("b", counter), items=items, source="markdown",
            ))
            continue

        if stripped[:1].isdigit() and ". " in stripped[:6]:
            ordered: list[str] = []
            while index < len(lines):
                item = lines[index].strip()
                if item[:1].isdigit() and ". " in item[:6]:
                    ordered.append(item.split(". ", 1)[1].strip())
                elif item[:1].isdigit() and ") " in item[:6]:
                    ordered.append(item.split(") ", 1)[1].strip())
                else:
                    break
                index += 1
            blocks.append(ContentBlock(
                type="ordered", id=_next_id("o", counter), items=ordered, source="markdown",
            ))
            continue

        paragraph_lines: list[str] = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if (not nxt or nxt.startswith(("#", ">", "|", "- ", "* ", "+ ", "```"))
                    or (nxt[:1].isdigit() and ". " in nxt[:6])):
                break
            paragraph_lines.append(nxt)
            index += 1
        blocks.append(ContentBlock(
            type="paragraph", id=_next_id("p", counter),
            text="\n".join(paragraph_lines), source="markdown",
        ))

    return ContentDocument(
        source=source, format="markdown", title=title, blocks=blocks,
        metadata={"parser": "markdown/stdlib"},
    )


# --------------------------------------------------------------------------- #
# docx (python-docx, lazily imported -- extras only)
# --------------------------------------------------------------------------- #
def parse_docx(path: str | Path) -> ContentDocument:
    """Word document -> ContentDocument. Requires ``pip install python-docx``."""
    try:
        import docx  # type: ignore[import-untyped]  # optional extra
    except ImportError as exc:  # pragma: no cover - optional extra
        raise RuntimeError(
            "python-docx is required for .docx parsing; "
            "install with pip install 'ppt-agent[docx]'"
        ) from exc

    document = docx.Document(str(path))
    blocks: list[ContentBlock] = []
    counter = [0]
    title: str | None = None

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = str(getattr(paragraph.style, "name", "") or "")
        if style_name.startswith("Heading"):
            try:
                level = int(style_name.split()[-1])
            except ValueError:
                level = 2
            if level == 1 and title is None:
                title = text
            blocks.append(ContentBlock(
                type="heading", id=_next_id("h", counter),
                text=text, level=min(level, 6), source="docx",
            ))
        elif style_name.startswith("List Bullet"):
            _append_list_item(blocks, counter, "bullets", text)
        elif style_name.startswith("List Number"):
            _append_list_item(blocks, counter, "ordered", text)
        else:
            blocks.append(ContentBlock(
                type="paragraph", id=_next_id("p", counter),
                text=text, source="docx",
            ))

    for table in document.tables:
        rows = [
            [cell.text.strip() for cell in row.cells]
            for row in table.rows
        ]
        headers = rows.pop(0) if rows else []
        blocks.append(ContentBlock(
            type="table", id=_next_id("tbl", counter),
            headers=headers, rows=rows, source="docx",
        ))

    return ContentDocument(
        source=str(path), format="docx", title=title, blocks=blocks,
        metadata={"parser": "docx/python-docx"},
    )


def _append_list_item(
    blocks: list[ContentBlock], counter: list[int], kind: str, text: str
) -> None:
    if blocks and blocks[-1].type == kind and blocks[-1].source == "docx":
        blocks[-1].items.append(text)
    else:
        blocks.append(ContentBlock(
            type=kind, id=_next_id("b" if kind == "bullets" else "o", counter),
            items=[text], source="docx",
        ))


# --------------------------------------------------------------------------- #
# pptx (python-pptx, lazily imported -- extras only)
# --------------------------------------------------------------------------- #
def parse_pptx(path: str | Path) -> ContentDocument:
    """Source deck -> ContentDocument (one heading + body blocks per slide)."""
    try:
        from pptx import Presentation  # type: ignore[import-untyped]  # optional extra
    except ImportError as exc:  # pragma: no cover - optional extra
        raise RuntimeError(
            "python-pptx is required for .pptx parsing; "
            "install with pip install 'ppt-agent[pptx]'"
        ) from exc

    prs = Presentation(str(path))
    blocks: list[ContentBlock] = []
    counter = [0]
    title: str | None = None

    for slide_no, slide in enumerate(prs.slides, 1):
        slide_texts: list[str] = []
        slide_title: str | None = None
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if slide_title is None:
                slide_title = text
            else:
                slide_texts.extend(part.strip() for part in text.split("\n") if part.strip())
            if slide_no == 1 and title is None:
                title = text
        if slide_title:
            blocks.append(ContentBlock(
                type="heading", id=_next_id("h", counter),
                text=slide_title, level=2, source="pptx",
            ))
        if slide_texts:
            blocks.append(ContentBlock(
                type="bullets", id=_next_id("b", counter),
                items=slide_texts, source="pptx",
            ))

    return ContentDocument(
        source=str(path), format="pptx", title=title, blocks=blocks,
        metadata={"parser": "pptx/python-pptx"},
    )


# --------------------------------------------------------------------------- #
# auto dispatch
# --------------------------------------------------------------------------- #
_PARSERS: dict[str, Any] = {
    ".md": ("text", parse_markdown),
    ".markdown": ("text", parse_markdown),
    ".docx": ("path", parse_docx),
    ".pptx": ("path", parse_pptx),
}


def parse_auto(path: str | Path) -> ContentDocument:
    """Parse any supported source file by extension."""
    source = Path(path)
    entry = _PARSERS.get(source.suffix.lower())
    if entry is None:
        raise ValueError(f"unsupported source format: {source.suffix!r}")
    mode, parser = entry
    if mode == "text":
        return parser(source.read_text(encoding="utf-8"), source=str(source))
    return parser(source)
