from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from .ir import Component, Presentation
from .styling import (
    DEFAULT_SLIDE_H_IN,
    DEFAULT_SLIDE_W_IN,
    EMU_PER_INCH,
    FLOW_GAP_IN,
    MARGIN_IN,
    TEXT_COMPONENT_TYPES,
    TITLE_PURPOSES,
    default_font_size,
    estimate_flow_height,
    font_style_of,
    group_children,
    has_geometry,
    normalize_alignment,
    normalize_color,
    normalize_opacity,
    resolve_layout,
    slide_size_inches,
)

# Backwards-compatible private aliases: these names shipped in V0.3.
_TEXT_TYPES = TEXT_COMPONENT_TYPES
_TITLE_PURPOSES = TITLE_PURPOSES

try:  # pragma: no cover - import guard is environment dependent
    from pptx import Presentation as _PptxPresentation
    from pptx.dml.color import RGBColor as _RGBColor
    from pptx.enum.shapes import MSO_SHAPE as _MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR as _MSO_ANCHOR
    from pptx.enum.text import PP_ALIGN as _PP_ALIGN
    from pptx.oxml.ns import qn as _qn
    from pptx.util import Inches as _Inches
    from pptx.util import Pt as _Pt

    _PPTX_AVAILABLE = True
except ImportError:  # pragma: no cover
    _PPTX_AVAILABLE = False


def _rgb(value: Any):
    if not _PPTX_AVAILABLE or value is None:
        return None
    text = normalize_color(value)
    if text is None:
        return None
    try:
        return _RGBColor.from_string(text)
    except (ValueError, TypeError):
        return None


def _opacity(fill: Any) -> float | None:
    return normalize_opacity(fill)


def _set_solid_alpha(shape: Any, opacity: float) -> None:
    sp_pr = shape._element.spPr
    solid = sp_pr.find(_qn("a:solidFill"))
    if solid is None:
        return
    clr = solid.find(_qn("a:srgbClr"))
    if clr is None:
        clr = solid.find(_qn("a:schemeClr"))
    if clr is None:
        return
    for existing in clr.findall(_qn("a:alpha")):
        clr.remove(existing)
    alpha = clr.makeelement(_qn("a:alpha"), {"val": str(int(max(0.0, min(1.0, opacity)) * 100000))})
    clr.append(alpha)


def _apply_fill(shape: Any, fill: Any) -> None:
    if not isinstance(fill, dict) or not fill:
        return
    fill_type = str(fill.get("type") or "").lower()
    rgb = _rgb(fill.get("rgb"))
    if rgb is None:
        if fill_type in ("none", "background") or fill.get("rgb") is None:
            try:
                shape.fill.background()
            except Exception:
                pass
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    opacity = _opacity(fill)
    if opacity is not None and opacity < 1.0:
        _set_solid_alpha(shape, opacity)


def _apply_line(shape: Any, line: Any) -> None:
    if not isinstance(line, dict) or not line:
        # python-pptx leaves the default theme outline on new autoshapes; an
        # unspecified line must mean "no outline", or every decoration gets a
        # blue hairline border.
        try:
            shape.line.fill.background()
        except Exception:
            pass
        return
    rgb = _rgb(line.get("rgb"))
    width = line.get("width_pt")
    if rgb is None and not isinstance(width, (int, float)):
        return
    try:
        if rgb is not None:
            shape.line.color.rgb = rgb
        if isinstance(width, (int, float)) and width > 0:
            shape.line.width = _Pt(float(width))
    except Exception:
        pass


def _apply_text(shape: Any, comp: Component, default_size_pt: float) -> None:
    text = comp.text
    if text is None and isinstance(comp.data, dict):
        text = comp.data.get("text")
    text = "" if text is None else str(text)
    font_style = font_style_of(comp)
    size_pt = font_style.get("size_pt", default_size_pt)
    tf = shape.text_frame
    tf.word_wrap = True
    try:
        tf.vertical_anchor = _MSO_ANCHOR.TOP
    except Exception:
        pass
    lines = text.split("\n")
    tf.text = lines[0] if lines else ""
    for line in lines[1:]:
        tf.add_paragraph().text = line
    alignment = None
    normalized = normalize_alignment(font_style.get("alignment"))
    if normalized:
        alignment = {
            "left": _PP_ALIGN.LEFT,
            "center": _PP_ALIGN.CENTER,
            "right": _PP_ALIGN.RIGHT,
            "justify": _PP_ALIGN.JUSTIFY,
        }[normalized]
    font_rgb = _rgb(font_style.get("rgb"))
    for paragraph in tf.paragraphs:
        if alignment is not None:
            paragraph.alignment = alignment
        for run in paragraph.runs:
            font = run.font
            if isinstance(size_pt, (int, float)) and size_pt > 0:
                font.size = _Pt(float(size_pt))
            if isinstance(font_style.get("bold"), bool):
                font.bold = font_style["bold"]
            if isinstance(font_style.get("italic"), bool):
                font.italic = font_style["italic"]
            if isinstance(font_style.get("name"), str):
                font.name = font_style["name"]
            if font_rgb is not None:
                font.color.rgb = font_rgb


def _has_geometry(comp: Component) -> bool:
    return has_geometry(comp)


def _box(comp: Component) -> tuple[Any, Any, Any, Any]:
    return _Inches(float(comp.x)), _Inches(float(comp.y)), _Inches(max(float(comp.w), 0.05)), _Inches(max(float(comp.h), 0.05))


def _estimate_flow_height(comp: Component, width_in: float, default_size_pt: float) -> float:
    return estimate_flow_height(comp, width_in, default_size_pt)


def _add_placeholder(slide: Any, box: tuple[Any, Any, Any, Any], label: str) -> Any:
    shape = slide.shapes.add_shape(_MSO_SHAPE.RECTANGLE, *box)
    try:
        shape.fill.background()
        shape.line.color.rgb = _RGBColor(0xAA, 0xAA, 0xAA)
        shape.line.width = _Pt(1)
    except Exception:
        pass
    tf = shape.text_frame
    tf.text = label
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = _Pt(11)
            run.font.color.rgb = _RGBColor(0x88, 0x88, 0x88)
    return shape


def _render_image(slide: Any, comp: Component, box: tuple[Any, Any, Any, Any]) -> Any:
    data = comp.data if isinstance(comp.data, dict) else {}
    path = data.get("path") if isinstance(data.get("path"), str) else None
    if path and Path(path).exists():
        return slide.shapes.add_picture(str(path), *box)
    blob = data.get("bytes") or data.get("image") or data.get("blob")
    if isinstance(blob, (bytes, bytearray)):
        return slide.shapes.add_picture(io.BytesIO(bytes(blob)), *box)
    return _add_placeholder(slide, box, f"[image] {data.get('name') or comp.id or ''}".strip())


def _render_table(slide: Any, comp: Component, box: tuple[Any, Any, Any, Any]) -> Any:
    data = comp.data if isinstance(comp.data, dict) else {}
    rows = data.get("rows")
    if isinstance(rows, list) and rows and all(isinstance(r, (list, tuple)) for r in rows):
        cols = max(len(r) for r in rows)
        if cols == 0:
            return _add_placeholder(slide, box, "[table]")
        table = slide.shapes.add_table(len(rows), cols, *box).table
        for r, row in enumerate(rows):
            for c in range(cols):
                cell = table.cell(r, c)
                cell.text = str(row[c]) if c < len(row) else ""
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = _Pt(12)
        return table
    return _add_placeholder(slide, box, f"[table] {data.get('name') or ''}".strip())


def _render_component(slide: Any, comp: Component, box: tuple[Any, Any, Any, Any], default_size_pt: float) -> None:
    ctype = (comp.type or "shape").lower()
    if ctype in _TEXT_TYPES:
        shape = slide.shapes.add_textbox(*box)
        _apply_text(shape, comp, default_size_pt)
        return
    if ctype == "image":
        _render_image(slide, comp, box)
        return
    if ctype == "table":
        _render_table(slide, comp, box)
        return
    if ctype == "chart":
        _add_placeholder(slide, box, f"[chart] {(comp.data or {}).get('name') if isinstance(comp.data, dict) else ''}".strip())
        return
    if ctype == "group":
        children = [child for child in group_children(comp) if _has_geometry(child)]
        if children:
            for child in children:
                _render_component(slide, child, _box(child), default_size_pt)
            return
        _add_placeholder(slide, box, f"[group] {comp.id or ''}".strip())
        return
    # generic shape / unknown type -> autoshape with fill/line
    style = comp.style if isinstance(comp.style, dict) else {}
    preset = _MSO_SHAPE.RECTANGLE
    if str(style.get("shape") or "").lower() in ("ellipse", "oval", "circle"):
        preset = _MSO_SHAPE.OVAL
    shape = slide.shapes.add_shape(preset, *box)
    fill = style.get("fill")
    if isinstance(fill, dict) and fill:
        _apply_fill(shape, fill)
    else:
        try:
            shape.fill.background()
        except Exception:
            pass
    _apply_line(shape, style.get("line"))
    if comp.text:
        _apply_text(shape, comp, default_size_pt)


def _default_size(comp: Component, slide_purpose: str) -> float:
    return default_font_size(comp, slide_purpose)


def _apply_background(slide: Any, slide_data: Any) -> None:
    if not isinstance(slide_data, dict):
        return
    background = slide_data.get("background")
    if not isinstance(background, dict):
        return
    fill = background.get("fill") if isinstance(background.get("fill"), dict) else background
    rgb = _rgb(fill.get("rgb")) if isinstance(fill, dict) else None
    if rgb is None:
        return
    try:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb
        opacity = _opacity(fill)
        if opacity is not None and opacity < 1.0:
            _set_solid_alpha(slide.background, opacity)
    except Exception:
        pass


def _blank_layout(prs: Any) -> Any:
    for layout in prs.slide_layouts:
        if (layout.name or "").strip().lower() == "blank":
            return layout
    return prs.slide_layouts[min(6, len(prs.slide_layouts) - 1)]


def _slide_size(presentation: Presentation) -> tuple[float, float]:
    return slide_size_inches(presentation)


def render_presentation(presentation: Presentation, output: str | Path, *, iteration: int | None = None) -> Path:
    """Render a Universal Presentation IR into an editable native PPTX file."""
    if not _PPTX_AVAILABLE:
        raise RuntimeError("python-pptx is required for rendering; install with pip install 'ppt-agent[pptx]'")

    from .design import design_presentation

    presentation = design_presentation(presentation)
    width_in, height_in = _slide_size(presentation)
    prs = _PptxPresentation()
    prs.slide_width = _Inches(width_in)
    prs.slide_height = _Inches(height_in)
    layout = _blank_layout(prs)

    for slide_ir in presentation.slides:
        slide = prs.slides.add_slide(layout)
        _apply_background(slide, slide_ir.data)
        for box in resolve_layout(slide_ir, width_in, height_in):
            _render_component(
                slide,
                box.component,
                (_Inches(box.x), _Inches(box.y), _Inches(box.w), _Inches(box.h)),
                box.font_pt,
            )

        if slide_ir.speaker_notes:
            slide.notes_slide.notes_text_frame.text = slide_ir.speaker_notes

    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destination))
    return destination


def make_ir_build(presentation: Presentation, output: str | Path):
    """Return a ``build`` callback compatible with ``delivery.run_repair_loop``."""

    def build(iteration: int, repair_requests: list[dict[str, Any]]) -> Path:
        return render_presentation(presentation, output, iteration=iteration)

    return build
