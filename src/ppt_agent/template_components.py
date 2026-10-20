"""Template component catalog and deterministic page-component ranking."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any, Iterable

SCHEMA = "template-component-catalog/v1"

_KEYWORDS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("org_chart", ("组织", "保障", "职责", "分工", "团队", "人员", "架构", "协同")),
    ("timeline", ("时间", "节点", "阶段", "推进", "计划", "里程碑", "月份", "旬")),
    ("architecture", ("技术架构", "分层", "平台", "体系", "集成", "底座", "系统")),
    ("kpi", ("指标", "进度", "完成", "比例", "成效", "数量", "市场")),
    ("cards", ("功能", "背景", "目标", "原则", "承诺", "风险", "能力")),
)


def _text_shape(sh: Any) -> bool:
    return bool(getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip())


def _font_size(sh: Any) -> float:
    sizes = [
        run.font.size.pt
        for paragraph in sh.text_frame.paragraphs
        for run in paragraph.runs
        if run.font.size
    ]
    return max(sizes) if sizes else 0.0


def _box(sh: Any) -> dict[str, float]:
    return {
        "left": round((sh.left or 0) / 914400, 4),
        "top": round((sh.top or 0) / 914400, 4),
        "width": round((sh.width or 0) / 914400, 4),
        "height": round((sh.height or 0) / 914400, 4),
    }


def _layout_kind(layout_name: str, index: int) -> str:
    name = (layout_name or "").lower()
    if index == 1 or any(token in name for token in ("封面", "title slide")):
        return "cover"
    if any(token in name for token in ("封底", "结束", "thank", "closing")):
        return "closing"
    if any(token in name for token in ("章节", "divider", "section")):
        return "section"
    if any(token in name for token in ("目录", "contents", "agenda")):
        return "toc"
    return "content"


def _semantic_kind(layout_kind: str, title: str, text: str, *, table_count: int) -> str:
    if layout_kind in {"cover", "closing", "section", "toc"}:
        return layout_kind
    haystack = f"{title}\n{text}".lower()
    if table_count and any(token in haystack for token in ("分工", "责任", "矩阵", "表")):
        return "table"
    for kind, words in _KEYWORDS:
        if any(word.lower() in haystack for word in words):
            return kind
    if table_count:
        return "table"
    return "cards"


def _tags(title: str, text: str) -> list[str]:
    haystack = f"{title}\n{text}"
    tags: list[str] = []
    for _kind, words in _KEYWORDS:
        for word in words:
            if word in haystack and word not in tags:
                tags.append(word)
    return tags


def _text_capacity(shape: Any, font_size: float) -> int:
    """Conservative CJK character capacity for one template text slot."""
    width = float((shape.width or 0) / 914400)
    height = float((shape.height or 0) / 914400)
    size = max(8.0, font_size or 10.0)
    chars_per_line = max(1, int(width * 72 / (size * 1.45)))
    lines = max(1, int(height * 72 / (size * 1.65)))
    return chars_per_line * lines


def build_template_component_catalog(
    template: str | Path,
    *,
    dna: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Catalog every template page, including one-off visual compositions."""
    from pptx import Presentation

    source = Path(template).expanduser().resolve()
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    prs = Presentation(str(source))
    dna_pages = (dna or {}).get("slides") or []
    components: list[dict[str, Any]] = []

    for index, slide in enumerate(prs.slides, start=1):
        text_shapes = [shape for shape in slide.shapes if _text_shape(shape)]
        text_shapes.sort(key=_font_size, reverse=True)
        title = text_shapes[0].text_frame.text.strip() if text_shapes else ""
        all_text = "\n".join(shape.text_frame.text.strip() for shape in text_shapes)
        table_count = sum(
            1 for shape in slide.shapes if getattr(shape, "has_table", False)
        )
        picture_count = sum(1 for shape in slide.shapes if shape.shape_type == 13)
        group_count = sum(1 for shape in slide.shapes if shape.shape_type == 6)
        page_kind = (
            str((dna_pages[index - 1] if index <= len(dna_pages) else {}).get("kind"))
            if dna
            else _layout_kind(slide.slide_layout.name, index)
        )
        if page_kind not in {"cover", "toc", "section", "content", "closing"}:
            page_kind = _layout_kind(slide.slide_layout.name, index)
        component_kind = _semantic_kind(
            page_kind, title, all_text, table_count=table_count,
        )
        slots = []
        for shape_index, shape in enumerate(slide.shapes):
            if not _text_shape(shape):
                continue
            bbox = _box(shape)
            slots.append({
                "shape_index": shape_index,
                "shape_id": int(getattr(shape, "shape_id", 0) or 0),
                "name": getattr(shape, "name", ""),
                "text": shape.text_frame.text.strip(),
                "role": "title" if shape is text_shapes[0] else "body",
                "font_size_pt": round(_font_size(shape), 2),
                "text_capacity": _text_capacity(shape, _font_size(shape)),
                **bbox,
            })
        body_slots = [slot for slot in slots if slot["role"] == "body"]
        components.append({
            "component_id": f"tpl-{index:02d}-{component_kind}",
            "template_slide": index,
            "page_kind": page_kind,
            "component_kind": component_kind,
            "title": title[:100],
            "tags": _tags(title, all_text),
            "slot_count": len(slots),
            "item_capacity": max(1, len(body_slots) or len(slots)),
            "text_capacity": sum(slot["text_capacity"] for slot in body_slots),
            "slots": slots,
            "table_count": table_count,
            "picture_count": picture_count,
            "group_count": group_count,
            "shape_count": len(slide.shapes),
            "copy_policy": "slide_clone",
            "source": "template",
        })

    return {
        "schema": SCHEMA,
        "template": str(source),
        "template_hash": digest,
        "component_count": len(components),
        "components": components,
    }


def rank_template_components(
    catalog: dict[str, Any],
    *,
    component_kind: str | None = None,
    tags: Iterable[str] = (),
    item_count: int = 0,
    text_chars: int = 0,
    exclude_slides: Iterable[int] = (),
    limit: int = 5,
) -> list[dict[str, Any]]:
    """Rank template candidates using role, tags, capacity and table fit."""
    wanted_tags = {str(tag) for tag in tags if tag}
    excluded = {int(value) for value in exclude_slides}
    ranked: list[dict[str, Any]] = []
    for component in catalog.get("components") or []:
        slide_index = int(component.get("template_slide") or 0)
        if slide_index in excluded:
            continue
        score = 0.0
        reasons: list[str] = []
        kind = str(component.get("component_kind") or "")
        if kind not in {"cover", "section", "closing", "toc"}:
            if str(component.get("page_kind") or "") != "content":
                continue
            if int(component.get("slot_count") or 0) < 2:
                continue
        if component_kind and kind == component_kind:
            score += 0.55
            reasons.append(f"kind={component_kind}")
        overlap = wanted_tags & set(component.get("tags") or [])
        if overlap:
            score += min(0.2, 0.05 * len(overlap))
            reasons.append("tags=" + ",".join(sorted(overlap)))
        capacity = int(component.get("item_capacity") or 0)
        if item_count and capacity:
            if capacity == item_count:
                score += 0.15
                reasons.append("capacity=exact")
            elif abs(capacity - item_count) <= 1:
                score += 0.08
                reasons.append("capacity=near")
        text_capacity = int(component.get("text_capacity") or 0)
        if text_chars and text_capacity:
            ratio = text_chars / text_capacity
            if ratio > 0.75:
                continue
            if ratio <= 0.5:
                score += 0.12
                reasons.append("text_capacity=good")
            elif ratio <= 0.75:
                score += 0.04
                reasons.append("text_capacity=borderline")
            else:
                score -= 0.35
                reasons.append("text_capacity=too_small")
        if component_kind == "table" and int(component.get("table_count") or 0):
            score += 0.1
            reasons.append("native_table")
        if score > 0:
            ranked.append({
                "component_id": component.get("component_id"),
                "template_slide": slide_index,
                "component_kind": kind,
                "title": component.get("title"),
                "score": round(min(score, 0.95), 3),
                "slot_count": component.get("slot_count"),
                "item_capacity": capacity,
                "text_capacity": text_capacity,
                "table_count": component.get("table_count"),
                "reasons": reasons,
                "source": "template",
            })
    ranked = [item for item in ranked if item["score"] >= 0.5]
    ranked.sort(key=lambda item: (-item["score"], item["template_slide"]))
    return ranked[: max(1, int(limit))]


__all__ = [
    "SCHEMA",
    "build_template_component_catalog",
    "rank_template_components",
]
