"""Editability policy for generated decks.

Generated slide content must stay native and editable.  In particular, a
picture must never stand in for text, a table, a chart or a screenshot taken
from a template/reference file.  Template-native media inherited through a
shell (for example a cover background) is not a violation by itself; pages
still need native editable content.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

SCHEMA = "editability-report/v1"

PROHIBITED_IMAGE_TYPES = frozenset({
    "image", "picture", "screenshot", "snapshot", "image_block", "photo",
})

_PROHIBITED_PICTURE_LABELS = (
    "screenshot", "screen shot", "snapshot", "capture", "reference image",
    "template image", "参考截图", "模板截图", "资料截图",
)
_MEDIA_KEY_MARKERS = ("image", "picture", "screenshot", "snapshot", "photo")
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}


def audit_pptx_editability(pptx: str | Path) -> dict[str, Any]:
    """Audit the saved deck without treating inherited media as content."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(pptx))
    pages: list[dict[str, Any]] = []
    total_editable = total_pictures = 0

    for index, slide in enumerate(prs.slides, 1):
        shapes = list(_walk_shapes(slide.shapes, MSO_SHAPE_TYPE.GROUP))
        text_shapes = sum(
            1 for shape in shapes
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()
        )
        table_count = sum(1 for shape in shapes if _is_table(shape))
        chart_count = sum(1 for shape in shapes if _is_chart(shape))
        picture_shapes = [
            shape for shape in shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        editable_shapes = [
            shape for shape in shapes
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE
        ]
        issues: list[str] = []
        if not editable_shapes:
            issues.append("page has no native editable shape")
        if not text_shapes and not table_count and not chart_count:
            issues.append("page has no native editable text, table, or chart")
        issues.extend(_picture_label_issues(picture_shapes, index))

        pages.append({
            "page": index,
            "passed": not issues,
            "issues": issues,
            "native_editable_shape_count": len(editable_shapes),
            "native_text_shape_count": text_shapes,
            "native_table_count": table_count,
            "native_chart_count": chart_count,
            "slide_local_picture_count": len(picture_shapes),
        })
        total_editable += len(editable_shapes)
        total_pictures += len(picture_shapes)

    return {
        "schema": SCHEMA,
        "policy": "all generated content is native and editable; no screenshot paste",
        "passed": bool(pages) and all(page["passed"] for page in pages),
        "slide_count": len(pages),
        "pages": pages,
        "native_editable_shape_count": total_editable,
        "slide_local_picture_count": total_pictures,
    }


def validate_presentation_editability(presentation: Any) -> dict[str, Any]:
    """Validate the presentation IR before it reaches a renderer."""
    pages: list[dict[str, Any]] = []
    for page_no, slide in enumerate(presentation.slides, 1):
        issues: list[str] = []
        editable = 0
        for component in _walk_components(slide.components):
            ctype = str(component.type or "").lower()
            if ctype in PROHIBITED_IMAGE_TYPES:
                issues.append(
                    f"component {component.id or ctype!r} is {ctype!r}; "
                    "use native text, table, shape, or chart content"
                )
            else:
                editable += 1
        if not editable:
            issues.append("page has no native editable component")
        pages.append({
            "page": page_no,
            "slide_id": slide.id,
            "passed": not issues,
            "issues": issues,
            "native_editable_component_count": editable,
        })

    return {
        "schema": SCHEMA,
        "policy": "all generated content is native and editable; no screenshot paste",
        "passed": bool(pages) and all(page["passed"] for page in pages),
        "slide_count": len(pages),
        "pages": pages,
    }


def validate_page_plan_editability(pages: list[dict[str, Any]]) -> None:
    """Reject page-plan fields that request raster slide content."""
    for page_no, spec in enumerate(pages, 1):
        _check_editable_payload(spec, [f"pages[{page_no}]"])


def _walk_shapes(shapes: Any, group_type: Any) -> Any:
    for shape in shapes:
        yield shape
        if shape.shape_type == group_type:
            yield from _walk_shapes(shape.shapes, group_type)


def _walk_components(components: Any) -> Any:
    from .styling import group_children

    for component in components:
        yield component
        if str(component.type or "").lower() == "group":
            yield from _walk_components(group_children(component))


def _is_table(shape: Any) -> bool:
    try:
        return bool(shape.has_table)
    except (AttributeError, ValueError):
        return False


def _is_chart(shape: Any) -> bool:
    try:
        return bool(shape.has_chart)
    except (AttributeError, ValueError):
        return False


def _picture_label_issues(pictures: list[Any], page: int) -> list[str]:
    issues: list[str] = []
    for shape in pictures:
        labels = [str(shape.name or "")]
        try:
            c_nv_pr = shape._element._nvXxPr.cNvPr
            labels.extend([c_nv_pr.get("descr", ""), c_nv_pr.get("title", "")])
        except AttributeError:
            pass
        label = " ".join(labels).lower()
        if any(marker in label for marker in _PROHIBITED_PICTURE_LABELS):
            issues.append(
                f"page {page} picture {shape.shape_id} "
                f"({str(shape.name or 'unnamed')!r}) appears to be a "
                "screenshot/template/reference image; replace it with native content"
            )
    return issues


def _check_editable_payload(value: Any, path: list[str]) -> None:
    if isinstance(value, (bytes, bytearray, memoryview)):
        raise ValueError(
            f"{'.'.join(path)} contains binary media; generated slide content "
            "must be native and editable"
        )
    if isinstance(value, dict):
        for key, item in value.items():
            key_text = str(key).lower()
            next_path = [*path, str(key)]
            if item is not None and any(marker in key_text for marker in _MEDIA_KEY_MARKERS):
                raise ValueError(
                    f"{'.'.join(next_path)} requests non-editable media; "
                    "use native text, table, shape, or chart content"
                )
            if key_text == "path" and isinstance(item, (str, Path)):
                suffix = Path(str(item)).suffix.lower()
                if suffix in _IMAGE_SUFFIXES:
                    raise ValueError(
                        f"{'.'.join(next_path)} points to a raster image; "
                        "replace screenshots with editable native content"
                    )
            _check_editable_payload(item, next_path)
    elif isinstance(value, (list, tuple)):
        for index, item in enumerate(value):
            _check_editable_payload(item, [*path, str(index)])


__all__ = [
    "PROHIBITED_IMAGE_TYPES",
    "SCHEMA",
    "audit_pptx_editability",
    "validate_presentation_editability",
    "validate_page_plan_editability",
]
